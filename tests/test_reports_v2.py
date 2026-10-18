"""Reports engine (services/reports_service.py) — numbers, consistency, permissions, isolation.

One October 2026 scenario covers: cash sale, credit sale, partial payment with a
discount, full return, cancellation, an unbooked order, posted and unposted
expenses, cash and credit purchases, a cancelled purchase, deposits, a supplier
payment, an expense withdrawal and an internal transfer. Every dashboard
figure must equal its detail report and FINORA's own accounting helpers.
"""
from __future__ import annotations

import os
import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

TENANT = f"test_reports_v2_{os.getpid()}"
OTHER_TENANT = f"test_reports_v2_other_{os.getpid()}"
OCT = {"period": "custom", "date_from": "2026-10-01", "date_to": "2026-10-31"}
JSON = {"Accept": "application/json", "X-Requested-With": "XMLHttpRequest"}


def _wipe() -> None:
    try:
        from extensions import db

        db.session.remove()
    except Exception:
        pass
    for slug in (TENANT, OTHER_TENANT):
        try:
            from extensions_tenant import clear_tenant_engine

            clear_tenant_engine(slug)
        except Exception:
            pass
        db_file = ROOT / "tenants" / f"{slug}.db"
        if db_file.exists():
            db_file.unlink()


def _grant(employee, *names):
    from models.role import Permission, Role

    role = Role(name=f"reports-role-{employee.username}")
    for name in names:
        role.permissions.append(Permission.query.filter_by(name=name).first() or Permission(name=name, description=name))
    employee.roles.append(role)


def _seed_main() -> dict:
    from flask import g
    from werkzeug.security import generate_password_hash

    from extensions import db
    from extensions_tenant import init_tenant_db
    from models.account_transaction import AccountTransaction
    from models.branch import Branch, BranchStock
    from models.customer import Customer
    from models.employee import Employee
    from models.expense import Expense
    from models.invoice import Invoice
    from models.order_item import OrderItem
    from models.product import Product
    from models.purchase import Purchase
    from models.supplier import Supplier
    from models.treasury_account import TreasuryAccount
    from utils.branch_migration import get_default_branch
    from utils.expense_posting import build_expense_withdraw_note
    from utils.payment_ledger import append_payment_ledger_delta
    from utils.treasury_calculations import record_treasury_transfer

    g.tenant = TENANT
    init_tenant_db(TENANT)
    branch = get_default_branch()
    other_branch = Branch(name="فرع البصرة", code="BSR", is_active=True)
    db.session.add(other_branch)
    db.session.flush()
    pw = generate_password_hash("x-pass-123")
    admin = Employee(name="Admin", username="rp_admin", password=pw, role="admin", is_active=True, branch_id=branch.id)
    seller = Employee(name="سارة", username="rp_seller", password=pw, role="sales", is_active=True, branch_id=branch.id)
    viewer = Employee(name="Viewer", username="rp_viewer", password=pw, role="sales", is_active=True, branch_id=branch.id)
    outsider = Employee(name="Outsider", username="rp_out", password=pw, role="sales", is_active=True, branch_id=branch.id)
    db.session.add_all([admin, seller, viewer, outsider])
    db.session.flush()
    _grant(viewer, "view_reports")
    customer = Customer(branch_id=branch.id, name="أحمد سالم", phone="07701234567", city="بغداد")
    supplier = Supplier(branch_id=branch.id, name="شركة الرافدين")
    db.session.add_all([customer, supplier])
    db.session.flush()
    product = Product(branch_id=branch.id, name="شاشة LG", sku="TV-LG", buy_price=100000, sale_price=150000,
                      quantity=10, low_stock_threshold=2, catalog_category="شاشات", active=True)
    stale = Product(branch_id=branch.id, name="مكيف قديم", sku="AC-OLD", buy_price=80000, sale_price=120000,
                    quantity=4, low_stock_threshold=1, active=True, created_at=datetime(2026, 1, 1))
    db.session.add_all([product, stale])
    db.session.flush()
    for p in (product, stale):
        db.session.add(BranchStock(branch_id=branch.id, product_id=p.id, quantity=p.quantity, opening_stock=p.quantity))

    def invoice(total, status, payment, *, paid=0, discount=0, day=2, employee=admin, deducted=True, qty=1):
        inv = Invoice(customer_id=customer.id, customer_name=customer.name, branch_id=branch.id,
                      employee_id=employee.id, employee_name=employee.name, total=total, discount_amount=discount,
                      paid_amount=paid, status=status, payment_status=payment, stock_is_deducted=deducted,
                      sale_type="cash" if payment == "مسدد" else "credit", created_at=datetime(2026, 10, day, 11, 0))
        db.session.add(inv)
        db.session.flush()
        db.session.add(OrderItem(invoice_id=inv.id, product_id=product.id, product_name=product.name,
                                 price=total // qty, quantity=qty, cost=100000, total=total))
        return inv

    cash_sale = invoice(150000, "تم التوصيل", "مسدد", paid=150000, day=2)
    credit_sale = invoice(300000, "تم التوصيل", "غير مسدد", day=3, employee=seller, qty=2)
    partial = invoice(140000, "تم التوصيل", "جزئي", paid=40000, discount=10000, day=4, employee=seller)
    invoice(150000, "راجع", "مرتجع", day=5)
    invoice(150000, "ملغي", "غير مسدد", day=6, deducted=False)
    invoice(150000, "تم الطلب", "غير مسدد", day=7, deducted=False)

    db.session.add_all([
        Expense(branch_id=branch.id, title="إيجار المحل", category="إيجار", amount=50000, cash_posted=True,
                expense_date=datetime(2026, 10, 4), employee_id=admin.id),
        Expense(branch_id=branch.id, title="مسودة", category="أخرى", amount=20000, cash_posted=False,
                expense_date=datetime(2026, 10, 4)),
    ])
    db.session.add_all([
        Purchase(supplier_id=supplier.id, product_id=product.id, quantity=2, price=100000, total=200000,
                 grand_total=200000, paid_total=200000, remaining_total=0, status="confirmed", branch_id=branch.id,
                 purchase_date=date(2026, 10, 3)),
        Purchase(supplier_id=supplier.id, product_id=product.id, quantity=1, price=100000, total=100000,
                 grand_total=100000, paid_total=0, remaining_total=100000, status="confirmed", branch_id=branch.id,
                 purchase_mode="credit", purchase_date=date(2026, 10, 6)),
        Purchase(supplier_id=supplier.id, product_id=product.id, quantity=1, price=50000, total=50000,
                 grand_total=50000, status="cancelled", branch_id=branch.id, purchase_date=date(2026, 10, 6)),
    ])
    cash = TreasuryAccount(name="الصندوق", account_type="cash", is_default=True, is_active=True, branch_id=branch.id)
    bank = TreasuryAccount(name="مصرف الرافدين", account_type="bank", is_default=False, is_active=True,
                           branch_id=branch.id)
    db.session.add_all([cash, bank])
    db.session.commit()

    db.session.add(AccountTransaction(branch_id=branch.id, type="deposit", amount=500000, note="إيداع رأس مال",
                                      treasury_account_id=cash.id, created_at=datetime(2026, 10, 1, 9, 0)))
    expense = Expense.query.filter_by(title="إيجار المحل").first()
    db.session.add(AccountTransaction(branch_id=branch.id, type="withdraw", amount=50000,
                                      note=build_expense_withdraw_note(expense), treasury_account_id=cash.id,
                                      created_at=datetime(2026, 10, 4, 12, 0)))
    db.session.add(AccountTransaction(branch_id=branch.id, type="withdraw", amount=200000, note="دفعة شراء #1",
                                      treasury_account_id=cash.id, created_at=datetime(2026, 10, 3, 12, 0)))
    db.session.commit()
    append_payment_ledger_delta(cash_sale.id, 150000, cash.id, recorded_at=datetime(2026, 10, 2, 11, 5))
    append_payment_ledger_delta(partial.id, 40000, bank.id, recorded_at=datetime(2026, 10, 4, 11, 5))
    db.session.commit()
    transfer = record_treasury_transfer(cash.id, bank.id, 30000, "تحويل داخلي")
    transfer.created_at = datetime(2026, 10, 5, 10, 0)
    db.session.commit()
    return {
        "admin_id": admin.id, "seller_id": seller.id, "viewer_id": viewer.id, "outsider_id": outsider.id,
        "branch_id": branch.id, "other_branch_id": other_branch.id, "credit_sale": credit_sale.id,
    }


def _seed_other() -> dict:
    from flask import g
    from werkzeug.security import generate_password_hash

    from extensions import db
    from extensions_tenant import init_tenant_db
    from models.employee import Employee
    from models.invoice import Invoice
    from utils.branch_migration import get_default_branch

    g.tenant = OTHER_TENANT
    init_tenant_db(OTHER_TENANT)
    branch = get_default_branch()
    admin = Employee(name="Other", username="rp_other", password=generate_password_hash("x-pass-123"),
                     role="admin", is_active=True, branch_id=branch.id)
    from models.customer import Customer

    customer = Customer(branch_id=branch.id, name="شركة أخرى", phone="07800000000")
    db.session.add_all([admin, customer])
    db.session.flush()
    db.session.add(Invoice(customer_id=customer.id, customer_name="شركة أخرى", branch_id=branch.id, total=999000, paid_amount=999000,
                           status="تم التوصيل", payment_status="مسدد", stock_is_deducted=True,
                           created_at=datetime(2026, 10, 2, 10, 0)))
    db.session.commit()
    return {"admin_id": admin.id}


@pytest.fixture(scope="module")
def ctx():
    _wipe()
    previous = os.environ.get("FINORA_LOCAL_WORKSPACE")
    os.environ["FINORA_LOCAL_WORKSPACE"] = "1"
    from app import app

    app.config["TESTING"] = True
    from extensions import db

    with app.app_context():
        ids = _seed_main()
        db.session.remove()
    with app.app_context():
        other = _seed_other()
        db.session.remove()
    yield {"app": app, "ids": ids, "other": other}
    if previous is None:
        os.environ.pop("FINORA_LOCAL_WORKSPACE", None)
    else:
        os.environ["FINORA_LOCAL_WORKSPACE"] = previous
    _wipe()


def _client(ctx, key="admin_id", role="admin", tenant=TENANT, ids=None):
    client = ctx["app"].test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = (ids or ctx["ids"])[key]
        sess["tenant_slug"] = tenant
        sess["role"] = role
        sess["name"] = key
    return client


def _get(client, path, **params):
    res = client.get(path, query_string={**OCT, **params}, headers=JSON)
    return res


def _ok(res):
    assert res.status_code == 200, res.get_data(as_text=True)[:600]
    return res.get_json()


def _kpis(body):
    return {k["key"]: k for k in body["kpis"]}


def test_sales_profit_and_expenses_are_correct_and_consistent(ctx):
    client = _client(ctx)
    dash = _ok(_get(client, "/reports/api/v2/dashboard"))
    kpi = _kpis(dash)
    # booked: cash 150k + credit 300k + partial 140k (after a 10k discount)
    assert kpi["sales"]["value"] == 590000
    sales = _ok(_get(client, "/reports/api/v2/section/sales"))["metrics"]
    assert sales["net"] == 590000 and sales["gross"] == 600000 and sales["discounts"] == 10000
    assert sales["invoices"] == 3 and sales["average"] == 590000 // 3
    assert sales["returns"] == 150000 and sales["cogs"] == 400000 and sales["gross_profit"] == 190000

    profit = _ok(_get(client, "/reports/api/v2/section/profit"))
    st = profit["statement"]
    assert st["gross_revenue"] == 590000 + 10000 + 150000
    assert st["net_revenue"] == kpi["sales"]["value"]
    assert st["cogs"] == 400000 and st["operating_expenses"] == 50000
    assert st["net_profit"] == 590000 - 400000 - 50000 - st["other_expenses"]
    assert kpi["net_profit"]["value"] == st["net_profit"]
    assert dash["sections"]["profit"]["net_profit"] == st["net_profit"]

    expenses = _ok(_get(client, "/reports/api/v2/section/expenses"))
    assert expenses["metrics"]["total"] == 50000 == kpi["expenses"]["value"]  # the unposted 20k is excluded
    assert [r["label"] for r in expenses["breakdown"]["rows"]] == ["إيجار"]
    # rising expenses are bad news even when the number goes up
    assert kpi["expenses"]["tone"] in ("bad", "neutral")


def test_debts_inventory_and_cash_match_finora_balances(ctx):
    from flask import g

    from utils.accounting_calculations import calculate_accounts_receivable, calculate_inventory_value
    from utils.treasury_calculations import calculate_treasury_balance, list_treasury_accounts

    client = _client(ctx)
    dash = _ok(_get(client, "/reports/api/v2/dashboard"))
    kpi = _kpis(dash)
    customers = _ok(_get(client, "/reports/api/v2/section/customers", tab="debts"))
    # credit 300k + partial remainder 100k; the unbooked order is not a debt
    assert kpi["receivables"]["value"] == 400000 == customers["metrics"]["receivables"]
    assert sum(r["remaining"] for r in customers["rows"]) == 400000
    assert sum(a["amount"] for a in customers["aging"]) == 400000
    assert kpi["receivables"]["change_pct"] is None  # no historical snapshot: never invented

    inventory = _ok(_get(client, "/reports/api/v2/section/inventory"))
    assert kpi["inventory"]["value"] == inventory["metrics"]["value"]
    assert sum(r["value"] for r in inventory["rows"]) == inventory["metrics"]["value"]
    stagnant = _ok(_get(client, "/reports/api/v2/section/inventory", tab="stagnant"))
    assert any(r["name"] == "مكيف قديم" for r in stagnant["rows"])

    cash = _ok(_get(client, "/reports/api/v2/section/cash"))
    cashflow = _ok(_get(client, "/reports/api/v2/section/cashflow"))
    assert kpi["cash"]["value"] == cash["metrics"]["closing"] == cashflow["flow"]["closing"]
    assert cash["metrics"]["transfers"] == 30000
    flow = cashflow["flow"]
    inflows = {r["key"]: r["amount"] for r in flow["inflows"]}
    outflows = {r["key"]: r["amount"] for r in flow["outflows"]}
    # credit sales are not cash: only the 150k + 40k actually collected
    assert inflows["sales_collected"] == 190000
    assert outflows["expenses_paid"] == 50000 and outflows["supplier_payments"] == 200000
    # an internal transfer is neither income nor expense
    assert flow["net_transfers"] == 0
    assert flow["closing"] == flow["opening"] + flow["inflow"] - flow["outflow"] + flow["net_transfers"]

    with ctx["app"].test_request_context():
        g.tenant = TENANT
        assert calculate_accounts_receivable() == 400000
        assert calculate_inventory_value() == inventory["metrics"]["value"]
        if date.today() <= date(2026, 10, 31):
            assert cash["metrics"]["closing"] == sum(calculate_treasury_balance(a.id) for a in list_treasury_accounts())


def test_invoices_employees_and_purchases(ctx):
    client = _client(ctx)
    inv = _ok(_get(client, "/reports/api/v2/section/invoices"))["metrics"]
    assert (inv["paid"], inv["partial"], inv["credit"], inv["cancelled"], inv["returned"]) == (1, 1, 2, 1, 1)
    returned = _ok(_get(client, "/reports/api/v2/section/invoices", state="returned"))
    assert [r["state"] for r in returned["rows"]] == ["مرتجعة"]

    employees = _ok(_get(client, "/reports/api/v2/section/employees"))["rows"]
    by_name = {r["name"]: r for r in employees}
    assert by_name["سارة"]["sales"] == 440000 and by_name["سارة"]["orders"] == 2
    assert by_name["سارة"]["collections"] == 40000 and by_name["سارة"]["discounts"] == 10000

    only_seller = _ok(_get(client, "/reports/api/v2/dashboard", employee_id=ctx["ids"]["seller_id"]))
    assert _kpis(only_seller)["sales"]["value"] == 440000

    purchases = _ok(_get(client, "/reports/api/v2/section/purchases"))
    assert purchases["metrics"]["amount"] == 300000  # the cancelled purchase is excluded
    assert purchases["metrics"]["paid"] == 200000 and purchases["metrics"]["remaining"] == 100000
    debts = _ok(_get(client, "/reports/api/v2/section/purchases", tab="debts"))
    assert [r["remaining"] for r in debts["rows"]] == [100000]


def test_permissions_are_enforced_on_the_server(ctx):
    viewer = _client(ctx, "viewer_id", role="sales")
    dash = _ok(_get(viewer, "/reports/api/v2/dashboard"))
    keys = set(_kpis(dash))
    assert "sales" in keys and "net_profit" not in keys and "cash" not in keys and "expenses" not in keys
    assert "profit" not in dash["sections"] and "cashflow" not in dash["sections"]
    for section in ("profit", "cashflow", "cash", "expenses"):
        assert _get(viewer, f"/reports/api/v2/section/{section}").status_code == 403
        assert _get(viewer, f"/reports/api/v2/export/{section}.xlsx").status_code == 403
    employees = _ok(_get(viewer, "/reports/api/v2/section/employees"))
    assert all(r["profit"] is None for r in employees["rows"])

    outsider = _client(ctx, "outsider_id", role="sales")
    assert _get(outsider, "/reports/api/v2/dashboard").status_code == 403
    # from a non-loopback address (loopback is FINORA's local-workspace auto sign-in)
    anonymous = ctx["app"].test_client()
    res = anonymous.get("/reports/api/v2/dashboard", headers=JSON, environ_base={"REMOTE_ADDR": "203.0.113.9"})
    assert res.status_code in (302, 401)


def test_branch_and_tenant_isolation(ctx):
    viewer = _client(ctx, "viewer_id", role="sales")
    res = _get(viewer, "/reports/api/v2/dashboard", branch_id=ctx["ids"]["other_branch_id"])
    assert res.status_code == 403

    other = _client(ctx, "admin_id", tenant=OTHER_TENANT, ids=ctx["other"])
    kpi = _kpis(_ok(_get(other, "/reports/api/v2/dashboard")))
    assert kpi["sales"]["value"] == 999000  # only its own company
    mine = _kpis(_ok(_get(_client(ctx), "/reports/api/v2/dashboard")))
    assert mine["sales"]["value"] == 590000


def test_empty_period_and_bad_filters(ctx):
    client = _client(ctx)
    empty = _ok(client.get("/reports/api/v2/dashboard", query_string={"period": "custom", "date_from": "2020-01-01",
                                                                       "date_to": "2020-01-31"}, headers=JSON))
    kpi = _kpis(empty)
    assert kpi["sales"]["value"] == 0 and kpi["net_profit"]["value"] == 0 and kpi["sales"]["change_pct"] is None
    assert len(empty["sales_series"]) == 31 and all(p["amount"] == 0 for p in empty["sales_series"])
    bad = client.get("/reports/api/v2/dashboard", query_string={"period": "custom", "date_from": "2026-10-31",
                                                                 "date_to": "2026-10-01"}, headers=JSON)
    assert bad.status_code == 400
    assert client.get("/reports/api/v2/section/nope", headers=JSON).status_code == 404


def test_exports_follow_the_same_numbers(ctx):
    from openpyxl import load_workbook

    client = _client(ctx)
    res = _get(client, "/reports/api/v2/export/sales.xlsx")
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.data))
    data = wb["البيانات"]
    rows = list(data.iter_rows(values_only=True))
    assert rows[0][1] == "رقم الفاتورة"
    assert rows[-1][0] == "الإجمالي" and rows[-1][7] == 590000  # الصافي column total
    html = _get(client, "/reports/api/v2/print/profit").get_data(as_text=True)
    assert "الأرباح والخسائر" in html and "= صافي الربح" in html
