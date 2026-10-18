from io import BytesIO
from openpyxl import load_workbook
from flask import g
from extensions import db
from extensions_tenant import init_tenant_db
from tests.test_hr_workspace import ctx, _client, create_person, command, _run, _grant, TENANT, JSON


def test_hr_own_payslip_contains_frozen_details_and_blocks_other_employees(ctx):
    client=_client(ctx)
    other=create_person(client,ctx)
    from modules.hr.models import HRPerson
    with ctx["app"].app_context():
        g.tenant=TENANT
        init_tenant_db(TENANT)
        own_id=HRPerson.query.filter_by(employee_id=ctx["ids"]["clerk_id"]).one().id
    detail=client.get(f"/hr/api/people/{own_id}",headers=JSON).get_json()["person"]
    changed=command(client,"employee.update",{"person_id":own_id,"version":detail["version"],"joined_on":"2026-01-01","salary":650000,"effective_from":"2026-01-01"})
    assert changed.status_code==200,changed.get_json()
    _run(client,ctx,"2026-02")
    _grant(ctx,"hr.self")
    clerk=_client(ctx,"clerk_id","cashier")
    exported=clerk.get(f"/hr/api/reports/payroll/xlsx?month=2026-02&person_id={own_id}",headers=JSON)
    assert exported.status_code==200,exported.get_json()
    workbook=load_workbook(BytesIO(exported.data))
    assert len(workbook.worksheets)==3
    assert workbook.worksheets[1].max_row==30
    assert workbook.worksheets[1]["G3"].value==0
    assert clerk.get(f"/hr/api/reports/payroll/xlsx?month=2026-02&person_id={other}",headers=JSON).status_code in (403,404)


def test_hr_report_exports_all_filtered_rows_and_blocks_spreadsheet_formulas(ctx):
    client = _client(ctx)
    create_person(client, ctx, name="=SUM(1)")
    from modules.hr.models import HRPerson, HRContract
    from datetime import date
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        for index in range(105):
            profile=HRPerson(branch_id=ctx["ids"]["branch_id"], name=f"Export test {index}", code=f"RPT-{index}", joined_on=date(2026,1,1))
            db.session.add(profile);db.session.flush()
            db.session.add(HRContract(person_id=profile.id,effective_from=date(2026,1,1),terms='{"salary":0,"salary_type":"monthly"}',created_by=ctx["ids"]["admin_id"]))
        db.session.commit()
    exported = client.get("/hr/api/reports/employees/xlsx?search=Export test", headers=JSON)
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data))
    assert workbook.active.max_row == 107 and workbook.active.sheet_view.rightToLeft
    formula = client.get("/hr/api/reports/employees/xlsx?search=%3DSUM", headers=JSON)
    sheet = load_workbook(BytesIO(formula.data)).active
    assert sheet["B3"].value == "'=SUM(1)" and sheet["B3"].data_type == "s"
    printed = client.get("/hr/api/reports/employees/print?search=Export test", headers=JSON)
    assert printed.status_code == 200 and b"Export test 104" in printed.data
    assert "no-store" in exported.headers["Cache-Control"]


def test_hr_report_requires_report_and_underlying_financial_permission(ctx):
    client = _client(ctx)
    client.get("/hr/api/lookups", headers=JSON)
    _grant(ctx, "hr.view")
    clerk = _client(ctx, "clerk_id", "cashier")
    assert clerk.get("/hr/api/reports/employees/xlsx", headers=JSON).status_code == 403
    _grant(ctx, "hr.view", "hr.reports.view")
    assert clerk.get("/hr/api/reports/payroll/xlsx?month=2026-06", headers=JSON).status_code == 403
    assert clerk.get("/hr/api/reports/employees/xlsx", headers=JSON).status_code == 200
    _grant(ctx,"hr.view","hr.reports.view","hr.payroll.view")
    assert clerk.get(f"/hr/api/statement?person_id={create_person(client,ctx)}",headers=JSON).status_code==403


def test_native_adjustments_respect_scope_and_frozen_hr_payroll(ctx,monkeypatch):
    from routes import hr as legacy_hr
    monkeypatch.setattr(legacy_hr,"_business_type",lambda:"hr_management")
    client=_client(ctx)
    client.get("/hr/api/lookups",headers=JSON)
    from modules.hr.models import HRPerson
    from models.hr import HRPayrollAdjustment
    with ctx["app"].app_context():
        g.tenant=TENANT
        init_tenant_db(TENANT)
        profile=HRPerson.query.filter_by(employee_id=ctx["ids"]["clerk_id"]).one()
        own_id,version=profile.id,profile.version
        adjustment=HRPayrollAdjustment(employee_id=ctx["ids"]["clerk_id"],period_year=2026,period_month=4,amount=1000,title="legacy adjustment")
        db.session.add(adjustment);db.session.commit()
        adjustment_id=adjustment.id
    changed=command(client,"employee.update",{"person_id":own_id,"version":version,"joined_on":"2026-01-01","salary":650000,"effective_from":"2026-01-01"})
    assert changed.status_code==200,changed.get_json()
    _run(client,ctx,"2026-04")
    assert client.post(f"/hr/payroll/adjustments/{adjustment_id}/delete",headers=JSON).status_code==409
    assert client.post("/hr/payroll",data={"employee_id":ctx["ids"]["clerk_id"],"period_year":2026,"period_month":4,"amount":2000,"title":"changed"},headers=JSON).status_code==409
