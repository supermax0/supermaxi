import base64
from uuid import uuid4
from tests.test_hr_workspace import ctx, command, create_person, _client, _grant, TENANT, JSON


def test_hr_private_identity_bank_fields_and_manager_cycles_are_protected(ctx):
    client=_client(ctx)
    first,second=create_person(client,ctx),create_person(client,ctx)
    response=command(client,"employee.update",{"person_id":first,"version":1,"birth_date":"1990-02-03","national_id":"PRIVATE-ID","bank_name":"PRIVATE-BANK","bank_account":"123456"})
    assert response.status_code==200,response.get_json()
    info=client.get(f"/hr/api/people/{first}",headers=JSON).get_json()["private_info"]
    assert info["national_id"]=="PRIVATE-ID" and info["bank_account"]=="123456"
    assert command(client,"employee.update",{"person_id":first,"version":2,"manager_id":second}).status_code==200
    assert command(client,"employee.update",{"person_id":second,"version":1,"manager_id":first}).status_code==400
    _grant(ctx,"hr.view","hr.employees.edit","hr.resources.manage")
    clerk=_client(ctx,"clerk_id","cashier")
    assert clerk.get(f"/hr/api/people/{first}",headers=JSON).get_json()["private_info"]=={}
    assert command(clerk,"employee.update",{"person_id":first,"version":3,"national_id":"LEAK"}).status_code==403
    assert command(clerk,"employee.update",{"person_id":first,"version":3,"bank_account":"999"}).status_code==403


def test_hr_state_changes_are_versioned_do_not_disable_login_implicitly_and_keep_final_exit_closed(ctx):
    client=_client(ctx)
    profile=create_person(client,ctx)
    result=command(client,"employee.state",{"person_id":profile,"version":1,"status":"suspended","reason":"إيقاف إداري"})
    assert result.status_code==200,result.get_json()
    assert command(client,"employee.state",{"person_id":profile,"version":1,"status":"active","reason":"تفعيل"}).status_code==409
    assert command(client,"employee.state",{"person_id":profile,"version":2,"status":"terminated","reason":"غير مسموح"}).status_code==400
    assert command(client,"employee.state",{"person_id":profile,"version":2,"status":"active","reason":"انتهى الإيقاف"}).status_code==200


def test_hr_private_documents_validate_content_and_require_document_permission(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    payload = {"person_id": profile_id, "title": "عقد", "kind": "contract", "filename": "contract.pdf",
               "mime": "application/pdf", "content": base64.b64encode(b"%PDF-1.4\nHR test").decode()}
    document = command(client, "document.create", payload)
    assert document.status_code == 200, document.get_json()
    document_id = document.get_json()["id"]
    assert command(client, "document.create", {**payload, "content": base64.b64encode(b"<script>").decode()}).status_code == 400
    assert client.get(f"/hr/api/documents/{document_id}", headers=JSON).get_json()["content"] == payload["content"]
    _grant(ctx, "hr.view", "hr.resources.manage")
    clerk = _client(ctx, "clerk_id", role="cashier")
    assert clerk.get(f"/hr/api/documents/{document_id}", headers=JSON).status_code == 403
    assert clerk.get(f"/hr/api/people/{profile_id}", headers=JSON).get_json()["documents"] == []


def test_hr_uses_existing_assets_prevents_duplicate_custody_and_requires_return_before_exit(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    other_id = create_person(client, ctx)
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from models.fixed_asset import FixedAsset
    from models.fixed_asset_category import FixedAssetCategory
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        category = FixedAssetCategory(name="HR custody test")
        db.session.add(category)
        db.session.flush()
        asset = FixedAsset(asset_code=f"HRT-{uuid4().hex[:8]}", name="حاسبة", category_id=category.id,
                           branch_id=ctx["ids"]["branch_id"], status="active")
        db.session.add(asset)
        db.session.commit()
        asset_id = asset.id
    assigned = command(client, "asset.assign", {"person_id": profile_id, "asset_id": asset_id, "note": "سليم"})
    assert assigned.status_code == 200, assigned.get_json()
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        from utils.fixed_assets_service import _assert_hr_custody_returned, FixedAssetError
        import pytest
        with pytest.raises(FixedAssetError,match="بعهدة"):
            _assert_hr_custody_returned(db.session.get(FixedAsset,asset_id))
    assert command(client, "asset.assign", {"person_id": other_id, "asset_id": asset_id}).status_code == 400
    exit_payload = {"person_id": profile_id, "version": 1, "ends_on": "2026-08-15", "reason": "استقالة"}
    assert command(client, "employee.terminate", exit_payload).status_code == 400
    assert command(client, "asset.return", {"custody_id": assigned.get_json()["id"], "version": 1, "reason": "سليم"}).status_code == 200
    assert command(client, "employee.terminate", exit_payload).status_code == 200
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        from modules.hr.models import HRPerson
        from modules.hr.finance import calculate_person
        calculation = calculate_person(db.session.get(HRPerson, profile_id), "2026-08")
        assert calculation["gross"] == 314516 and calculation["ends_on"] == "2026-08-15"
    exported=client.get(f"/hr/api/reports/settlement/xlsx?person_id={profile_id}",headers=JSON)
    assert exported.status_code==200,exported.get_json()
    from io import BytesIO
    from openpyxl import load_workbook
    rows=dict(load_workbook(BytesIO(exported.data)).active.iter_rows(min_row=3,max_col=2,values_only=True))
    assert rows["آخر يوم عمل"]=="2026-08-15" and rows["الأجر المحتسب حتى آخر يوم عمل"]==314516
    assert "لا ينشئ دفعة" in rows["الصرف"]
    printed=client.get(f"/hr/api/reports/settlement/print?person_id={profile_id}",headers=JSON)
    assert printed.status_code==200 and "كشف نهاية الخدمة" in printed.get_data(as_text=True)
    _grant(ctx,"hr.view","hr.reports.view")
    assert _client(ctx,"clerk_id","cashier").get(f"/hr/api/reports/settlement/xlsx?person_id={profile_id}",headers=JSON).status_code==403


def test_hr_monthly_allowances_and_wage_cuts_enter_calculation_once(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    for kind, amount in (("allowance", 50000), ("deduction", 20000)):
        response = command(client, "payroll.adjustment", {"person_id": profile_id, "month": "2026-06", "kind": kind,
                            "amount": amount, "title": "تسوية معتمدة"})
        assert response.status_code == 200, response.get_json()
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from modules.hr.models import HRPerson
    from modules.hr.finance import calculate_person
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        calculation = calculate_person(db.session.get(HRPerson, profile_id), "2026-06")
        assert calculation["gross"] == 700000 and calculation["net"] == 680000
        assert calculation["wage_deductions"] == 20000 and len(calculation["adjustments"]) == 2
