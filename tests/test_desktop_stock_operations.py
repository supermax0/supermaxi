"""Physical balances, FIFO, non-cash loss, atomic audit, replay and source isolation."""
import json
from datetime import datetime,timedelta
from uuid import uuid4
import pytest
from tests.test_finance_operations import finance,balance


@pytest.fixture
def stock(finance):
    f=finance
    from extensions import db
    from models.branch import Branch,BranchStock
    from models.product import Product
    from models.inventory_lot import InventoryLot
    f['get']('/api/finance/meta')
    destination=Branch(code='TARGET',name='الموقع الثاني',is_active=True)
    db.session.add(destination);db.session.flush()
    product=Product(name='صنف نقل',sku='STOCK-OPS',branch_id=f['branch'].id,quantity=6,
        buy_price=250,sale_price=400,active=True,low_stock_threshold=2,catalog_group_id=uuid4().hex)
    db.session.add(product);db.session.flush()
    db.session.add(BranchStock(branch_id=f['branch'].id,product_id=product.id,quantity=6,opening_stock=6))
    lots=[InventoryLot(product_id=product.id,branch_id=f['branch'].id,quantity=2,remaining_quantity=2,
        unit_cost=100,created_at=datetime.utcnow()-timedelta(days=2)),
        InventoryLot(product_id=product.id,branch_id=f['branch'].id,quantity=4,remaining_quantity=4,
        unit_cost=250,created_at=datetime.utcnow()-timedelta(days=1))]
    db.session.add_all(lots);db.session.commit()
    def call(path,body=None,*,employee=None,view_all=True,tenant=None,shared=False):
        from flask import g,session
        user=employee or f['admin']
        with f['app'].test_request_context(path,method='POST' if body is not None else 'GET',json=body):
            session.update(user_id=user.id,employee_id=user.id,role=user.role,tenant_slug=tenant or f['tenant'])
            g.tenant=f['tenant'];g.branch=f['branch'];g.branches=[f['branch'],destination]
            g.view_all_branches=view_all;g.independent_branches_enabled=not shared;g.branch_scope_ready=True
            endpoint,args=f['app'].url_map.bind('localhost').match(path.split('?')[0],method='POST' if body is not None else 'GET')
            return f['app'].make_response(f['app'].view_functions[endpoint](**args))
    def get(path,**kwargs):
        response=call(path,**kwargs);assert response.status_code==200,response.get_json()
        return response.get_json()['data']
    f.update(destination=destination,product=product,lots=lots,request=call,read=get)
    return f


def draft(f,quantity=3,variant_name=''):
    body={'operation_id':str(uuid4()),'from_location':f"branch:{f['branch'].id}",
        'to_location':f"branch:{f['destination'].id}",'reason':'إعادة توزيع مخزون',
        'lines':[{'product_id':f['product'].id,'variant_name':variant_name,'quantity':quantity}]}
    response=f['request']('/api/v1/desktop/stock-operations/branch/transfers',body)
    assert response.status_code==200,response.get_json()
    return response.get_json()['data']['id'],body


def action(f,document_id,phase,key=None):
    detail=f['read'](f'/api/v1/desktop/stock-operations/branch/transfers/{document_id}')
    body={'operation_id':key or str(uuid4()),'revision':detail['revision'],'reason':'تحقق وتسليم فعلي'}
    return f['request'](f'/api/v1/desktop/stock-operations/branch/transfers/{document_id}/{phase}',body),body


def source(f):return f['read'](f"/api/v1/desktop/stock-operations/stocks?location=branch:{f['branch'].id}")['items'][0]


def test_transfer_send_receive_preserves_fifo_value_cash_and_exact_replay(stock):
    from models.branch import BranchStock,StockTransfer
    from models.inventory_lot import InventoryLot
    from models.stock_operations import StockTransferCostLayer
    from models.journal_entry import JournalEntry
    from utils.inventory_lots import current_lot_inventory_value
    f=stock;cash=balance(f);journals=JournalEntry.query.count()
    assert current_lot_inventory_value()==1200
    from utils.accounting_calculations import calculate_inventory_value
    assert calculate_inventory_value()==1200
    document,create=draft(f)
    assert source(f)['quantity']==6
    assert f['request']('/api/v1/desktop/stock-operations/branch/transfers',create).get_json()['data']['id']==document
    sent,payload=action(f,document,'send');assert sent.status_code==200,sent.get_json()
    assert source(f)['quantity']==3
    assert current_lot_inventory_value()==1200
    assert calculate_inventory_value()==1200
    assert [lot.remaining_quantity for lot in f['lots']]==[0,3]
    replay=f['request'](f'/api/v1/desktop/stock-operations/branch/transfers/{document}/send',payload)
    assert replay.status_code==200 and source(f)['quantity']==3
    received,receive=action(f,document,'receive');assert received.status_code==200,received.get_json()
    assert f['request'](f'/api/v1/desktop/stock-operations/branch/transfers/{document}/receive',receive).status_code==200
    row=StockTransfer.query.get(document);destination=row.lines[0].destination_product_id
    assert BranchStock.query.filter_by(branch_id=f['destination'].id,product_id=destination).one().quantity==3
    assert sorted((lot.remaining_quantity,lot.unit_cost) for lot in InventoryLot.query.filter_by(branch_id=f['destination'].id))==[(1,250),(2,100)]
    assert StockTransferCostLayer.query.filter_by(transfer_id=document,phase='receive').count()==2
    assert current_lot_inventory_value()==1200
    assert calculate_inventory_value()==1200
    assert balance(f)==cash and JournalEntry.query.count()==journals


def test_sent_cancellation_restores_original_layers_once_and_received_cannot_cancel(stock):
    f=stock;document,_=draft(f)
    assert action(f,document,'send')[0].status_code==200
    response,payload=action(f,document,'cancel');assert response.status_code==200,response.get_json()
    assert source(f)['quantity']==6 and [lot.remaining_quantity for lot in f['lots']]==[2,4]
    assert f['request'](f'/api/v1/desktop/stock-operations/branch/transfers/{document}/cancel',payload).status_code==200
    other,_=draft(f);assert action(f,other,'send')[0].status_code==200
    assert action(f,other,'receive')[0].status_code==200
    assert action(f,other,'cancel')[0].status_code==409


def test_unreconciled_or_missing_legacy_fifo_is_blocked_without_stock_change(stock):
    f=stock;document,_=draft(f)
    f['lots'][0].remaining_quantity=1;f['db'].session.commit()
    response,_=action(f,document,'send')
    assert response.status_code==409 and source(f)['quantity']==6
    assert f['read'](f'/api/v1/desktop/stock-operations/branch/transfers/{document}')['status']=='draft'


def test_audit_failure_rolls_back_receipt_layers_and_stock_and_same_operation_retries(stock):
    from models.activity_log import ActivityLog
    from models.stock_operations import StockTransferCostLayer
    from sqlalchemy import event
    f=stock;document,_=draft(f)
    detail=f['read'](f'/api/v1/desktop/stock-operations/branch/transfers/{document}')
    body={'operation_id':str(uuid4()),'revision':detail['revision'],'reason':'تسليم موثق'}
    def fail(*args):raise RuntimeError('audit unavailable')
    event.listen(ActivityLog,'before_insert',fail)
    try:
        with pytest.raises(RuntimeError,match='audit unavailable'):
            f['request'](f'/api/v1/desktop/stock-operations/branch/transfers/{document}/send',body)
    finally:event.remove(ActivityLog,'before_insert',fail)
    assert source(f)['quantity']==6 and StockTransferCostLayer.query.count()==0
    assert f['request'](f'/api/v1/desktop/stock-operations/branch/transfers/{document}/send',body).status_code==200


def test_writeoff_and_reversal_post_exact_non_cash_cost_and_seal_history(stock):
    from models.stock_operations import InventoryWriteOff
    from models.journal_entry import JournalEntry
    from models.stock_movement import StockMovement
    from sqlalchemy import text
    f=stock;before=source(f);cash=balance(f)
    body={'operation_id':str(uuid4()),'location_key':f"branch:{f['branch'].id}",'product_id':f['product'].id,
        'variant_name':'','quantity':3,'stock_revision':before['revision'],'reason':'تلف فعلي موثق'}
    reply=f['request']('/api/v1/desktop/stock-operations/branch/writeoffs',body)
    assert reply.status_code==200,reply.get_json()
    document=reply.get_json()['data']['id']
    row=InventoryWriteOff.query.get(document)
    assert row.cost==450 and row.sealed and source(f)['quantity']==3
    assert JournalEntry.query.get(row.journal_entry_id).amount==450
    assert StockMovement.query.get(row.movement_id).total_cost==450
    from utils.inventory_writeoff_totals import inventory_writeoff_loss
    from utils.accounting_calculations import calculate_net_profit
    assert inventory_writeoff_loss()==450 and calculate_net_profit()==-450
    from modules.finance.common import today
    from services.reports_service import profit_and_loss,ReportFilters
    report=profit_and_loss(ReportFilters('today',today(),today()),today(),today())
    assert report['inventory_losses']==450 and report['net_profit']==-450
    from services.reports_export import profit_lines
    assert next(line for line in profit_lines({'statement':report}) if 'خسائر المخزون' in line['label'])['amount']==-450
    from services.reports_export import to_xlsx,to_html
    from openpyxl import load_workbook
    from io import BytesIO
    exported={'statement':report,'margins':{'gross_margin_pct':None,'net_margin_pct':None}}
    columns=[('label','البند',False),('amount','المبلغ',True)];lines=profit_lines(exported)
    workbook=load_workbook(BytesIO(to_xlsx('profit',exported,lines,columns,[],'Test')),read_only=True)
    sheet=workbook[workbook.sheetnames[-1]]
    assert list(sheet.values)[-1][1]==-450
    captured={}
    def render(name,**data):captured.update(data);return 'html'
    to_html('profit',exported,lines,columns,[],'Test',None,render)
    assert captured['totals']==[]
    trial=f['read']('/api/finance/reports/trial-balance',view_all=False)
    assert trial['balanced'] and trial['expenses']==450 and trial['profit']==-450
    from models.employee import Employee
    from flask import g,session
    staff=Employee(name='Other branch',username=uuid4().hex,password='test',role='cashier',is_active=True,branch_id=f['destination'].id)
    f['db'].session.add(staff);f['db'].session.commit()
    with f['app'].test_request_context('/reports/api/v2/section/profit'):
        session.update(user_id=staff.id,tenant_slug=f['tenant'])
        g.branch=f['destination'];g.branches=[f['branch'],f['destination']];g.view_all_branches=False;g.independent_branches_enabled=False
        assert inventory_writeoff_loss()==0
        session['user_id']=f['admin'].id
        assert inventory_writeoff_loss()==0
        from modules.finance.reports import trial_balance
        assert trial_balance({})['expenses']==0
    assert f['request']('/api/v1/desktop/stock-operations/branch/writeoffs',body).status_code==200
    detail=f['read'](f'/api/v1/desktop/stock-operations/writeoffs/{document}')
    reverse={'operation_id':str(uuid4()),'location_key':detail['location_key'],'revision':detail['revision'],'reason':'تصحيح تلف سجل بالخطأ'}
    response=f['request'](f'/api/v1/desktop/stock-operations/writeoffs/{document}/reverse',reverse)
    assert response.status_code==200,response.get_json()
    assert source(f)['quantity']==6 and [lot.remaining_quantity for lot in f['lots']]==[2,4]
    assert balance(f)==cash
    assert inventory_writeoff_loss()==0 and calculate_net_profit()==0
    assert profit_and_loss(ReportFilters('today',today(),today()),today(),today())['net_profit']==0
    assert f['request'](f'/api/v1/desktop/stock-operations/writeoffs/{document}/reverse',reverse).status_code==200
    with pytest.raises(Exception,match='immutable'):
        f['db'].session.execute(text('DELETE FROM inventory_writeoff WHERE id=:id'),{'id':document})
    f['db'].session.rollback()
    row=InventoryWriteOff.query.get(document);row.reason='تغيير مباشر'
    with pytest.raises(Exception,match='محفوظ'):f['db'].session.flush()
    f['db'].session.rollback()


def test_reorder_threshold_is_scoped_versioned_and_replayed_without_mutating_stock(stock):
    f=stock;body={'operation_id':str(uuid4()),'location_key':f"branch:{f['branch'].id}",
        'product_id':f['product'].id,'variant_name':'','minimum':7,'target':12,'rule_version':0,'reason':'حاجة تشغيلية'}
    reply=f['request']('/api/v1/desktop/stock-operations/branch/rules',body)
    assert reply.status_code==200,reply.get_json()
    row=source(f);assert (row['minimum'],row['target'],row['suggested'],row['quantity'])==(7,12,6,6)
    assert f['request']('/api/v1/desktop/stock-operations/branch/rules',body).status_code==200
    changed={**body,'operation_id':str(uuid4()),'minimum':9}
    assert f['request']('/api/v1/desktop/stock-operations/branch/rules',changed).status_code==409
    assert f['read'](f"/api/v1/desktop/stock-operations/stocks?location=branch:{f['branch'].id}&reorder=1")['total']==1


def test_scopes_permissions_invalid_quantity_and_changed_key_are_rejected(stock):
    from models.employee import Employee
    from models.role import Role
    f=stock;staff=Employee(name='Denied',username=uuid4().hex,password='test',role='cashier',is_active=True,branch_id=f['branch'].id)
    staff.roles=[Role(name='no-stock-'+uuid4().hex)];f['db'].session.add(staff);f['db'].session.commit()
    assert f['request']('/api/v1/desktop/stock-operations/meta',employee=staff).status_code==403
    assert f['request']('/api/v1/desktop/stock-operations/meta',tenant='other-company').status_code==403
    assert f['request'](f"/api/v1/desktop/stock-operations/stocks?location=branch:{f['destination'].id}",view_all=False).status_code==403
    document,body=draft(f)
    assert f['request']('/api/v1/desktop/stock-operations/branch/transfers',{**body,'reason':'تغيير نفس المعرف'}).status_code==409
    bad={**body,'operation_id':str(uuid4()),'lines':[{'product_id':f['product'].id,'quantity':True,'variant_name':''}]}
    assert f['request']('/api/v1/desktop/stock-operations/branch/transfers',bad).status_code in (400,409)
    assert f['request']('/api/v1/desktop/stock-operations/meta?kind=warehouse').status_code==403


def test_reads_do_not_create_operation_tables_or_change_lots(stock):
    from models.stock_operations import InventoryWriteOff,StockReorderRule,StockTransferCostLayer
    from sqlalchemy import inspect,text
    f=stock
    for table in [InventoryWriteOff.__table__,StockReorderRule.__table__,StockTransferCostLayer.__table__]:
        table.drop(bind=f['db'].session.connection(),checkfirst=True)
    f['db'].session.commit()
    for path in ['/api/v1/desktop/stock-operations/meta',f"/api/v1/desktop/stock-operations/stocks?location=branch:{f['branch'].id}",
        '/api/v1/desktop/stock-operations/transfers',f"/api/v1/desktop/stock-operations/writeoffs?location=branch:{f['branch'].id}"]:
        assert f['request'](path).status_code==200
    assert not inspect(f['db'].session.connection()).has_table('stock_transfer_cost_layer')
    assert [lot.remaining_quantity for lot in f['lots']]==[2,4]


def test_option_transfer_preserves_selected_color_size_and_fifo(stock):
    from models.product_color_variant import ProductColorVariant
    from models.branch import StockTransfer
    from models.inventory_lot import InventoryLot
    f=stock;option='أسود · L'
    f['product'].meta_json=json.dumps({'has_colors':True})
    variant=ProductColorVariant(product_id=f['product'].id,color_name=option,quantity=6,buy_price=250,sale_price=400)
    f['db'].session.add(variant)
    for lot in f['lots']:lot.variant_color=option
    f['db'].session.commit()
    rows=f['read'](f"/api/v1/desktop/stock-operations/stocks?location=branch:{f['branch'].id}")['items']
    assert rows[0]['variant_id']==variant.id and rows[0]['variant_name']==option
    document,_=draft(f,quantity=3,variant_name=option)
    sent,_=action(f,document,'send');assert sent.status_code==200,sent.get_json()
    assert variant.quantity==3
    received,_=action(f,document,'receive');assert received.status_code==200,received.get_json()
    target=StockTransfer.query.get(document).lines[0].destination_product_id
    assert ProductColorVariant.query.filter_by(product_id=target,color_name=option).one().quantity==3
    assert {(lot.remaining_quantity,lot.unit_cost,lot.variant_color) for lot in InventoryLot.query.filter_by(branch_id=f['destination'].id)}=={(2,100,option),(1,250,option)}


def test_writeoff_preview_closed_period_stale_revision_and_print_escape(stock):
    from utils.financial_period_guard import close_financial_period
    from modules.finance.common import today
    f=stock;row=source(f);place=f"branch:{f['branch'].id}"
    preview=f['read'](f"/api/v1/desktop/stock-operations/valuation?location={place}&product_id={f['product'].id}&quantity=3")
    assert preview['cost']==450 and preview['eligible'] and preview['stock_revision']==row['revision']
    body={'operation_id':str(uuid4()),'location_key':place,'product_id':f['product'].id,'variant_name':'','quantity':3,'stock_revision':'wrong','reason':'<script>alert(1)</script>'}
    assert f['request']('/api/v1/desktop/stock-operations/branch/writeoffs',body).status_code==409
    date=today();close_financial_period(date.year,date.month);f['db'].session.commit()
    body['stock_revision']=row['revision']
    assert f['request']('/api/v1/desktop/stock-operations/branch/writeoffs',body).status_code==409
    assert source(f)['quantity']==6
    transfer,_=draft(f)
    assert action(f,transfer,'send')[0].status_code==409
    assert source(f)['quantity']==6
    from models.financial_period_close import FinancialPeriodClose
    FinancialPeriodClose.query.delete();f['db'].session.commit()
    response=f['request']('/api/v1/desktop/stock-operations/branch/writeoffs',body)
    assert response.status_code==200,response.get_json()
    document=response.get_json()['data']['id']
    ticket=f['request'](f'/api/v1/desktop/stock-operations/branch/writeoffs/{document}/print')
    assert ticket.status_code==200
    assert '<script>alert(1)</script>' not in ticket.get_data(as_text=True)
    assert '&lt;script&gt;' in ticket.get_data(as_text=True)
    activity=f['read']('/api/v1/desktop/inventory-activity')
    assert any(item['source_type']=='inventory_writeoff' and item['document_id']==document and item['can_open_source'] for item in activity['items'])


def test_warehouse_optional_creation_transfer_cost_proof_and_movement_history(stock):
    from models.tenant import Tenant
    from modules.wholesale.models.warehouse import WarehouseStock
    f=stock
    tenant=Tenant.query.first()
    if tenant:tenant.business_type='wholesale'
    else:f['db'].session.add(Tenant(name='Wholesale fixture',business_type='wholesale'))
    f['db'].session.commit()
    ids=[]
    for code,branch in [('W1',f['branch']),('W2',f['destination'])]:
        response=f['request']('/api/v1/desktop/stock-operations/warehouse/locations',{'operation_id':str(uuid4()),'name':code,'code':code,'branch_id':branch.id,'warehouse_kind':'main','reason':'مخزن فعلي جديد'})
        assert response.status_code==200,response.get_json();ids.append(response.get_json()['data']['id'])
    f['db'].session.add(WarehouseStock(warehouse_id=ids[0],product_id=f['product'].id,quantity=6,opening_stock=6));f['db'].session.commit()
    body={'operation_id':str(uuid4()),'from_location':f'warehouse:{ids[0]}','to_location':f'warehouse:{ids[1]}','reason':'توزيع مخازن','lines':[{'product_id':f['product'].id,'quantity':3,'variant_name':''}]}
    response=f['request']('/api/v1/desktop/stock-operations/warehouse/transfers',body)
    assert response.status_code==200,response.get_json();document=response.get_json()['data']['id']
    def phase(name):
        row=f['read'](f'/api/v1/desktop/stock-operations/warehouse/transfers/{document}')
        return f['request'](f'/api/v1/desktop/stock-operations/warehouse/transfers/{document}/{name}',{'operation_id':str(uuid4()),'revision':row['revision'],'reason':'تم التحقق الفعلي'})
    response=phase('send');assert response.status_code==200,response.get_json()
    assert WarehouseStock.query.filter_by(warehouse_id=ids[1]).one().in_transit_quantity==3
    f['product'].buy_price=999;f['db'].session.commit()
    response=phase('receive');assert response.status_code==200,response.get_json()
    detail=f['read'](f'/api/v1/desktop/stock-operations/warehouse/transfers/{document}')
    assert detail['lines'][0]['cost']==750
    target=WarehouseStock.query.filter_by(warehouse_id=ids[1]).one();assert target.quantity==3 and target.in_transit_quantity==0
    history=f['read'](f'/api/v1/desktop/stock-operations/warehouse/movements?location=warehouse:{ids[1]}')
    assert history['items'][0]['document_id']==document and history['items'][0]['cost']==750
    assert not f['read']('/api/v1/desktop/stock-operations/meta?kind=warehouse')['can_writeoff']
    response=f['request']('/api/v1/desktop/stock-operations/warehouse/transfers',{**body,'operation_id':str(uuid4()),'lines':[{'product_id':f['product'].id,'quantity':1,'variant_name':''}]})
    assert response.status_code==200,response.get_json();document=response.get_json()['data']['id']
    assert phase('send').status_code==200
    cancelled=phase('cancel');assert cancelled.status_code==200,cancelled.get_json()
    assert WarehouseStock.query.filter_by(warehouse_id=ids[0]).one().quantity==3
    assert target.quantity==3 and target.in_transit_quantity==0


def test_shared_catalog_financial_value_keeps_fifo_and_in_transit_once(stock):
    from utils.accounting_calculations import calculate_inventory_value
    from flask import g
    f=stock;g.independent_branches_enabled=False
    assert calculate_inventory_value()==1200
    document,_=draft(f)
    assert action(f,document,'send')[0].status_code==200
    g.independent_branches_enabled=False
    assert calculate_inventory_value()==1200
    assert action(f,document,'receive')[0].status_code==200
    g.independent_branches_enabled=False
    assert calculate_inventory_value()==1200


def test_offline_device_allocation_is_not_available_for_transfer_or_writeoff(stock):
    from utils.desktop_stock_reservations import DesktopStockAllocation,ensure_reservation_schema
    from models.branch import BranchStock
    f=stock;ensure_reservation_schema();row=BranchStock.query.filter_by(product_id=f['product'].id,branch_id=f['branch'].id).one()
    f['db'].session.add(DesktopStockAllocation(id=str(uuid4()),grant_id=str(uuid4()),product_id=f['product'].id,
        requested_product_id=f['product'].id,branch_id=f['branch'].id,stock_id=row.id,quantity=5));f['db'].session.commit()
    available=source(f);assert available['available']==1 and available['reserved']==5
    body={'operation_id':str(uuid4()),'location_key':f"branch:{f['branch'].id}",'product_id':f['product'].id,'variant_name':'','quantity':2,'stock_revision':available['revision'],'reason':'تلف غير متاح'}
    assert f['request']('/api/v1/desktop/stock-operations/branch/writeoffs',body).status_code==409
    assert source(f)['quantity']==6


def test_oversized_accounting_loss_is_rejected_without_rounding_or_stock_mutation(stock):
    f=stock
    for lot in f['lots']:lot.unit_cost=5_000_000_000_000_000
    f['db'].session.commit();row=source(f);place=f"branch:{f['branch'].id}"
    preview=f['read'](f"/api/v1/desktop/stock-operations/valuation?location={place}&product_id={f['product'].id}&quantity=2")
    assert not preview['eligible'] and preview['cost'] is None
    body={'operation_id':str(uuid4()),'location_key':place,'product_id':f['product'].id,'quantity':2,'variant_name':'','stock_revision':row['revision'],'reason':'اختبار مجال الكلفة'}
    assert f['request']('/api/v1/desktop/stock-operations/branch/writeoffs',body).status_code==400
    assert source(f)['quantity']==6 and [r.remaining_quantity for r in f['lots']]==[2,4]

