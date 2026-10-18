"""HR operates against isolated tenants and the real FINORA staff session."""
from uuid import uuid4
from datetime import datetime, timedelta, timezone
import pytest
from tests.test_desktop_data_api import ctx, _client, TENANT, JSON


def command(client, action, payload, command_id=None):
    return client.post("/hr/api/commands", json={"id": command_id or str(uuid4()), "action": action, "payload": payload}, headers=JSON)


def create_person(client, ctx, **extra):
    payload = {"name": f"HR Person {uuid4().hex[:8]}", "joined_on": "2026-01-01", "branch_id": ctx["ids"]["branch_id"],
               "salary": 650000, "salary_type": "monthly", **extra}
    response = command(client, "employee.create", payload)
    assert response.status_code == 200, response.get_json()
    return response.get_json()["id"]


def test_hr_creates_profile_without_login_and_replays_exact_command(ctx):
    client = _client(ctx)
    request_id = str(uuid4())
    payload = {"name": "HR Independent Profile", "joined_on": "2026-01-01", "branch_id": ctx["ids"]["branch_id"], "salary": 650000}
    first = command(client, "employee.create", payload, request_id)
    assert first.status_code == 200, first.get_json()
    assert command(client, "employee.create", payload, request_id).get_json() == first.get_json()
    assert command(client, "employee.create", {**payload, "name": "Different"}, request_id).status_code == 409
    details = client.get(f"/hr/api/people/{first.get_json()['id']}", headers=JSON).get_json()
    assert details["person"]["employee_id"] is None
    assert details["person"]["salary"] == 650000
    assert len(details["contracts"]) == 1
    page = client.get("/hr/api/page?search=HR Independent Profile", headers=JSON).get_json()
    assert page["total"] == 1
    assert "password" not in str(details)


def test_hr_update_requires_current_version_and_effective_financial_date(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    payload = {"person_id": profile_id, "version": 1, "job_title": "محاسب"}
    assert command(client, "employee.update", payload).status_code == 200
    assert command(client, "employee.update", payload).status_code == 409
    assert command(client, "employee.update", {"person_id": profile_id, "version": 2, "salary": 700000}).status_code == 400
    changed = command(client, "employee.update", {"person_id": profile_id, "version": 2, "salary": 700000, "effective_from": "2026-10-01"})
    assert changed.status_code == 200, changed.get_json()
    row = client.get(f"/hr/api/people/{profile_id}", headers=JSON).get_json()
    assert row["person"]["version"] == 3 and len(row["contracts"]) == 2


def _grant(ctx, *permissions):
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from models.employee import Employee
    from models.role import Permission, Role
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        role = Role(name=f"hr-test-{uuid4().hex[:12]}")
        role.permissions = Permission.query.filter(Permission.name.in_(permissions)).all()
        employee = db.session.get(Employee, ctx["ids"]["clerk_id"])
        employee.roles = [role]
        db.session.add(role)
        db.session.commit()


def test_hr_directory_permission_does_not_expose_salary_or_allow_payments(ctx):
    admin = _client(ctx)
    profile_id = create_person(admin, ctx)
    _grant(ctx, "hr.view")
    clerk = _client(ctx, "clerk_id", role="cashier")
    row = clerk.get(f"/hr/api/people/{profile_id}", headers=JSON).get_json()
    assert "salary" not in row["person"] and row["contracts"] == []
    assert row["attendance"] is None
    assert clerk.get("/hr/api/overview", headers=JSON).get_json()["salaries"] is None
    assert clerk.get("/hr/api/workspace?tab=payroll", headers=JSON).status_code == 403
    assert clerk.get(f"/hr/api/statement?person_id={profile_id}",headers=JSON).status_code==403
    assert command(clerk, "payroll.pay", {"item_id": 1, "amount": 1, "treasury_account_id": 1}).status_code == 403


def test_hr_attendance_and_overnight_shift_use_one_event_and_unpaid_break(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    shift = command(client, "shift.create", {"branch_id": ctx["ids"]["branch_id"], "name": "وردية ليلية",
                    "starts": "22:00", "ends": "06:00", "weekdays": list(range(7)), "break_minutes": 30,
                    "grace_minutes": 10, "overtime_enabled": True})
    assert shift.status_code == 200, shift.get_json()
    now = datetime.now(timezone.utc)
    local_date = (now + timedelta(hours=3)).date()
    # The attendance event test uses the real yesterday-night shift; assigning
    # yesterday is deliberately forbidden, so seed its approved historic assignment.
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from modules.hr.models import HRShiftAssignment
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        db.session.add(HRShiftAssignment(person_id=profile_id, shift_id=shift.get_json()["id"], effective_from=local_date - timedelta(days=2)))
        db.session.commit()
    yesterday = local_date - timedelta(days=1)
    clock_in = datetime.fromisoformat(f"{yesterday}T22:20:00+03:00")
    clock_out = datetime.fromisoformat(f"{local_date}T06:15:00+03:00")
    if clock_out > now:
        clock_in -= timedelta(days=1)
        clock_out -= timedelta(days=1)
    payload = {"person_id": profile_id, "event": "check_in", "occurred_at": clock_in.isoformat()}
    event_id = str(uuid4())
    first = command(client, "attendance.record", payload, event_id)
    assert first.status_code == 200, first.get_json()
    assert command(client, "attendance.record", payload, event_id).get_json() == first.get_json()
    assert command(client, "attendance.record", payload).status_code == 400
    out = command(client, "attendance.record", {"person_id": profile_id, "event": "check_out", "occurred_at": clock_out.isoformat()})
    assert out.status_code == 200, out.get_json()
    row = out.get_json()["attendance"]
    assert row["worked_minutes"] == 445 and row["late_minutes"] == 10 and row["overtime_minutes"] == 15


def test_hr_scoped_branch_and_self_access(ctx):
    admin = _client(ctx)
    profile_id = create_person(admin, ctx)
    _grant(ctx, "hr.self")
    clerk = _client(ctx, "clerk_id", role="cashier")
    assert clerk.get(f"/hr/api/people/{profile_id}", headers=JSON).status_code == 404
    page = clerk.get("/hr/api/page", headers=JSON).get_json()
    assert page["total"] == 1 and page["rows"][0]["employee_id"] == ctx["ids"]["clerk_id"]
    assert command(clerk, "employee.create", {"name": "Forbidden"}).status_code == 403


def _fund(ctx, amount=1000000, account_type="cash"):
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from models.treasury_account import TreasuryAccount
    from models.account_transaction import AccountTransaction
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        account = TreasuryAccount(name=f"HR Fund {uuid4().hex[:8]}", account_type=account_type,
                                  branch_id=ctx["ids"]["branch_id"], is_active=True)
        db.session.add(account)
        db.session.flush()
        db.session.add(AccountTransaction(branch_id=account.branch_id, type="deposit", amount=amount, treasury_account_id=account.id, note="HR test opening"))
        db.session.commit()
        return account.id


def _run(client, ctx, month):
    prepared = command(client, "payroll.prepare", {"month": month, "branch_id": ctx["ids"]["branch_id"]})
    assert prepared.status_code == 200, prepared.get_json()
    run_id = prepared.get_json()["id"]
    reviewed = command(client, "payroll.decide", {"run_id": run_id, "version": 1, "status": "reviewed"})
    assert reviewed.status_code == 200, reviewed.get_json()
    approved = command(client, "payroll.decide", {"run_id": run_id, "version": 2, "status": "approved"})
    assert approved.status_code == 200, approved.get_json()
    return run_id


def test_hr_payroll_accrual_does_not_move_cash_and_partial_payment_is_exactly_once(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    account_id = _fund(ctx)
    run_id = _run(client, ctx, "2026-02")
    rows = client.get(f"/hr/api/workspace?tab=payroll&month=2026-02&person_id={profile_id}", headers=JSON).get_json()["rows"]
    assert len(rows) == 1 and rows[0]["net"] == 650000
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from models.account_transaction import AccountTransaction
    from models.journal_entry import JournalEntry
    from modules.hr.models import HRPayrollItem
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        assert AccountTransaction.query.filter_by(treasury_account_id=account_id, type="withdraw").count() == 0
        assert JournalEntry.query.filter_by(reference_type="hr_payroll_accrual", reference_id=rows[0]["id"]).one().amount == 650000
    payload = {"item_id": rows[0]["id"], "treasury_account_id": account_id, "amount": 200000}
    command_id = str(uuid4())
    paid = command(client, "payroll.pay", payload, command_id)
    assert paid.status_code == 200 and paid.get_json()["remaining"] == 450000, paid.get_json()
    assert command(client, "payroll.pay", payload, command_id).get_json() == paid.get_json()
    statement=client.get(f"/hr/api/statement?person_id={profile_id}",headers=JSON)
    assert statement.status_code==200,statement.get_json()
    balances=statement.get_json()
    assert (balances["due"],balances["paid"],balances["balance"],balances["advance_remaining"])==(650000,200000,450000,0)
    assert len(balances["rows"])==2 and all(r["document"]["kind"]=="journal" for r in balances["rows"])
    later=client.get(f"/hr/api/statement?person_id={profile_id}&date_from=2080-01-01",headers=JSON).get_json()
    assert later["opening"]==450000 and later["balance"]==450000 and later["rows"]==[]
    assert client.get(f"/hr/api/statement?person_id={profile_id}&page_size=1&page=2",headers=JSON).get_json()["rows"][0]["balance"]==450000
    assert command(client, "payroll.pay", {**payload, "amount": 450001}).status_code == 400
    assert command(client, "employee.update", {"person_id": profile_id, "version": 1, "salary": 700000, "effective_from": "2026-02-01"}).status_code == 400
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        assert AccountTransaction.query.filter_by(treasury_account_id=account_id, type="withdraw").count() == 1
        assert db.session.get(HRPayrollItem, rows[0]["id"]).paid == 200000
        assert JournalEntry.query.filter_by(reference_type="hr_payroll_payment", reference_id=paid.get_json()["id"]).one().amount == 200000
    from io import BytesIO
    from openpyxl import load_workbook
    exported=client.get(f"/hr/api/reports/statement/xlsx?person_id={profile_id}",headers=JSON)
    assert exported.status_code==200,exported.get_json()
    financial_book=load_workbook(BytesIO(exported.data))
    assert len(financial_book.worksheets)==3 and financial_book.active["F4"].value==450000
    assert financial_book.worksheets[1]["E3"].value==450000
    report=client.get("/hr/api/reports/payroll/xlsx?month=2026-02",headers=JSON)
    assert report.status_code==200,report.get_json()
    reconciliation=load_workbook(BytesIO(report.data)).worksheets[1]
    values=next(row for row in reconciliation.iter_rows(min_row=3,values_only=True) if row[0]==client.get(f"/hr/api/people/{profile_id}",headers=JSON).get_json()["person"]["code"])
    assert values[2:5]==(650000,650000,650000)
    assert values[5:9]==(200000,200000,200000,200000) and values[9]=="مطابق"


def test_hr_advance_is_receivable_then_deducted_once_from_approved_payroll(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    account_id = _fund(ctx)
    request = command(client, "advance.request", {"person_id": profile_id, "amount": 300000, "instalments": 3, "first_month": "2026-03", "reason": "test"})
    assert request.status_code == 200, request.get_json()
    advance_id = request.get_json()["id"]
    assert command(client, "advance.pay", {"advance_id": advance_id, "version": 1, "treasury_account_id": account_id}).status_code == 400
    assert command(client, "advance.decide", {"advance_id": advance_id, "version": 1, "status": "approved"}).status_code == 200
    paid = command(client, "advance.pay", {"advance_id": advance_id, "version": 2, "treasury_account_id": account_id})
    assert paid.status_code == 200, paid.get_json()
    assert command(client, "advance.pay", {"advance_id": advance_id, "version": 3, "treasury_account_id": account_id}).status_code == 400
    _run(client, ctx, "2026-03")
    item = client.get(f"/hr/api/workspace?tab=payroll&month=2026-03&person_id={profile_id}", headers=JSON).get_json()["rows"][0]
    assert item["gross"] == 650000 and item["deductions"] == 100000 and item["net"] == 550000
    advance = client.get(f"/hr/api/workspace?tab=advances&person_id={profile_id}", headers=JSON).get_json()["rows"][0]
    assert advance["repaid"] == 100000 and advance["remaining"] == 200000
    statement=client.get(f"/hr/api/statement?person_id={profile_id}",headers=JSON).get_json()
    assert statement["balance"]==550000 and statement["adjustment"]==100000 and statement["paid"]==0
    assert statement["advance_remaining"]==200000
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from models.journal_entry import JournalEntry
    from models.expense import Expense
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        assert JournalEntry.query.filter_by(reference_type="hr_advance_payment", reference_id=advance_id).one().amount == 300000
        assert JournalEntry.query.filter_by(reference_type="hr_advance_deduction", reference_id=item["id"]).one().amount == 100000
        assert not Expense.query.filter(Expense.title.contains(f"HR-A{advance_id}")).first()


def test_hr_stale_payroll_can_be_recalculated_but_approved_payroll_is_frozen(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    prepared = command(client, "payroll.prepare", {"month": "2026-04", "branch_id": ctx["ids"]["branch_id"]})
    assert prepared.status_code == 200, prepared.get_json()
    run_id = prepared.get_json()["id"]
    updated = command(client, "employee.update", {"person_id": profile_id, "version": 1,
                       "salary": 700000, "effective_from": "2026-04-01"})
    assert updated.status_code == 200, updated.get_json()
    assert command(client, "payroll.decide", {"run_id": run_id, "version": 1, "status": "reviewed"}).status_code == 409
    recalculated = command(client, "payroll.recalculate", {"run_id": run_id, "version": 1, "reason": "تصحيح عقد"})
    assert recalculated.status_code == 200, recalculated.get_json()
    assert command(client, "payroll.decide", {"run_id": run_id, "version": 2, "status": "reviewed"}).status_code == 200
    assert command(client, "payroll.decide", {"run_id": run_id, "version": 3, "status": "approved"}).status_code == 200
    row = client.get(f"/hr/api/workspace?tab=payroll&month=2026-04&person_id={profile_id}", headers=JSON).get_json()["rows"][0]
    assert row["gross"] == 700000
    assert command(client, "payroll.recalculate", {"run_id": run_id, "version": 4, "reason": "invalid"}).status_code == 400


def test_hr_insufficient_funds_rolls_back_payroll_amount_and_receipt(ctx):
    client = _client(ctx)
    profile_id = create_person(client, ctx)
    account_id = _fund(ctx, amount=100)
    _run(client, ctx, "2026-05")
    row = client.get(f"/hr/api/workspace?tab=payroll&month=2026-05&person_id={profile_id}", headers=JSON).get_json()["rows"][0]
    command_id = str(uuid4())
    result = command(client, "payroll.pay", {"item_id": row["id"], "treasury_account_id": account_id, "amount": 1000}, command_id)
    assert result.status_code == 400, result.get_json()
    from flask import g
    from extensions import db
    from extensions_tenant import init_tenant_db
    from modules.hr.models import HRPayrollItem, HRCommand
    from models.account_transaction import AccountTransaction
    with ctx["app"].app_context():
        g.tenant = TENANT
        init_tenant_db(TENANT)
        assert db.session.get(HRPayrollItem, row["id"]).paid == 0
        assert db.session.get(HRCommand, command_id) is None
        assert AccountTransaction.query.filter_by(treasury_account_id=account_id, type="withdraw").count() == 0
