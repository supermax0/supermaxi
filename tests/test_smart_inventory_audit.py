"""Real tenant databases: counting, authorization, movement conflicts and retry receipts."""
import os
from pathlib import Path
from uuid import uuid4

import pytest


@pytest.fixture
def audit_workspace():
    from app import app
    from extensions import db
    from extensions_tenant import init_tenant_db, clear_tenant_engine
    from flask import g, session
    from models.branch import Branch, BranchStock
    from models.employee import Employee
    from models.product import Product
    from models.role import Role, Permission
    from modules.inventory_audit.common import PERMISSIONS
    from modules.inventory_audit.schema import ensure_schema

    tenant = f"test_smart_inventory_audit_{os.getpid()}_{uuid4().hex[:8]}"
    with app.app_context():
        g.tenant = tenant
        init_tenant_db(tenant)
        branches = [Branch(code="A1", name="الأول", is_active=True, is_default=True),
                    Branch(code="A2", name="الثاني", is_active=True)]
        db.session.add_all(branches)
        db.session.flush()
        admin = Employee(name="Admin", username=uuid4().hex, password="test", role="admin", is_active=True, branch_id=branches[0].id)
        staff = Employee(name="Counter", username=uuid4().hex, password="test", role="cashier", is_active=True, branch_id=branches[0].id)
        db.session.add_all([admin, staff])
        db.session.flush()
        ensure_schema()
        role = Role(name="audit-blind-" + uuid4().hex)
        role.permissions = Permission.query.filter(Permission.name.in_(["inventory.audit.view", "inventory.audit.count"])).all()
        staff.roles = [role]
        group = str(uuid4())
        products = [Product(name="شاشة 55", sku="TV55", barcode="88055", catalog_group_id=group,
                            branch_id=branches[0].id, buy_price=100, sale_price=200, quantity=12, active=True),
                    Product(name="شاشة 55", sku="TV55", barcode="88055", catalog_group_id=group,
                            branch_id=branches[1].id, buy_price=100, sale_price=200, quantity=4, active=True),
                    Product(name="صنف صفري", sku="ZERO", branch_id=branches[0].id,
                            buy_price=50, sale_price=90, quantity=0, active=True)]
        db.session.add_all(products)
        db.session.flush()
        stocks = [BranchStock(branch_id=branches[0].id, product_id=products[0].id, quantity=12),
                  BranchStock(branch_id=branches[1].id, product_id=products[1].id, quantity=4),
                  BranchStock(branch_id=branches[0].id, product_id=products[2].id, quantity=0)]
        db.session.add_all(stocks)
        db.session.commit()

        def call(path, body=None, *, employee=admin, all_branches=True, session_tenant=None):
            with app.test_request_context(path, method="POST" if body is not None else "GET", json=body):
                g.tenant = tenant
                session.update(user_id=employee.id, role=employee.role, tenant_slug=session_tenant or tenant)
                g.branch, g.branches, g.view_all_branches = branches[0], branches, all_branches
                g.independent_branches_enabled, g.branch_scope_ready = True, True
                endpoint, kwargs = app.url_map.bind("localhost").match(path.split("?")[0], method="POST" if body is not None else "GET")
                return app.make_response(app.view_functions[endpoint](**kwargs))

        def create(*, employee=admin, locations=None):
            response = call("/api/v1/desktop/inventory-audits", {
                "idempotency_key": str(uuid4()), "ledger_kind": "branch",
                "location_keys": locations or [f"branch:{branches[0].id}"]}, employee=employee)
            assert response.status_code == 200, response.get_json()
            return response.get_json()["data"]["session_id"]

        yield {"call": call, "create": create, "admin": admin, "staff": staff, "products": products,
               "branches": branches, "stocks": stocks, "db": db, "tenant": tenant}
        db.session.remove()
    clear_tenant_engine(tenant)
    (Path(__file__).resolve().parents[1] / "tenants" / f"{tenant}.db").unlink(missing_ok=True)


def _command(f, session_id, action, **body):
    return f["call"](f"/api/v1/desktop/inventory-audits/{session_id}/commands",
                     {"idempotency_key": str(uuid4()), "action": action, **body})


def _ticket(f, session_id, line):
    response = _command(f, session_id, "begin_count", line_id=line["id"], version=line["version"])
    assert response.status_code == 200, response.get_json()
    return response.get_json()["data"]


def _lines(f, session_id, **kwargs):
    return f["call"](f"/api/v1/desktop/inventory-audits/{session_id}/lines", **kwargs).get_json()["data"]["items"]


def test_count_and_retry_preserve_stock_and_one_history_event(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    ticket = _ticket(f, sid, line)
    body = {"action": "save_count", "idempotency_key": str(uuid4()), "line_id": line["id"],
            "version": ticket["version"], "ticket": ticket["ticket"], "quantity": 10, "notes": "عد الرف"}
    path = f"/api/v1/desktop/inventory-audits/{sid}/commands"
    first, retry = f["call"](path, body), f["call"](path, body)
    assert first.status_code == retry.status_code == 200
    assert first.get_json() == retry.get_json()
    assert f["stocks"][0].quantity == 12
    result = _lines(f, sid)[0]
    assert result["difference"] == -2 and result["status"] == "shortage"
    from modules.inventory_audit.models import InventoryCountEvent, InventoryCountLine
    assert InventoryCountEvent.query.filter_by(action="save_count").count() == 1
    assert InventoryCountLine.query.filter_by(session_id=sid).count() == 2
    assert f["call"](path, {**body, "quantity": 11}).status_code == 409


def test_pause_resume_keeps_counts_stock_and_idempotent_history(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    ticket = _ticket(f, sid, line)
    assert _command(f, sid, "save_count", line_id=line["id"], version=ticket["version"],
                    ticket=ticket["ticket"], quantity=10).status_code == 200
    path = f"/api/v1/desktop/inventory-audits/{sid}"
    version = f["call"](path).get_json()["data"]["version"]
    pause = {"action": "pause", "idempotency_key": str(uuid4()), "version": version}
    first = f["call"](path + "/commands", pause)
    retry = f["call"](path + "/commands", pause)
    assert first.status_code == retry.status_code == 200 and first.get_json() == retry.get_json()
    paused = f["call"](path).get_json()["data"]
    assert paused["status"] == "paused" and paused["stats"]["counted"] == 1
    assert f["call"]("/api/v1/desktop/inventory-audits?status=paused").get_json()["data"]["total"] == 1
    counted = _lines(f, sid)[0]
    for action, body in [("begin_count", {"line_id": line["id"], "version": counted["version"]}),
                         ("submit", {"version": paused["version"], "acknowledge_partial": True}),
                         ("approve", {"version": paused["version"]}), ("post", {"version": paused["version"]})]:
        assert _command(f, sid, action, **body).status_code == 409
    assert _command(f, sid, "resume", version=version).status_code == 409
    assert _command(f, sid, "resume", version=paused["version"]).status_code == 200
    ticket = _ticket(f, sid, counted)
    assert _command(f, sid, "save_count", line_id=line["id"], version=ticket["version"],
                    ticket=ticket["ticket"], quantity=11).status_code == 200
    assert _lines(f, sid)[0]["counted_quantity"] == 11
    assert f["stocks"][0].quantity == 12
    from models.journal_entry import JournalEntry
    from models.account_transaction import AccountTransaction
    from modules.inventory_audit.models import InventoryCountEvent
    assert JournalEntry.query.count() == AccountTransaction.query.count() == 0
    assert InventoryCountEvent.query.filter_by(session_id=sid, action="pause").count() == 1


def test_pause_requires_closed_entry_and_invalidates_pre_pause_offline_lease(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    prepared = _command(f, sid, "prepare_offline", line_ids=[line["id"]]).get_json()["data"]["records"][0]
    ticket = _ticket(f, sid, line)
    path = f"/api/v1/desktop/inventory-audits/{sid}"
    version = f["call"](path).get_json()["data"]["version"]
    assert _command(f, sid, "pause", version=version).get_json()["code"] == "open_count"
    assert f["call"](path).get_json()["data"]["status"] == "counting"
    assert _command(f, sid, "cancel_count", line_id=line["id"], version=ticket["version"], ticket=ticket["ticket"]).status_code == 200
    prepared = _command(f, sid, "prepare_offline", line_ids=[line["id"]]).get_json()["data"]["records"][0]
    version = f["call"](path).get_json()["data"]["version"]
    assert _command(f, sid, "pause", version=version).status_code == 200
    version = f["call"](path).get_json()["data"]["version"]
    assert _command(f, sid, "resume", version=version).status_code == 200
    stale = _command(f, sid, "save_offline_count", line_id=line["id"], version=prepared["version"],
                     ticket=prepared["ticket"], quantity=10, counted_at=prepared["issued_at"])
    assert stale.status_code == 409 and stale.get_json()["code"] == "conflict"
    assert _lines(f, sid)[0]["counted_quantity"] is None
    assert _command(f, sid, "prepare_offline", line_ids=[line["id"]]).status_code == 200


def test_pause_resume_requires_count_permission(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    version = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]["version"]
    from models.role import Role, Permission
    role = Role(name="audit-read-only-" + uuid4().hex)
    role.permissions = Permission.query.filter_by(name="inventory.audit.view").all()
    f["staff"].roles = [role]
    f["db"].session.commit()
    for action in ("pause", "resume"):
        response = f["call"](f"/api/v1/desktop/inventory-audits/{sid}/commands", {
            "idempotency_key": str(uuid4()), "action": action, "version": version}, employee=f["staff"])
        assert response.status_code == 403


def test_stock_motion_during_count_detected_even_when_quantity_returns(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    ticket = _ticket(f, sid, line)
    from utils.branch_stock_service import set_branch_stock
    set_branch_stock(f["branches"][0].id, f["products"][0].id, 11)
    f["db"].session.commit()
    set_branch_stock(f["branches"][0].id, f["products"][0].id, 12)
    f["db"].session.commit()
    response = _command(f, sid, "save_count", line_id=line["id"], version=ticket["version"],
                        ticket=ticket["ticket"], quantity=12)
    assert response.status_code == 409 and response.get_json()["code"] == "stock_conflict"
    assert _lines(f, sid)[0]["counted_quantity"] is None


def test_blind_counter_cannot_read_system_cost_or_other_branches(audit_workspace):
    f = audit_workspace
    sid = f["create"](employee=f["staff"])
    line = _lines(f, sid, employee=f["staff"])[0]
    assert line["baseline_quantity"] is None and line["expected_quantity"] is None
    assert "unit_cost" not in line["product"]
    response = f["call"]("/api/v1/desktop/inventory-audits", {
        "idempotency_key": str(uuid4()), "location_keys": [f'branch:{f["branches"][1].id}']}, employee=f["staff"])
    assert response.status_code == 403
    other_sid = f["create"](locations=[f'branch:{f["branches"][1].id}'])
    assert f["call"](f"/api/v1/desktop/inventory-audits/{other_sid}", employee=f["staff"]).status_code == 403
    assert f["call"](f"/api/v1/desktop/inventory-audits/{sid}", session_tenant="other-tenant").status_code == 403


def test_optimistic_count_conflict_and_partial_review(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    ticket = _ticket(f, sid, line)
    assert _command(f, sid, "begin_count", line_id=line["id"], version=line["version"]).status_code == 409
    saved = _command(f, sid, "save_count", line_id=line["id"], version=ticket["version"], ticket=ticket["ticket"], quantity=12)
    assert saved.status_code == 200
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    assert row["stats"]["counted"] == row["stats"]["uncounted"] == 1
    assert _command(f, sid, "submit", version=row["version"]).get_json()["code"] == "incomplete"
    assert _command(f, sid, "submit", version=row["version"], acknowledge_partial=True).status_code == 200
    assert _command(f, sid, "begin_count", line_id=line["id"], version=ticket["version"] + 1).status_code == 409


def test_zero_stock_surplus_and_count_history_immutable(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = next(item for item in _lines(f, sid) if item["product"]["sku"] == "ZERO")
    ticket = _ticket(f, sid, line)
    assert _command(f, sid, "save_count", line_id=line["id"], version=ticket["version"], ticket=ticket["ticket"], quantity=1).status_code == 200
    result = next(item for item in _lines(f, sid) if item["id"] == line["id"])
    assert result["difference"] == 1 and result["status"] == "surplus"
    from sqlalchemy import text
    from sqlalchemy.exc import DatabaseError
    with pytest.raises(DatabaseError):
        f["db"].session.execute(text("DELETE FROM inventory_count_event"))
    f["db"].session.rollback()


def _count(f, sid, line, quantity):
    ticket = _ticket(f, sid, line)
    response = _command(f, sid, "save_count", line_id=line["id"], version=ticket["version"], ticket=ticket["ticket"], quantity=quantity)
    assert response.status_code == 200, response.get_json()
    return next(item for item in _lines(f, sid) if item["id"] == line["id"])


def _approve(f, sid):
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    response = _command(f, sid, "submit", version=row["version"], acknowledge_partial=True)
    assert response.status_code == 200, response.get_json()
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    response = _command(f, sid, "approve", version=row["version"])
    assert response.status_code == 200, response.get_json()
    return f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]


@pytest.mark.parametrize("quantity,journal_amount", [(10, 180), (12, 0), (15, 300)])
def test_explicit_fifo_basis_repair_keeps_physical_journal_separate_and_retry_once(audit_workspace, quantity, journal_amount):
    from models.inventory_lot import InventoryLot
    from models.journal_entry import JournalEntry
    from models.account_transaction import AccountTransaction
    from modules.inventory_audit.models import InventoryCountEvent
    from modules.finance.models import FinanceVoucher
    import json
    f = audit_workspace
    lots = [InventoryLot(product_id=f["products"][0].id, branch_id=f["branches"][0].id,
                        quantity=5, remaining_quantity=5, unit_cost=70),
            InventoryLot(product_id=f["products"][0].id, branch_id=f["branches"][0].id,
                        quantity=12, remaining_quantity=12, unit_cost=90),
            InventoryLot(product_id=f["products"][1].id, branch_id=f["branches"][1].id,
                        quantity=4, remaining_quantity=4, unit_cost=800)]
    f["db"].session.add_all(lots); f["db"].session.commit()
    sid = f["create"](); line = _count(f, sid, _lines(f, sid)[0], quantity)
    if quantity != 12:
        assert _command(f, sid, "review_line", line_id=line["id"], version=line["version"],
            resolution="loss" if quantity < 12 else "surplus", reason="عد فعلي مؤكد").status_code == 200
    row = _approve(f, sid)
    p = f"/api/v1/desktop/inventory-audits/{sid}"
    preview = f["call"](p + "?valuation=true").get_json()["data"]["valuation_review"]
    assert preview["removed_units"] == 5 and preview["affected_lines"] == 1
    assert [lot.remaining_quantity for lot in lots] == [5, 12, 4]
    if quantity < 12:
        assert _command(f, sid, "post", version=row["version"]).get_json()["code"] == "valuation_conflict"
    body = {"action":"post", "idempotency_key":str(uuid4()), "version":row["version"],
            "reconcile_valuation":True, "valuation_fingerprint":preview["fingerprint"], "valuation_reason":"مطابقة FIFO وفق رصيد الفرع"}
    response = f["call"](p + "/commands", body)
    assert response.status_code == 200, response.get_json()
    assert f["call"](p + "/commands", body).get_json() == response.get_json()
    f["db"].session.refresh(f["stocks"][0])
    assert f["stocks"][0].quantity == quantity
    total = sum(l.remaining_quantity for l in InventoryLot.query.filter_by(product_id=f["products"][0].id, branch_id=f["branches"][0].id))
    assert total == quantity and lots[2].remaining_quantity == 4
    assert sum(e.amount for e in JournalEntry.query.filter_by(reference_type="inventory_count")) == journal_amount
    assert AccountTransaction.query.count() == FinanceVoucher.query.count() == 0
    event = InventoryCountEvent.query.filter_by(session_id=sid, action="post").one()
    detail = json.loads(event.detail_json)
    assert detail["valuation_reconciliation"]["items"][0]["removed_value"] == 350
    assert detail["valuation_reconciliation"]["items"][0]["changes"][0]["lot_id"] == lots[0].id
    printed = f["call"](p + "/print").get_data(as_text=True)
    assert "مطابقة أساس طبقات التكلفة" in printed
    from models.role import Permission
    role = f["staff"].roles[0]
    role.permissions.extend(Permission.query.filter(Permission.name.in_(["inventory.audit.system", "inventory.audit.export"])).all())
    f["db"].session.commit()
    hidden = f["call"](p + "/history", employee=f["staff"]).get_json()["data"]["items"]
    assert "valuation_reconciliation" not in next(e for e in hidden if e["action"] == "post")["detail"]
    assert "مطابقة أساس طبقات التكلفة" not in f["call"](p + "/print", employee=f["staff"]).get_data(as_text=True)
    from openpyxl import load_workbook
    from io import BytesIO
    exported = f["call"](p + "/xlsx")
    exported.direct_passthrough = False
    workbook = load_workbook(BytesIO(exported.data))
    assert "مطابقة طبقات التكلفة" in workbook.sheetnames


def test_fifo_basis_fills_legacy_coverage_and_rejects_stale_snapshot_and_permission(audit_workspace):
    from models.inventory_lot import InventoryLot
    from models.journal_entry import JournalEntry
    f = audit_workspace
    lot = InventoryLot(product_id=f["products"][0].id, branch_id=f["branches"][0].id,
                       quantity=7, remaining_quantity=7, unit_cost=70)
    f["db"].session.add(lot); f["db"].session.commit()
    sid = f["create"](); line = _count(f, sid, _lines(f, sid)[0], 10)
    assert _command(f, sid, "review_line", line_id=line["id"], version=line["version"], resolution="loss", reason="مؤكد").status_code == 200
    row = _approve(f, sid); p = f"/api/v1/desktop/inventory-audits/{sid}"
    assert f["call"](p + "?valuation=true", employee=f["staff"]).status_code == 403
    preview = f["call"](p + "?valuation=true").get_json()["data"]["valuation_review"]
    assert preview["added_units"] == 5
    assert _command(f, sid, "post", version=row["version"], reconcile_valuation=True,
        valuation_fingerprint=preview["fingerprint"], valuation_reason="").status_code == 400
    lot.unit_cost = 80; f["db"].session.commit()
    assert _command(f, sid, "post", version=row["version"], reconcile_valuation=True,
        valuation_fingerprint=preview["fingerprint"], valuation_reason="مطابقة").get_json()["code"] == "valuation_conflict"
    assert lot.remaining_quantity == 7 and JournalEntry.query.count() == 0
    preview = f["call"](p + "?valuation=true").get_json()["data"]["valuation_review"]
    assert _command(f, sid, "post", version=row["version"], reconcile_valuation=True,
        valuation_fingerprint=preview["fingerprint"], valuation_reason="مطابقة").status_code == 200
    assert sum(x.remaining_quantity for x in InventoryLot.query.filter_by(product_id=f["products"][0].id)) == 10
    assert JournalEntry.query.filter_by(reference_type="inventory_count").one().amount == 160


def test_fifo_basis_and_stock_rollback_together_if_posting_fails(audit_workspace, monkeypatch):
    from models.inventory_lot import InventoryLot
    from models.journal_entry import JournalEntry
    from modules.inventory_audit.common import AuditError
    from modules.inventory_audit.models import InventoryCountEvent
    import modules.inventory_audit.posting as posting
    f = audit_workspace
    lot = InventoryLot(product_id=f["products"][0].id, branch_id=f["branches"][0].id,
                       quantity=17, remaining_quantity=17, unit_cost=100)
    f["db"].session.add(lot); f["db"].session.commit()
    sid = f["create"](); line = _count(f, sid, _lines(f, sid)[0], 10)
    assert _command(f, sid, "review_line", line_id=line["id"], version=line["version"], resolution="loss", reason="مؤكد").status_code == 200
    row = _approve(f, sid); p = f"/api/v1/desktop/inventory-audits/{sid}"
    preview = f["call"](p + "?valuation=true").get_json()["data"]["valuation_review"]
    def fail(*args, **kwargs): raise AuditError("فشل اختباري بعد مطابقة الطبقات", "test_failure", 409)
    monkeypatch.setattr(posting, "adjust_branch_stock", fail)
    assert _command(f, sid, "post", version=row["version"], reconcile_valuation=True,
        valuation_fingerprint=preview["fingerprint"], valuation_reason="مطابقة").get_json()["code"] == "test_failure"
    f["db"].session.refresh(lot); f["db"].session.refresh(f["stocks"][0])
    assert lot.remaining_quantity == 17 and f["stocks"][0].quantity == 12
    assert JournalEntry.query.count() == 0 and InventoryCountEvent.query.filter_by(action="post").count() == 0


def test_post_records_option_audit_when_legacy_aggregate_already_zero(audit_workspace):
    from models.product_color_variant import ProductColorVariant
    from models.inventory_audit import InventoryQuantityAudit
    from models.stock_movement import StockMovement
    f = audit_workspace
    product=f["products"][0]; product.meta_json='{"has_colors":true}'; product.quantity=0
    f["stocks"][0].quantity=0
    option=ProductColorVariant(product_id=product.id,color_name="أبيض",quantity=1,buy_price=100,sale_price=200)
    f["db"].session.add(option);f["db"].session.commit()
    sid=f["create"](); line=next(x for x in _lines(f,sid) if x["product"]["variant_id"]==option.id)
    line=_count(f,sid,line,0)
    assert _command(f,sid,"review_line",line_id=line["id"],version=line["version"],resolution="loss",reason="عد مؤكد").status_code==200
    row=_approve(f,sid)
    assert _command(f,sid,"post",version=row["version"]).status_code==200
    f["db"].session.refresh(option);f["db"].session.refresh(f["stocks"][0])
    assert option.quantity==f["stocks"][0].quantity==0
    movement=StockMovement.query.filter_by(source_type="inventory_count",source_line_id=line["id"]).one()
    evidence=InventoryQuantityAudit.query.filter_by(movement_id=movement.id).one()
    assert (evidence.before_quantity,evidence.after_quantity,evidence.delta_quantity)==(1,0,-1)
    assert evidence.get_metadata()["variant_id"]==option.id


def test_archived_catalog_copy_cannot_hide_active_canonical_branch_stock(audit_workspace, monkeypatch):
    from models.branch import BranchStock
    from models.product import Product
    from utils.branch_stock_service import _branch_product_id, get_branch_stock
    f=audit_workspace; canonical=f["products"][0]; archived=f["products"][1]
    monkeypatch.setattr("modules.inventory_audit.stock.independent_branches_enabled", lambda: False)
    archived.active=False;f["stocks"][1].quantity=0
    raw=BranchStock(branch_id=f["branches"][1].id,product_id=canonical.id,quantity=11)
    f["db"].session.add(raw);f["db"].session.commit()
    assert _branch_product_id(f["branches"][1].id,canonical.id)==canonical.id
    assert get_branch_stock(f["branches"][1].id,canonical.id)==11
    sid=f["create"](locations=[f'branch:{f["branches"][1].id}'])
    line=next(l for l in _lines(f,sid) if l["product"]["id"]==canonical.id)
    line=_count(f,sid,line,0)
    assert line["expected_quantity"]==11
    assert _command(f,sid,"review_line",line_id=line["id"],version=line["version"],resolution="loss",reason="مطابقة الجرد").status_code==200
    row=_approve(f,sid)
    assert _command(f,sid,"post",version=row["version"]).status_code==200
    f["db"].session.refresh(raw);f["db"].session.refresh(f["stocks"][0])
    assert raw.quantity==0 and f["stocks"][0].quantity==12


def test_approved_post_fifo_and_journal_are_atomic_and_retry_once(audit_workspace):
    f = audit_workspace
    from models.inventory_lot import InventoryLot
    from models.stock_movement import StockMovement
    from models.journal_entry import JournalEntry
    from models.account_transaction import AccountTransaction
    from modules.finance.models import FinanceVoucher
    lots = [InventoryLot(product_id=f["products"][0].id, branch_id=f["branches"][0].id,
                quantity=1, remaining_quantity=1, unit_cost=70),
            InventoryLot(product_id=f["products"][0].id, branch_id=f["branches"][0].id,
                quantity=11, remaining_quantity=11, unit_cost=90)]
    f["db"].session.add_all(lots)
    f["db"].session.commit()
    sid = f["create"]()
    line = _count(f, sid, _lines(f, sid)[0], 10)
    assert _command(f, sid, "review_line", line_id=line["id"], version=line["version"], resolution="loss", reason="عجز مؤكد بعد إعادة العد").status_code == 200
    row = _approve(f, sid)
    assert f["stocks"][0].quantity == 12  # approval itself never posts
    body = {"idempotency_key": str(uuid4()), "action": "post", "version": row["version"]}
    path = f"/api/v1/desktop/inventory-audits/{sid}/commands"
    before_cash, before_vouchers = AccountTransaction.query.count(), FinanceVoucher.query.count()
    response = f["call"](path, body)
    assert response.status_code == 200, response.get_json()
    assert f["call"](path, body).get_json() == response.get_json()
    f["db"].session.refresh(f["stocks"][0])
    assert f["stocks"][0].quantity == 10
    assert [lot.remaining_quantity for lot in lots] == [0, 10]
    movement = StockMovement.query.filter_by(source_type="inventory_count").one()
    assert movement.quantity == 2 and movement.total_cost == 160
    entry = JournalEntry.query.filter_by(reference_type="inventory_count").one()
    assert entry.amount == 160 and entry.debit_account.account_type == "expense" and entry.credit_account.account_type == "asset"
    assert AccountTransaction.query.count() == before_cash and FinanceVoucher.query.count() == before_vouchers
    assert _command(f, sid, "post", version=row["version"]).status_code == 409


def test_cross_location_opposites_are_possible_and_never_auto_transfer(audit_workspace):
    f = audit_workspace
    from models.branch import StockTransfer
    sid = f["create"](locations=[f'branch:{branch.id}' for branch in f["branches"]])
    first = next(line for line in _lines(f, sid) if line["product"]["id"] == f["products"][0].id)
    second = next(line for line in _lines(f, sid) if line["product"]["id"] == f["products"][1].id)
    first, second = _count(f, sid, first, 9), _count(f, sid, second, 6)
    before = StockTransfer.query.count()
    assert _command(f, sid, "investigate", line_id=first["id"]).status_code == 200
    data = f["call"](f'/api/v1/desktop/inventory-audits/{sid}/lines/{first["id"]}/details').get_json()["data"]["investigation"]
    assert data["classification"] == "possible" and data["auto_transfer"] is False
    assert data["company_comparison"]["net_difference"] == -1
    assert data["evidence"][0]["references"][0]["quantity"] == 2
    assert StockTransfer.query.count() == before
    assert [stock.quantity for stock in f["stocks"][:2]] == [12, 4]


def test_unexplained_review_cannot_post_and_stock_motion_after_approval_rejects(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _count(f, sid, _lines(f, sid)[0], 11)
    assert _command(f, sid, "review_line", line_id=line["id"], version=line["version"], resolution="unexplained", reason="نحتاج مراجعة").status_code == 200
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    assert _command(f, sid, "submit", version=row["version"], acknowledge_partial=True).status_code == 200
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    assert _command(f, sid, "approve", version=row["version"]).get_json()["code"] == "unreviewed"
    line = _lines(f, sid)[0]
    assert _command(f, sid, "review_line", line_id=line["id"], version=line["version"], resolution="loss", reason="تحقق المسؤول").status_code == 200
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    assert _command(f, sid, "approve", version=row["version"]).status_code == 200
    from utils.branch_stock_service import set_branch_stock
    set_branch_stock(f["branches"][0].id, f["products"][0].id, 13)
    f["db"].session.commit()
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    assert _command(f, sid, "post", version=row["version"]).get_json()["code"] == "stock_conflict"


def test_cancel_open_count_keeps_previous_observation_and_unblocks_review(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _count(f, sid, _lines(f, sid)[0], 12)
    ticket = _ticket(f, sid, line)
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    assert _command(f, sid, "submit", version=row["version"], acknowledge_partial=True).get_json()["code"] == "open_count"
    assert _command(f, sid, "cancel_count", line_id=line["id"], version=ticket["version"], ticket=ticket["ticket"]).status_code == 200
    assert _lines(f, sid)[0]["counted_quantity"] == 12
    row = f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    assert _command(f, sid, "submit", version=row["version"], acknowledge_partial=True).status_code == 200


def _offline(f, sid, line, **kwargs):
    response = _command(f, sid, "prepare_offline", line_ids=[line["id"]], **kwargs)
    assert response.status_code == 200, response.get_json()
    return response.get_json()["data"]["records"][0]


def test_offline_count_uses_snapshot_time_and_exact_retry_without_stock_posting(audit_workspace):
    from datetime import datetime, timezone
    from modules.inventory_audit.models import InventoryCountEvent
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    record = _offline(f, sid, line)
    counted_at = datetime.now(timezone.utc).isoformat()
    payload = {"action":"save_offline_count", "idempotency_key":str(uuid4()), "line_id":line["id"],
               "version":record["version"], "ticket":record["ticket"], "quantity":9, "counted_at":counted_at, "notes":"عد محلي"}
    path = f"/api/v1/desktop/inventory-audits/{sid}/commands"
    result = f["call"](path, payload)
    assert result.status_code == 200, result.get_json()
    assert f["call"](path, payload).get_json() == result.get_json()
    assert _lines(f, sid)[0]["counted_quantity"] == 9
    assert f["stocks"][0].quantity == 12
    assert InventoryCountEvent.query.filter_by(action="save_offline_count").count() == 1
    event = InventoryCountEvent.query.filter_by(action="save_offline_count").one()
    import json
    assert json.loads(event.detail_json)["offline"] is True


def test_offline_net_zero_movements_conflict_and_actor_ticket_isolation(audit_workspace):
    from datetime import datetime, timezone
    from utils.branch_stock_service import set_branch_stock
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    record = _offline(f, sid, line)
    payload = {"action":"save_offline_count", "idempotency_key":str(uuid4()), "line_id":line["id"],
               "version":record["version"], "ticket":record["ticket"], "quantity":10,
               "counted_at":datetime.now(timezone.utc).isoformat()}
    path = f"/api/v1/desktop/inventory-audits/{sid}/commands"
    assert f["call"](path, payload, employee=f["staff"]).get_json()["code"] == "count_ticket"
    set_branch_stock(f["branches"][0].id, f["products"][0].id, 11)
    f["db"].session.commit()
    set_branch_stock(f["branches"][0].id, f["products"][0].id, 12)
    f["db"].session.commit()
    assert f["call"](path, payload).get_json()["code"] == "stock_conflict"
    assert _lines(f, sid)[0]["counted_quantity"] is None


def test_independent_offline_snapshots_do_not_overwrite_online_ticket_and_stale_version_rejected(audit_workspace):
    from datetime import datetime, timezone
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f, sid)[0]
    first, second = _offline(f, sid, line), _offline(f, sid, line)
    assert first["ticket"] != second["ticket"]
    online = _ticket(f, sid, line)
    assert _command(f, sid, "save_offline_count", line_id=line["id"], version=first["version"], ticket=first["ticket"],
                    quantity=10, counted_at=datetime.now(timezone.utc).isoformat()).get_json()["code"] == "conflict"
    assert _command(f, sid, "save_count", line_id=line["id"], version=online["version"], ticket=online["ticket"], quantity=12).status_code == 200


def test_reports_have_saved_scope_zero_vs_uncounted_and_do_not_expose_blind_cost(audit_workspace):
    from io import BytesIO
    from openpyxl import load_workbook
    from models.role import Permission
    f = audit_workspace
    role = f["staff"].roles[0]
    role.permissions.append(Permission.query.filter_by(name="inventory.audit.export").one())
    f["db"].session.commit()
    sid = f["create"]()
    line = _count(f, sid, _lines(f, sid)[0], 0)
    print_page = f["call"](f"/api/v1/desktop/inventory-audits/{sid}/print")
    assert print_page.status_code == 200 and print_page.mimetype == "text/html"
    assert "إجمالي النقص" in print_page.get_data(as_text=True)
    blind = f["call"](f"/api/v1/desktop/inventory-audits/{sid}/print", employee=f["staff"])
    assert blind.status_code == 200
    assert "تكلفة التسوية د.ع" not in blind.get_data(as_text=True) and "النظام عند العد" not in blind.get_data(as_text=True)
    sheet_response = f["call"](f"/api/v1/desktop/inventory-audits/{sid}/xlsx")
    sheet_response.direct_passthrough = False
    workbook = load_workbook(BytesIO(sheet_response.get_data()))
    sheet = workbook["جرد المخزون"]
    assert sheet["H2"].value == 0 and sheet["H3"].value is None
    assert sheet["I2"].value == 12 and sheet["J2"].value == -12
    assert f["stocks"][0].quantity == 12 and line["counted_quantity"] == 0


def test_explicit_transfer_draft_is_linked_idempotent_and_does_not_change_balances(audit_workspace):
    from models.branch import StockTransfer
    f = audit_workspace
    sid = f["create"](locations=[f'branch:{branch.id}' for branch in f["branches"]])
    a = next(line for line in _lines(f, sid) if line["product"]["id"] == f["products"][0].id)
    b = next(line for line in _lines(f, sid) if line["product"]["id"] == f["products"][1].id)
    a, b = _count(f, sid, a, 9), _count(f, sid, b, 6)
    payload = {"action":"transfer_draft","idempotency_key":str(uuid4()),"line_id":a["id"],"version":a["version"],
               "target_line_id":b["id"],"quantity":2,"reason":"راجعت سجل إرسال غير مكتمل", "acknowledge_draft":True}
    path = f"/api/v1/desktop/inventory-audits/{sid}/commands"
    result = f["call"](path, payload)
    assert result.status_code == 200, result.get_json()
    assert f["call"](path, payload).get_json() == result.get_json()
    transfer = f["db"].session.get(StockTransfer, result.get_json()["data"]["transfer_id"])
    assert transfer.status == "draft" and transfer.lines[0].quantity == 2
    assert [stock.quantity for stock in f["stocks"][:2]] == [12,4]


def test_count_variants_preserves_sibling_qty_and_posts_one_source_per_size(audit_workspace):
    import json
    from models.product_color_variant import ProductColorVariant
    f = audit_workspace
    product = f["products"][0]
    product.meta_json = json.dumps({"has_colors":True})
    options = [ProductColorVariant(product_id=product.id,color_name="أسود / L",quantity=5,buy_price=100),
               ProductColorVariant(product_id=product.id,color_name="أسود / XL",quantity=7,buy_price=120)]
    f["db"].session.add_all(options)
    f["db"].session.commit()
    sid = f["create"]()
    line = next(line for line in _lines(f, sid) if line["product"]["variant_name"] == "أسود / L")
    line = _count(f, sid, line, 3)
    assert line["difference"] == -2
    assert _command(f,sid,"review_line",line_id=line["id"],version=line["version"],resolution="loss",reason="عد المقاس المحدد").status_code == 200
    row = _approve(f,sid)
    result = _command(f,sid,"post",version=row["version"])
    assert result.status_code == 200,result.get_json()
    assert [option.quantity for option in options] == [3,7]
    f["db"].session.refresh(product)
    assert product.quantity == 10 and f["stocks"][0].quantity == 10


def test_warehouse_on_hand_count_ignores_reservation_subtraction_and_only_posts_warehouse_ledger(audit_workspace):
    from modules.wholesale.models.warehouse import Warehouse,WarehouseStock,WarehouseMovement
    f = audit_workspace
    warehouse = Warehouse(branch_id=f["branches"][0].id,code="W1",name="مخزن",kind="main",is_active=True)
    f["db"].session.add(warehouse)
    f["db"].session.flush()
    stock = WarehouseStock(warehouse_id=warehouse.id,product_id=f["products"][0].id,quantity=12,reserved_quantity=3,in_transit_quantity=2)
    f["db"].session.add(stock)
    f["db"].session.commit()
    result = f["call"]("/api/v1/desktop/inventory-audits",{"idempotency_key":str(uuid4()),"ledger_kind":"warehouse","location_keys":[f"warehouse:{warehouse.id}"]})
    assert result.status_code == 200,result.get_json()
    sid = result.get_json()["data"]["session_id"]
    line = _count(f,sid,_lines(f,sid)[0],10)
    assert line["expected_quantity"] == 12 and line["difference"] == -2
    assert _command(f,sid,"review_line",line_id=line["id"],version=line["version"],resolution="loss",reason="عد الرف الفعلي").status_code == 200
    row = _approve(f,sid)
    result = _command(f,sid,"post",version=row["version"])
    assert result.status_code == 200,result.get_json()
    assert stock.quantity == 10 and stock.reserved_quantity == 3 and stock.in_transit_quantity == 2
    assert f["stocks"][0].quantity == 12
    movement = WarehouseMovement.query.filter_by(reference_type="inventory_count").one()
    assert movement.quantity_before == 12 and movement.quantity_after == 10 and movement.quantity == -2


def test_legacy_audit_cannot_bypass_approval_or_create_fake_cash(audit_workspace):
    from models.account_transaction import AccountTransaction
    f = audit_workspace
    before = AccountTransaction.query.count()
    response = f["call"]("/inventory/save-audit",{"items":[{"id":f["products"][0].id,"actual_qty":3}]})
    assert response.status_code == 409 and response.get_json()["code"] == "audit_workflow_required"
    assert f["stocks"][0].quantity == 12 and AccountTransaction.query.count() == before


def test_readonly_product_timeline_contains_actual_count_once(audit_workspace):
    from utils.inventory_movements import get_product_inventory_movements
    f = audit_workspace
    sid = f["create"]()
    line = _count(f,sid,_lines(f,sid)[0],10)
    assert _command(f,sid,"review_line",line_id=line["id"],version=line["version"],resolution="loss",reason="جرد مثبت").status_code == 200
    row = _approve(f,sid)
    assert _command(f,sid,"post",version=row["version"]).status_code == 200
    movements = get_product_inventory_movements(f["products"][0].id,branch_id=f["branches"][0].id)
    counted = [movement for movement in movements if movement["type"] == "inventory_count_adjustment"]
    assert len(counted) == 1 and counted[0]["quantity_out"] == 2 and counted[0]["balance_after"] == 10
    assert f["stocks"][0].quantity == 10  # reading never reposts


def test_closed_period_and_posted_result_journal_history_are_immutable(audit_workspace):
    from models.financial_period_close import FinancialPeriodClose
    from models.journal_entry import JournalEntry
    from modules.finance.common import today
    from sqlalchemy import text
    from sqlalchemy.exc import DatabaseError
    f = audit_workspace
    sid = f["create"]()
    line = _count(f,sid,_lines(f,sid)[0],11)
    assert _command(f,sid,"review_line",line_id=line["id"],version=line["version"],resolution="loss",reason="عجز مثبت").status_code == 200
    row = _approve(f,sid)
    close = FinancialPeriodClose(period_year=today().year,period_month=today().month,closed_by=f["admin"].id)
    f["db"].session.add(close);f["db"].session.commit()
    result = _command(f,sid,"post",version=row["version"])
    assert result.status_code == 400 and result.get_json()["code"] == "period_closed"
    assert f["stocks"][0].quantity == 12 and JournalEntry.query.filter_by(reference_type="inventory_count").count() == 0
    f["db"].session.delete(close);f["db"].session.commit()
    assert _command(f,sid,"post",version=row["version"]).status_code == 200
    for statement in ["UPDATE inventory_count_session SET title='changed' WHERE id=:id",
                      "UPDATE inventory_count_line SET counted_quantity=0 WHERE session_id=:id",
                      "DELETE FROM inventory_count_event WHERE session_id=:id",
                      "UPDATE journal_entry SET amount=1 WHERE reference_type='inventory_count' AND reference_id=:id"]:
        with pytest.raises(DatabaseError):
            f["db"].session.execute(text(statement),{"id":sid});f["db"].session.commit()
        f["db"].session.rollback()


def test_snapshot_and_actor_names_remain_saved_after_catalog_and_employee_rename(audit_workspace):
    f = audit_workspace
    sid = f["create"]()
    line = _count(f,sid,_lines(f,sid)[0],12)
    original = line["product"]["name"]
    f["products"][0].name = "اسم جديد"
    f["products"][0].barcode = "NEW-CODE"
    f["admin"].name = "New actor name"
    f["db"].session.commit()
    saved = _lines(f,sid)[0]
    assert saved["product"]["name"] == original and saved["product"]["current_barcode"] == "NEW-CODE"
    history = f["call"](f"/api/v1/desktop/inventory-audits/{sid}/history").get_json()["data"]["items"]
    assert all(event["actor_name"] == "Admin" for event in history)
    response = f["call"](f"/api/v1/desktop/inventory-audits/{sid}/lines?location_key=branch:{f['branches'][0].id}&q=NEW-CODE&exact=true")
    assert response.get_json()["data"]["total"] == 1


def test_id_model_search_and_private_metadata_do_not_expose_prices(audit_workspace):
    import json
    f = audit_workspace
    f["products"][0].meta_json = json.dumps({"model":"MODEL-55", "private_price":986543210})
    f["db"].session.commit()
    sid = f["create"]()
    path = f"/api/v1/desktop/inventory-audits/{sid}/lines"
    assert f["call"](path+"?q=MODEL-55").get_json()["data"]["total"] == 1
    assert f["call"](path+f"?q={f['products'][0].id}&exact=true").get_json()["data"]["total"] == 1
    assert f["call"](path+"?q=986543210",employee=f["staff"]).get_json()["data"]["total"] == 0


def test_expired_offline_ticket_does_not_apply_and_other_counter_needs_reasoned_release(audit_workspace):
    from datetime import datetime,timedelta,timezone
    from modules.inventory_audit.models import InventoryCountOfflineTicket
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f,sid)[0]
    record = _offline(f,sid,line)
    ticket = f["db"].session.get(InventoryCountOfflineTicket,record["ticket"])
    ticket.created_at -= timedelta(hours=25)
    f["db"].session.commit()
    result = _command(f,sid,"save_offline_count",line_id=line["id"],version=line["version"],ticket=record["ticket"],quantity=9,counted_at=datetime.now(timezone.utc).isoformat())
    assert result.get_json()["code"] == "snapshot_expired"
    own = _ticket(f,sid,line)
    response = f["call"](f"/api/v1/desktop/inventory-audits/{sid}/commands",{"idempotency_key":str(uuid4()),"action":"begin_count","line_id":line["id"],"version":own["version"]},employee=f["staff"])
    assert response.get_json()["code"] == "count_in_progress"
    assert _command(f,sid,"release_ticket",line_id=line["id"],version=own["version"],reason="إغلاق إدخال متروك بعد مراجعة الموظف").status_code == 200
    assert _lines(f,sid)[0]["counted_quantity"] is None


@pytest.mark.parametrize("invalid",[-1,1.25,True,"12",None,1000000001])
def test_invalid_physical_quantities_never_change_count_or_stock(audit_workspace,invalid):
    f = audit_workspace
    sid = f["create"]()
    line = _lines(f,sid)[0]
    ticket = _ticket(f,sid,line)
    result = _command(f,sid,"save_count",line_id=line["id"],version=ticket["version"],ticket=ticket["ticket"],quantity=invalid)
    assert result.status_code == 400
    assert _lines(f,sid)[0]["counted_quantity"] is None and f["stocks"][0].quantity == 12


def _source_invoice(f, status, *, warehouse_id=None, agent=None):
    from models.customer import Customer
    from models.invoice import Invoice
    from models.order_item import OrderItem
    customer = Customer(name="عميل دليل الجرد", phone="07700000881")
    f["db"].session.add(customer)
    f["db"].session.flush()
    invoice = Invoice(customer_id=customer.id, customer_name=customer.name, employee_id=f["admin"].id,
        branch_id=f["branches"][0].id, status=status, payment_status="غير مسدد", total=200,
        paid_amount=0, barcode="AUDIT-"+uuid4().hex[:8], warehouse_id=warehouse_id,
        delivery_agent_id=agent.id if agent else None)
    f["db"].session.add(invoice)
    f["db"].session.flush()
    item = OrderItem(invoice_id=invoice.id, product_id=f["products"][0].id, product_name="شاشة 55",
        price=200, quantity=1, cost=100, total=200, fulfillment_branch_id=f["branches"][0].id, warehouse_id=warehouse_id)
    f["db"].session.add(item)
    f["db"].session.flush()
    return invoice, item


def test_real_orders_returns_and_agent_links_are_possible_evidence_not_custody_or_loss(audit_workspace):
    from models.delivery_agent import DeliveryAgent
    from models.stock_movement import StockMovement
    from models.role import Permission
    f = audit_workspace
    agent = DeliveryAgent(name="مندوب الاختبار", branch_id=f["branches"][0].id)
    f["db"].session.add(agent)
    f["db"].session.flush()
    late, _ = _source_invoice(f,"جاري الشحن",agent=agent)
    returned, return_item = _source_invoice(f,"راجع")
    f["db"].session.add(StockMovement(movement_key="SOURCE-"+uuid4().hex,movement_type="order_sale",direction="out",
        product_id=f["products"][0].id, branch_id=f["branches"][0].id,quantity=1,unit_cost=100,total_cost=100,
        source_type="invoice", source_id=returned.id,source_line_id=return_item.id))
    f["db"].session.commit()
    sid = f["create"]()
    line = _count(f,sid,_lines(f,sid)[0],10)
    assert _command(f,sid,"investigate",line_id=line["id"]).status_code == 200
    path=f'/api/v1/desktop/inventory-audits/{sid}/lines/{line["id"]}/details'
    analysis = f["call"](path).get_json()["data"]["investigation"]
    evidence={entry["kind"]:entry for entry in analysis["evidence"]}
    assert {ref["id"] for ref in evidence["orders"]["references"]} == {late.id, returned.id}
    assert evidence["returns"]["references"][0]["incoming_movements"] == 0
    assert evidence["couriers"]["references"][0]["verified_custody"] is False
    assert analysis["classification"] == "possible" and analysis["auto_transfer"] is False
    assert analysis["coverage"]["serials"] == "not_available"
    assert f["stocks"][0].quantity == 12
    role=f["staff"].roles[0]
    role.permissions.extend(Permission.query.filter(Permission.name.in_(["inventory.audit.system","inventory.audit.investigate"])).all())
    f["db"].session.commit()
    masked=f["call"](path,employee=f["staff"]).get_json()["data"]["investigation"]
    assert not any(entry["kind"] in {"orders","returns","couriers"} for entry in masked["evidence"])
    assert masked["coverage"]["orders"] == "restricted"


def test_warehouse_investigation_uses_its_own_sale_movements_and_purchase_location(audit_workspace):
    from modules.wholesale.models.warehouse import Warehouse,WarehouseStock,WarehouseMovement
    from models.supplier import Supplier
    from models.purchase import Purchase
    from models.purchase_item import PurchaseItem
    f=audit_workspace
    warehouse=Warehouse(branch_id=f["branches"][0].id,code="DOCW",name="مخزن المستندات",is_active=True)
    f["db"].session.add(warehouse)
    f["db"].session.flush()
    f["db"].session.add(WarehouseStock(warehouse_id=warehouse.id,product_id=f["products"][0].id,quantity=12))
    invoice,_=_source_invoice(f,"راجع",warehouse_id=warehouse.id)
    f["db"].session.add(WarehouseMovement(warehouse_id=warehouse.id,product_id=f["products"][0].id,
        movement_type="out",quantity=-1,quantity_before=13,quantity_after=12,
        reference_type="wholesale_sale",reference_id=invoice.id))
    supplier=Supplier(name="مورد دليل الجرد")
    f["db"].session.add(supplier)
    f["db"].session.flush()
    purchases=[]
    for site in (None,warehouse.id):
        purchase=Purchase(supplier_id=supplier.id,product_id=f["products"][0].id,branch_id=f["branches"][0].id,
            warehouse_id=site,quantity=1,price=100,total=100,status="draft",stock_applied=False)
        f["db"].session.add(purchase)
        f["db"].session.flush()
        f["db"].session.add(PurchaseItem(purchase_id=purchase.id,product_id=f["products"][0].id,quantity=1))
        purchases.append(purchase.id)
    f["db"].session.commit()
    response=f["call"]("/api/v1/desktop/inventory-audits",{"idempotency_key":str(uuid4()),"ledger_kind":"warehouse","location_keys":[f"warehouse:{warehouse.id}"]})
    sid=response.get_json()["data"]["session_id"]
    line=_count(f,sid,_lines(f,sid)[0],11)
    assert _command(f,sid,"investigate",line_id=line["id"]).status_code == 200
    analysis=f["call"](f'/api/v1/desktop/inventory-audits/{sid}/lines/{line["id"]}/details').get_json()["data"]["investigation"]
    evidence={entry["kind"]:entry for entry in analysis["evidence"]}
    assert evidence["orders"]["references"][0]["outgoing_movements"] == 1
    assert evidence["returns"]["references"][0]["id"] == invoice.id
    assert [ref["id"] for ref in evidence["purchases"]["references"]] == [purchases[1]]
    branch_sid=f["create"]()
    branch_line=_count(f,branch_sid,_lines(f,branch_sid)[0],10)
    assert _command(f,branch_sid,"investigate",line_id=branch_line["id"]).status_code == 200
    branch_analysis=f["call"](f'/api/v1/desktop/inventory-audits/{branch_sid}/lines/{branch_line["id"]}/details').get_json()["data"]["investigation"]
    branch_evidence={entry["kind"]:entry for entry in branch_analysis["evidence"]}
    assert "orders" not in branch_evidence and "returns" not in branch_evidence
    assert [ref["id"] for ref in branch_evidence["purchases"]["references"]] == [purchases[0]]


def test_linked_transfer_draft_requires_resolution_and_cannot_be_created_twice(audit_workspace):
    f=audit_workspace
    sid=f["create"](locations=[f'branch:{branch.id}' for branch in f["branches"]])
    first=next(line for line in _lines(f,sid) if line["product"]["id"]==f["products"][0].id)
    second=next(line for line in _lines(f,sid) if line["product"]["id"]==f["products"][1].id)
    first,second=_count(f,sid,first,10),_count(f,sid,second,6)
    body=dict(line_id=first["id"],version=first["version"],target_line_id=second["id"],quantity=2,reason="مراجعة انتقال محتمل",acknowledge_draft=True)
    response=_command(f,sid,"transfer_draft",**body)
    assert response.status_code==200,response.get_json()
    duplicate=_command(f,sid,"transfer_draft",**body)
    assert duplicate.status_code==409
    current=f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]
    response=_command(f,sid,"submit",version=current["version"],acknowledge_partial=True)
    assert response.status_code==409 and response.get_json()["code"]=="pending_transfer"
    assert [row.quantity for row in f["stocks"]]==[12,4,0]


@pytest.mark.parametrize("ledger_kind",["branch","warehouse"])
def test_stock_reservations_block_posting_before_any_sibling_or_journal_changes(audit_workspace,ledger_kind):
    from models.journal_entry import JournalEntry
    from modules.wholesale.models.warehouse import Warehouse,WarehouseStock
    from utils.desktop_stock_reservations import DesktopStockAllocation,ensure_reservation_schema
    f=audit_workspace
    if ledger_kind=="branch":
        ensure_reservation_schema()
        f["db"].session.add(DesktopStockAllocation(id=str(uuid4()),grant_id=str(uuid4()),product_id=f["products"][0].id,
            requested_product_id=f["products"][0].id,branch_id=f["branches"][0].id,stock_id=f["stocks"][0].id,quantity=2))
        f["db"].session.commit()
        sid=f["create"]()
        expected_code="offline_escrow"
    else:
        warehouse=Warehouse(branch_id=f["branches"][0].id,code="HELD",name="مخزن محجوز",is_active=True)
        f["db"].session.add(warehouse)
        f["db"].session.flush()
        stock=WarehouseStock(warehouse_id=warehouse.id,product_id=f["products"][0].id,quantity=12,reserved_quantity=11)
        f["db"].session.add(stock)
        f["db"].session.commit()
        response=f["call"]("/api/v1/desktop/inventory-audits",{"idempotency_key":str(uuid4()),"ledger_kind":"warehouse","location_keys":[f"warehouse:{warehouse.id}"]})
        sid=response.get_json()["data"]["session_id"]
        expected_code="reserved_stock"
    line=_count(f,sid,_lines(f,sid)[0],10)
    assert _command(f,sid,"review_line",line_id=line["id"],version=line["version"],resolution="loss",reason="عد الموجود الفعلي").status_code==200
    row=_approve(f,sid)
    response=_command(f,sid,"post",version=row["version"])
    assert response.status_code==409 and response.get_json()["code"]==expected_code
    assert f["stocks"][0].quantity==12
    if ledger_kind=="warehouse": assert stock.quantity==12 and stock.reserved_quantity==11
    assert JournalEntry.query.filter_by(reference_type="inventory_count").count()==0
    assert f["call"](f"/api/v1/desktop/inventory-audits/{sid}").get_json()["data"]["status"]=="approved"
