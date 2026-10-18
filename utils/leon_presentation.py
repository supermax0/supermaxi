"""Structured LEON chat presentation and export helpers.

The assistant narrative stays conversational while this module turns verified
tool results into a small, stable UI contract.  It never queries the database;
therefore it cannot widen the permission scope enforced by ai_assistant_tools.
"""
from __future__ import annotations

import csv
import io
import json
import os
from datetime import datetime
from html import escape
from typing import Any


def _number(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _metric(label: str, value: Any, *, kind: str = "number", tone: str = "purple") -> dict:
    return {"label": label, "value": _number(value) if kind in {"number", "money"} else str(value or ""), "kind": kind, "tone": tone}


def _metric_widget(title: str, items: list[dict], *, subtitle: str = "") -> dict:
    return {"type": "metric_grid", "title": title, "subtitle": subtitle, "items": items[:6]}


def _bar_widget(title: str, values: dict | list, *, value_key: str = "value", money: bool = True) -> dict | None:
    rows: list[dict] = []
    if isinstance(values, dict):
        for label, raw in list(values.items())[:12]:
            value = raw.get(value_key, raw.get("total", 0)) if isinstance(raw, dict) else raw
            rows.append({"label": str(label), "value": _number(value)})
    elif isinstance(values, list):
        for raw in values[:12]:
            if not isinstance(raw, dict):
                continue
            label = raw.get("label") or raw.get("name") or raw.get("product") or raw.get("customer") or raw.get("branch") or "—"
            rows.append({"label": str(label), "value": _number(raw.get(value_key, raw.get("total", 0)))})
    if not rows:
        return None
    return {"type": "bar_list", "title": title, "kind": "money" if money else "number", "items": rows}


def _table_widget(title: str, columns: list[tuple[str, str, str]], rows: list[dict]) -> dict | None:
    if not rows:
        return None
    safe_columns = [{"key": key, "label": label, "kind": kind} for key, label, kind in columns]
    safe_rows = []
    for raw in rows[:20]:
        if not isinstance(raw, dict):
            continue
        safe_rows.append({key: raw.get(key, "") for key, _, _ in columns})
    if not safe_rows:
        return None
    return {"type": "table", "title": title, "columns": safe_columns, "rows": safe_rows}


def _widgets_for_tool(name: str, data: dict) -> list[dict]:
    if not isinstance(data, dict) or data.get("error") or data.get("restricted"):
        message = data.get("error") or data.get("message") if isinstance(data, dict) else ""
        return [{"type": "notice", "tone": "warning", "title": "النتيجة غير متاحة", "message": str(message or "تعذّر قراءة هذه البيانات.")}]

    widgets: list[dict] = []
    if name == "get_sales_summary":
        widgets.append(_metric_widget("ملخص المبيعات", [
            _metric("الطلبات", data.get("orders_count"), tone="blue"),
            _metric("إجمالي المبيعات", data.get("total_sales"), kind="money"),
            _metric("المحصّل", data.get("collected_cash"), kind="money", tone="green"),
            _metric("المتبقي", data.get("remaining_receivable"), kind="money", tone="amber"),
        ], subtitle=str(data.get("period") or "")))
        bar = _bar_widget("توزيع المبيعات", data.get("breakdown") or {}, value_key="total")
        if bar:
            widgets.append(bar)
    elif name == "get_profit_summary":
        net = _number(data.get("net_profit"))
        widgets.append(_metric_widget("الأرباح", [
            _metric("المبيعات", data.get("sales_total"), kind="money"),
            _metric("تكلفة البضاعة", data.get("cogs"), kind="money", tone="blue"),
            _metric("المصاريف", data.get("expenses"), kind="money", tone="amber"),
            _metric("صافي الربح", net, kind="money", tone="green" if net >= 0 else "red"),
        ], subtitle=str(data.get("period") or "")))
    elif name == "get_expenses":
        widgets.append(_metric_widget("المصاريف", [
            _metric("الإجمالي", data.get("total"), kind="money", tone="red"),
            _metric("عدد الحركات", data.get("count"), tone="blue"),
        ]))
        bar = _bar_widget("المصاريف حسب الفئة", data.get("by_category") or {})
        if bar:
            widgets.append(bar)
        table = _table_widget("آخر المصاريف", [("date", "التاريخ", "text"), ("title", "المصروف", "text"), ("amount", "المبلغ", "money")], data.get("recent") or [])
        if table:
            widgets.append(table)
    elif name in {"get_treasury_accounts", "get_cash_movements"}:
        if name == "get_treasury_accounts":
            widgets.append(_metric_widget("السيولة والخزائن", [
                _metric("إجمالي السيولة", data.get("total_liquidity"), kind="money", tone="green"),
                _metric("الصندوق", data.get("cash_box_balance"), kind="money"),
                _metric("البنوك", data.get("bank_total"), kind="money", tone="blue"),
            ]))
        else:
            widgets.append(_metric_widget("حركة الصندوق", [
                _metric("رصيد البداية", data.get("opening_balance"), kind="money"),
                _metric("الداخل", data.get("total_in"), kind="money", tone="green"),
                _metric("الخارج", data.get("total_out"), kind="money", tone="red"),
                _metric("الرصيد الحالي", data.get("current_balance_now"), kind="money", tone="blue"),
            ]))
            table = _table_widget("آخر حركات الصندوق", [("date", "التاريخ", "text"), ("type", "النوع", "text"), ("amount", "المبلغ", "money"), ("description", "الوصف", "text")], data.get("movements") or [])
            if table:
                widgets.append(table)
    elif name == "get_supplier_debts":
        widgets.append(_metric_widget("ديون الموردين", [
            _metric("إجمالي المتبقي", data.get("total_remaining_debt"), kind="money", tone="amber"),
            _metric("عدد الموردين", data.get("suppliers_count"), tone="blue"),
        ]))
        table = _table_widget("أرصدة الموردين", [("name", "المورد", "text"), ("total_debt", "الدين", "money"), ("total_paid", "المسدّد", "money"), ("remaining", "المتبقي", "money")], data.get("suppliers") or [])
        if table:
            widgets.append(table)
    elif name == "get_purchases_summary":
        widgets.append(_metric_widget("المشتريات", [
            _metric("عدد الفواتير", data.get("count"), tone="blue"),
            _metric("الإجمالي", data.get("grand_total"), kind="money"),
            _metric("المدفوع", data.get("paid_total"), kind="money", tone="green"),
            _metric("المتبقي", data.get("remaining_total"), kind="money", tone="amber"),
        ]))
    elif name == "get_customers_receivables":
        widgets.append(_metric_widget("ذمم الزبائن", [
            _metric("إجمالي الذمم", data.get("total_receivables"), kind="money", tone="amber"),
            _metric("زبائن عليهم دين", data.get("customers_with_debt"), tone="blue"),
            _metric("أقساط متأخرة", (data.get("overdue_installments") or {}).get("amount"), kind="money", tone="red"),
        ]))
        bar = _bar_widget("أعلى الذمم", data.get("top_debtors") or [], value_key="remaining")
        if bar:
            widgets.append(bar)
    elif name in {"get_low_stock", "search_products", "get_product_stock", "get_top_products"}:
        if name == "get_low_stock":
            widgets.append(_metric_widget("تنبيه المخزون", [_metric("أصناف منخفضة", data.get("count"), tone="red")]))
            table = _table_widget("المخزون المنخفض", [("product", "المنتج", "text"), ("branch", "الفرع", "text"), ("qty", "المتوفر", "number"), ("threshold", "حد التنبيه", "number")], data.get("low_stock") or [])
        elif name == "get_top_products":
            table = _table_widget("أفضل المنتجات", [("product", "المنتج", "text"), ("quantity_sold", "الكمية", "number"), ("revenue", "الإيراد", "money"), ("profit", "الربح", "money")], data.get("top_products") or [])
        elif name == "search_products":
            table = _table_widget("نتائج المنتجات", [("name", "المنتج", "text"), ("total_quantity", "المخزون", "number"), ("sale_price", "سعر البيع", "money")], data.get("products") or [])
        else:
            product = data.get("product") or {}
            widgets.append(_metric_widget(str(product.get("name") or "مخزون المنتج"), [
                _metric("الإجمالي", product.get("total_quantity"), tone="blue"),
                _metric("سعر البيع", product.get("sale_price"), kind="money"),
            ]))
            table = _table_widget("المخزون حسب الفرع", [("branch", "الفرع", "text"), ("system_qty", "النظام", "number"), ("reserved_ordered_or_shipping", "المحجوز", "number"), ("salable_qty", "القابل للبيع", "number")], data.get("branch_stock") or [])
        if table:
            widgets.append(table)
    elif name == "get_employees_performance":
        rows = data.get("performance") or []
        widgets.append(_metric_widget("أداء الفريق", [
            _metric("الموظفون", data.get("employees_count", len(rows)), tone="blue"),
            _metric("إجمالي الطلبات", sum(_number(row.get("orders")) for row in rows), tone="purple"),
            _metric("إجمالي المبيعات", sum(_number(row.get("total_sales")) for row in rows), kind="money", tone="green"),
        ], subtitle=str(data.get("period") or "")))
        table = _table_widget("ترتيب الموظفين", [("name", "الموظف", "text"), ("orders", "الطلبات", "number"), ("total_sales", "المبيعات", "money"), ("collected", "المحصّل", "money")], rows)
        if table:
            widgets.append(table)
    elif name == "get_financial_overview":
        treasury = data.get("treasury") or {}
        widgets.append(_metric_widget("النظرة المالية الشاملة", [
            _metric("إجمالي السيولة", treasury.get("total_liquidity"), kind="money", tone="green"),
            _metric("قيمة المخزون", data.get("inventory_value"), kind="money"),
            _metric("ذمم الزبائن", data.get("accounts_receivable"), kind="money", tone="blue"),
            _metric("ديون الموردين", data.get("supplier_debts"), kind="money", tone="amber"),
            _metric("ذمم شركات النقل", data.get("shipping_receivables"), kind="money", tone="purple"),
        ]))
        if isinstance(data.get("period_profit"), dict):
            widgets.extend(_widgets_for_tool("get_profit_summary", data["period_profit"]))
    elif name == "search_orders":
        table = _table_widget("الطلبات", [("invoice_id", "الطلب", "number"), ("customer", "الزبون", "text"), ("status", "الحالة", "text"), ("total", "الإجمالي", "money")], data.get("orders") or [])
        if table:
            widgets.append(table)
    elif name == "get_order_details":
        widgets.append(_metric_widget(f"الطلب #{data.get('invoice_id', '')}", [
            _metric("الإجمالي", data.get("total"), kind="money"),
            _metric("المسدّد", data.get("effective_paid"), kind="money", tone="green"),
            _metric("الحالة", data.get("status"), kind="text", tone="blue"),
        ], subtitle=str(data.get("customer") or "")))
        table = _table_widget("أصناف الطلب", [("name", "المنتج", "text"), ("quantity", "الكمية", "number"), ("price", "السعر", "money"), ("line_total", "الإجمالي", "money")], data.get("items") or [])
        if table:
            widgets.append(table)
    elif name == "run_accounting_audit":
        summary = data.get("summary") or {}
        widgets.append(_metric_widget("التدقيق المحاسبي", [
            _metric("فروقات المخزون", summary.get("stock_imbalances_count"), tone="red"),
            _metric("فواتير غير متطابقة", summary.get("invoice_total_mismatches_count"), tone="amber"),
            _metric("هوامش سالبة", summary.get("negative_margin_items_count"), tone="red"),
            _metric("حالات متناقضة", summary.get("status_inconsistencies_count"), tone="amber"),
        ]))
    return widgets


def _overview_widget(snapshot: dict) -> dict | None:
    items: list[dict] = []
    if snapshot.get("recent_30d_sales") is not None:
        items.append(_metric("مبيعات 30 يوم", snapshot.get("recent_30d_sales"), kind="money", tone="green"))
    products = snapshot.get("products") or {}
    if isinstance(products, dict) and not products.get("restricted"):
        items.append(_metric("المنتجات النشطة", products.get("active"), tone="blue"))
        if products.get("stock_value") is not None:
            items.append(_metric("قيمة المخزون", products.get("stock_value"), kind="money"))
    financial = snapshot.get("financial") or {}
    if isinstance(financial, dict) and not financial.get("restricted"):
        cash = financial.get("cash") or {}
        if isinstance(cash, dict) and cash.get("balance") is not None:
            items.append(_metric("رصيد الصندوق", cash.get("balance"), kind="money", tone="purple"))
        if financial.get("supplier_debts") is not None:
            items.append(_metric("ديون الموردين", financial.get("supplier_debts"), kind="money", tone="amber"))
    return _metric_widget("نظرة سريعة من النظام", items) if items else None


def build_leon_presentation(*, session_id: int, message: str, reply: str, tool_runs: list[dict] | None, local_findings: dict, action_plan: dict | None = None) -> dict:
    widgets: list[dict] = []
    used: set[str] = set()
    for run in tool_runs or []:
        name = str(run.get("name") or "")
        if not name or name in used:
            continue
        used.add(name)
        widgets.extend(_widgets_for_tool(name, run.get("result") or {}))
        if len(widgets) >= 5:
            break
    if not widgets:
        overview = _overview_widget(local_findings.get("snapshot") or {})
        if overview:
            widgets.append(overview)
    if action_plan:
        widgets.append({
            "type": "action_plan",
            "title": action_plan.get("title") or "خطة تحتاج موافقة",
            "message": action_plan.get("summary") or "راجع الخطة قبل التنفيذ.",
            "risk": action_plan.get("risk_level") or "medium",
            "plan_id": action_plan.get("id"),
            "status": action_plan.get("status") or "draft",
        })

    return {
        "version": 1,
        "title": "نتيجة LEON",
        "query": str(message or "")[:500],
        "summary": str(reply or "")[:1200],
        "widgets": widgets[:6],
        "files": [
            {"format": "xlsx", "label": "ملف Excel", "url": f"/assistant/api/chat/sessions/{session_id}/export.xlsx"},
            {"format": "pdf", "label": "ملخص PDF", "url": f"/assistant/api/chat/sessions/{session_id}/export.pdf"},
            {"format": "csv", "label": "بيانات CSV", "url": f"/assistant/api/chat/sessions/{session_id}/export.csv"},
        ] if widgets else [],
        "suggestions": ["قارنها بالفترة السابقة", "اعرض التفاصيل", "ما أهم تنبيه؟"],
        "generated_at": datetime.utcnow().isoformat(),
    }


def _flatten_presentation(presentation: dict) -> list[list[Any]]:
    rows: list[list[Any]] = [["LEON", presentation.get("title") or "نتيجة LEON"], ["السؤال", presentation.get("query") or ""], ["الملخص", presentation.get("summary") or ""], []]
    for widget in presentation.get("widgets") or []:
        rows.append([widget.get("title") or "قسم"])
        if widget.get("type") == "metric_grid":
            rows.append(["المؤشر", "القيمة", "النوع"])
            rows.extend([[item.get("label"), item.get("value"), item.get("kind")] for item in widget.get("items") or []])
        elif widget.get("type") == "bar_list":
            rows.append(["البند", "القيمة"])
            rows.extend([[item.get("label"), item.get("value")] for item in widget.get("items") or []])
        elif widget.get("type") == "table":
            columns = widget.get("columns") or []
            rows.append([column.get("label") for column in columns])
            rows.extend([[item.get(column.get("key"), "") for column in columns] for item in widget.get("rows") or []])
        elif widget.get("message"):
            rows.append([widget.get("message")])
        rows.append([])
    return rows


def export_presentation(presentation: dict, file_format: str) -> tuple[bytes, str, str]:
    file_format = str(file_format or "xlsx").lower()
    rows = _flatten_presentation(presentation)
    stamp = datetime.utcnow().strftime("%Y%m%d-%H%M")
    if file_format == "json":
        return json.dumps(presentation, ensure_ascii=False, indent=2).encode("utf-8"), "application/json", f"leon-summary-{stamp}.json"
    if file_format == "csv":
        text_buffer = io.StringIO(newline="")
        writer = csv.writer(text_buffer)
        writer.writerows(rows)
        return text_buffer.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", f"leon-summary-{stamp}.csv"
    if file_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ملخص LEON"
        sheet.sheet_view.rightToLeft = True
        for row_index, row in enumerate(rows, start=1):
            sheet.append(row)
            if len(row) == 1 and row[0]:
                sheet.cell(row_index, 1).font = Font(bold=True, color="FFFFFF")
                sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor="6F5DF7")
            for cell in sheet[row_index]:
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        for column in "ABCDEFGH":
            sheet.column_dimensions[column].width = 22
        sheet.column_dimensions["A"].width = 30
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"leon-summary-{stamp}.xlsx"
    if file_format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        font_name = "Helvetica"
        for font_path in (r"C:\Windows\Fonts\arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
            if os.path.exists(font_path):
                font_name = "LeonArabic"
                if font_name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(font_name, font_path))
                break
        output = io.BytesIO()
        document = SimpleDocTemplate(output, pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title="LEON Summary")
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("LeonTitle", parent=styles["Title"], fontName=font_name, fontSize=18, leading=24, alignment=2, textColor=colors.HexColor("#2B254F"))
        body_style = ParagraphStyle("LeonBody", parent=styles["BodyText"], fontName=font_name, fontSize=9, leading=15, alignment=2)
        story = [Paragraph("LEON — FINORA AI", title_style), Spacer(1, 5 * mm), Paragraph(escape(str(presentation.get("summary") or "")), body_style), Spacer(1, 5 * mm)]
        for widget in presentation.get("widgets") or []:
            story.append(Paragraph(escape(str(widget.get("title") or "")), title_style))
            flat = _flatten_presentation({"widgets": [widget]})[1:]
            clean = [[str(cell if cell is not None else "") for cell in row] for row in flat if row]
            if clean:
                table = Table(clean, repeatRows=1, hAlign="RIGHT")
                table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font_name), ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#DDD9F6")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0EEFF")), ("ALIGN", (0, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
                story.extend([table, Spacer(1, 4 * mm)])
        document.build(story)
        return output.getvalue(), "application/pdf", f"leon-summary-{stamp}.pdf"
    raise ValueError("صيغة الملف غير مدعومة")
