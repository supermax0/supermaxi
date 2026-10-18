"""Permission-scoped HR exports using the same data and approval snapshots as the screen."""
from io import BytesIO
from datetime import datetime
from flask import abort, render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .common import BAGHDAD, audit, can, integer, own, person, require
from .employees import list_people

REPORTS = {
    "employees": ("سجل الموظفين", [("code", "الرقم الوظيفي"), ("name", "الموظف"), ("department", "القسم"), ("job_title", "الوظيفة"), ("branch", "الفرع"), ("phone", "الهاتف"), ("joined_on", "تاريخ الانضمام"), ("status", "الحالة")]),
    "attendance": ("الحضور والدوام", [("code", "الرقم الوظيفي"), ("person_name", "الموظف"), ("state", "الحالة"), ("check_in", "الحضور"), ("check_out", "الانصراف"), ("worked_minutes", "دقائق العمل"), ("late_minutes", "التأخير"), ("overtime_minutes", "الإضافي")]),
    "payroll": ("كشف الرواتب", [("code", "الرقم الوظيفي"), ("person_name", "الموظف"), ("status", "حالة المسير"), ("gross", "الإجمالي"), ("deductions", "الاستقطاعات"), ("net", "الصافي"), ("paid", "المصروف"), ("remaining", "المتبقي")]),
    "advances": ("السلف وذمم الموظفين", [("code", "الرقم الوظيفي"), ("person_name", "الموظف"), ("amount", "السلفة"), ("repaid", "المسدد"), ("remaining", "المتبقي"), ("instalments", "الأقساط"), ("first_month", "بداية الأقساط"), ("status", "الحالة")]),
    "commissions": ("كشف العمولات والتسويات", [("code", "الرقم الوظيفي"), ("person_name", "الموظف"), ("invoice_id", "الفاتورة"), ("amount", "العمولة أو التسوية"), ("paid", "المصروف"), ("remaining", "المتبقي"), ("status", "الحالة")]),
    "leave": ("طلبات الإجازات", [("code", "الرقم الوظيفي"), ("person_name", "الموظف"), ("type", "نوع الإجازة"), ("starts", "من"), ("ends", "إلى"), ("days", "أيام العمل"), ("status", "الحالة")]),
    "resources": ("السجلات الوظيفية", [("code", "الرقم الوظيفي"), ("person_name", "الموظف"), ("kind", "نوع السجل"), ("title", "التفاصيل"), ("date", "التاريخ")]),
    "settlement": ("كشف نهاية الخدمة", [("label", "البيان"), ("value", "القيمة")]),
    "statement": ("كشف الحساب المالي للموظف", [("date", "التاريخ"), ("label", "العملية"), ("due", "المستحق"), ("paid", "المدفوع"), ("adjustment", "التسويات والاستقطاعات"), ("balance", "الرصيد"), ("source", "المستند")]),
}
LABELS = {"active":"نشط", "suspended":"موقوف", "terminated":"منتهي الخدمة", "pending":"بانتظار الاعتماد",
    "approved":"معتمد", "reviewed":"مراجع", "draft":"مسودة", "paid":"مصروف", "repaid":"مسدد", "repaying":"قيد التسديد",
    "void":"ملغى", "adjustment":"تسوية تحتاج مراجعة", "present":"حاضر", "late":"متأخر", "absent":"غائب", "leave":"إجازة",
    "scheduled":"لم يبدأ الدوام", "unscheduled":"بدون وردية", "rejected":"مرفوض", "review":"تقييم", "goal":"هدف", "warning":"إجراء إداري", "note":"ملاحظة", "contract_revision":"تعديل عقد","cancelled":"ملغى","holiday":"عطلة","monthly":"شهري","monthly_commission":"شهري وعمولات","weekly":"أسبوعي","daily":"يومي","hourly":"بالساعة","commission":"عمولات"}


def payroll_reconciliation(rows):
    from extensions import db
    from .models import HRPayrollItem, HRPayrollPayment
    from models.expense import Expense
    from models.employee_payment import EmployeePayment
    from models.account_transaction import AccountTransaction
    from models.journal_entry import JournalEntry
    ids=[r["id"] for r in rows if r["status"]=="approved"]
    if not ids:
        return []
    expenses={item.id:int(amount or 0) for item,amount in db.session.query(HRPayrollItem,Expense.amount).outerjoin(Expense,Expense.id==HRPayrollItem.expense_id).filter(HRPayrollItem.id.in_(ids)).all()}
    sums={kind:{} for kind in ("hr_payroll_accrual","hr_wage_deduction")}
    for reference,kind,amount in db.session.query(JournalEntry.reference_id,JournalEntry.reference_type,db.func.sum(JournalEntry.amount)).filter(JournalEntry.reference_id.in_(ids),JournalEntry.reference_type.in_(list(sums))).group_by(JournalEntry.reference_id,JournalEntry.reference_type).all():
        sums[kind][reference]=int(amount or 0)
    cash={};ledger={}
    for item_id,paid,cash_amount in db.session.query(HRPayrollPayment.payroll_item_id,db.func.sum(EmployeePayment.amount),db.func.sum(AccountTransaction.amount)).join(EmployeePayment,EmployeePayment.id==HRPayrollPayment.payment_id).join(AccountTransaction,AccountTransaction.id==HRPayrollPayment.transaction_id).filter(HRPayrollPayment.payroll_item_id.in_(ids)).group_by(HRPayrollPayment.payroll_item_id).all():
        cash[item_id]=(int(paid or 0),int(cash_amount or 0))
    for item_id,amount in db.session.query(HRPayrollPayment.payroll_item_id,db.func.sum(JournalEntry.amount)).join(JournalEntry,db.and_(JournalEntry.reference_id==HRPayrollPayment.payment_id,JournalEntry.reference_type=="hr_payroll_payment")).filter(HRPayrollPayment.payroll_item_id.in_(ids)).group_by(HRPayrollPayment.payroll_item_id).all():
        ledger[item_id]=int(amount or 0)
    result=[]
    for row in rows:
        if row["id"] not in expenses:continue
        id_=row["id"];expected=row["gross"]-row.get("calculation",{}).get("wage_deductions",0)
        gl=sums["hr_payroll_accrual"].get(id_,0)-sums["hr_wage_deduction"].get(id_,0)
        native,withdrawn=cash.get(id_,(0,0));payments=ledger.get(id_,0)
        differences=[expected-expenses[id_],expected-gl,row["paid"]-native,row["paid"]-withdrawn,row["paid"]-payments]
        result.append([row["code"],row["person_name"],expected,expenses[id_],gl,row["paid"],native,withdrawn,payments,"مطابق" if not any(differences) else "فرق يحتاج مراجعة: "+str(differences)])
    return result


def export_report(kind, format_, args):
    if kind not in REPORTS or format_ not in ("xlsx", "print"):
        abort(404)
    if not can("hr.reports.view"):
        # Staff can obtain their own payslip without access to other HR reports.
        if kind not in ("payroll", "settlement", "statement") or not args.get("person_id") or not can("hr.self") or not own(person(args["person_id"])):
            require("hr.reports.view")
    from .routes import workspace_data
    query = {k: v for k, v in args.items() if k in ("search", "status", "branch_id", "department_id", "person_id", "month", "date", "sort", "date_from", "date_to")}
    query.update(tab=kind, page_size=100)
    rows = []
    statement_data=None
    if kind == "statement":
        from .statements import statement
        statement_data=statement(query,paginate=False)
        rows=[{**r,"date":datetime.fromisoformat(r["date"].replace("Z","+00:00")).astimezone(BAGHDAD).strftime("%Y-%m-%d %H:%M"),"source":r["document"]["number"]} for r in statement_data["rows"]]
    elif kind == "settlement":
        from .employees import details
        employee=person(integer(args.get("person_id"),minimum=1))
        if not own(employee):
            for permission in ("hr.payroll.view","hr.commission.view","hr.advance.view"):
                require(permission)
        detail=details(employee.id)
        exit_row=detail["offboarding"]
        if not exit_row:abort(404)
        frozen=exit_row["settlement"]
        earned=frozen.get("earned_payroll",{})
        entries=[("الموظف",employee.name),("الرقم الوظيفي",employee.code),("آخر يوم عمل",exit_row["ends_on"]),("سبب إنهاء الخدمة",exit_row["reason"]),
            ("رواتب سابقة معتمدة لم تُصرف",frozen.get("approved_salary_due",0)),("الأجر المحتسب حتى آخر يوم عمل",earned.get("net",0)),
            ("العمولات المعتمدة المتبقية",frozen.get("commission_due",0)),("عمولات تحتاج مراجعة",frozen.get("commission_pending",0)),
            ("سلف متبقية",frozen.get("advance_remaining",0)),("اتفاق سداد السلف","موثّق" if frozen.get("advance_agreement") else "—"),
            ("ملاحظات التسوية",frozen.get("note","")),("الصرف","الكشف لا ينشئ دفعة. يُعتمد الأجر عند إقفال الفترة ويُصرف من الخزينة؛ رصيد الإجازات يحتاج بدل تسوية بحسب السياسة.")]
        entries.extend(("رصيد إجازة: "+leave["type"],leave["days"]) for leave in frozen.get("leave_balances",[]))
        rows=[{"label":label,"value":value} for label,value in entries]
        query["date"]=exit_row["ends_on"]
    else:
        for page in range(1, 1001):
            query["page"] = page
            part = list_people(query) if kind == "employees" else workspace_data(query)
            rows.extend(part["rows"])
            if page >= part["total_pages"]:
                break
        else:
            raise ValueError("التقرير يتجاوز الحد؛ حدد فرعاً أو قسماً لتقليص النطاق")
    title, columns = REPORTS[kind]
    breakdown=[]
    if statement_data:
        title+=" · "+statement_data["name"]+" · "+statement_data["code"]
        breakdown.append({"title":"ملخص الفترة","columns":["رصيد أول الفترة","المستحق","المدفوع","التسويات","رصيد نهاية الفترة"],"rows":[[statement_data[k] for k in ("opening","due","paid","adjustment","balance")]]})
        breakdown.append({"title":"ذمم السلف الحالية (منفصلة عن الأجور)","columns":["السلفة","المبلغ","المسدد","المتبقي","السبب"],"rows":[[r[k] for k in ("id","amount","repaid","remaining","reason")] for r in statement_data["advances"]]})
    if kind=="payroll" and not args.get("person_id"):
        breakdown.append({"title":"مطابقة الرواتب مع المصروف والقيود والخزينة","columns":["الرقم","الموظف","المصروف المحتسب","المصروف المسجل","قيد المصروف الصافي","المصروف للموظف","سجل الدفعات","المنصرف من الخزينة","قيد الصرف","نتيجة المطابقة"],"rows":payroll_reconciliation(rows)})
    if kind=="payroll" and args.get("person_id"):
        title="قسيمة راتب الموظف"
        for row in rows:
            calc=row.get("calculation",{})
            breakdown.append({"title":row["person_name"]+" · "+row["code"],"columns":["التاريخ","نوع الأجر","الأجر","دقائق العمل","المبلغ المحتسب","بدل إضافي","حسم دوام"],
                "rows":[[LABELS.get(str(d.get(key,"")),d.get(key,"")) if key=="type" else d.get(key,"") for key in ("date","type","rate","minutes","amount","overtime_bonus","attendance_penalty")] for d in calc.get("days",[])]})
            entries=[[a.get("title",""),"بدل" if a.get("kind")=="allowance" else "استقطاع",a.get("amount",0)] for a in calc.get("adjustments",[])]
            entries.extend([["قسط سلفة #"+str(a["advance_id"]),"سداد ذمة",a["amount"]] for a in calc.get("advance_repayments",[])])
            if calc.get("legacy_paid"):entries.append(["راتب سبق صرفه","من السجل السابق",calc["legacy_paid"]])
            breakdown.append({"title":"تفصيل البدلات والاستقطاعات والسلف","columns":["البيان","النوع","المبلغ"],"rows":entries})
    data = [[LABELS.get(str(r.get(key, "")), r.get(key, "")) if key in ("status", "state", "kind") else
             r.get("detail", {}).get("title", "") if key == "title" else r.get(key, "") for key, _ in columns] for r in rows]
    generated = datetime.now(BAGHDAD).strftime("%Y-%m-%d %H:%M")
    audit("report.export", kind, detail={"format": format_, "rows": len(rows), "filters": query})
    from extensions import db
    db.session.commit()
    if format_ == "print":
        return render_template("hr/report.html", title=title, columns=columns, rows=data, generated=generated,
            period=query.get("month") or query.get("date") or "", currency="الدينار العراقي",breakdown=breakdown)
    book = Workbook()
    sheet = book.active
    sheet.title = "FINORA HR"
    sheet.sheet_view.rightToLeft = True
    sheet.append([title, generated, query.get("month") or query.get("date") or ""])
    sheet.append([label for _, label in columns])
    for values in data:
        sheet.append(["'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v for v in values])
    for cell in sheet[2]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="801A43")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
        for cell in column:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.##"
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(max(2, sheet.max_row), len(columns)).coordinate}"
    for index,part in enumerate(breakdown):
        detail=book.create_sheet("التفصيل "+str(index+1))
        detail.sheet_view.rightToLeft=True
        detail.append([part["title"]]);detail.append(part["columns"])
        for row in part["rows"]:
            detail.append(["'"+v if isinstance(v,str) and v.startswith(("=","+","-","@")) else v for v in row])
        for cells in detail.columns:
            detail.column_dimensions[cells[0].column_letter].width=24
        detail.freeze_panes="A3"
    output = BytesIO()
    book.save(output)
    output.seek(0)
    response = send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True, download_name=f"FINORA-HR-{kind}.xlsx")
    response.headers["Cache-Control"] = "private, no-store"
    return response
