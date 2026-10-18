"""One saved-result report for print/PDF and XLSX, with identical branch/cost permissions."""
from io import BytesIO
import json
from datetime import datetime
from flask import render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .common import AuditError, can, require
from .service import history, line_data, lines, session_data
from .models import InventoryCountEvent

STATES = {"counting":"قيد الجرد", "paused":"متوقف مؤقتاً", "pending_review":"بانتظار المراجعة", "approved":"معتمد", "posted":"مرحّل"}


def data(row):
    require("inventory.audit.export")
    if lines(row).count() > 20000:
        raise AuditError("التقرير يتجاوز 20 ألف صنف؛ قسم الجرد إلى نطاقات أصغر")
    basis = None
    if can("inventory.audit.cost") and can("inventory.audit.system"):
        posted = InventoryCountEvent.query.filter_by(session_id=row.id, action="post").first()
        if posted:
            basis = json.loads(posted.detail_json).get("valuation_reconciliation")
    return {"audit":session_data(row), "rows":[line_data(line, inspect_current=False) for line in lines(row)],
            "valuation_basis": basis,
            "show_system":can("inventory.audit.system"), "show_cost":can("inventory.audit.cost"),
            "state":STATES[row.status], "history":history(row, page_size=100)["items"]}


def print_page(row):
    return render_template("inventory_smart_audit_print.html", format_time=format_time, **data(row))


def format_time(value):
    if not value:
        return "لم يجرد"
    return datetime.fromisoformat(value.replace("Z", "+00:00")).strftime("%Y/%m/%d %H:%M:%S") + " UTC"


def spreadsheet(row):
    content = data(row)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "جرد المخزون"
    sheet.sheet_view.rightToLeft = True
    headers = ["الجلسة", "الحالة", "الموقع", "المنتج", "SKU", "المتغير", "وقت العد", "العدد الفعلي"]
    if content["show_system"]:
        headers += ["رصيد النظام عند العد", "الفرق", "التسوية المرحّلة"]
    headers += ["نتيجة المراجعة", "السبب", "ملاحظة", "مستبعد"]
    if content["show_cost"]:
        headers += ["تكلفة التسوية د.ع"]
    sheet.append(headers)
    locations = {loc["key"]:f'{loc["branch_name"]} / {loc["name"]}' for loc in content["audit"]["locations"]}
    for line in content["rows"]:
        values = [row.number, content["state"], locations[line["location_key"]], line["product"]["name"],
                  line["product"]["sku"], line["product"]["variant_name"], format_time(line["counted_at"]), line["counted_quantity"]]
        if content["show_system"]:
            values += [line["expected_quantity"], line["difference"], line["posted_delta"]]
        values += [line["resolution"], line["resolution_note"], line["notes"], "نعم" if line["excluded"] else "لا"]
        if content["show_cost"]:
            values += [line["posted_cost"]]
        # Imported names/notes are literal text, never spreadsheet formulas.
        sheet.append([("'" + value if value.startswith(("=", "+", "-", "@")) else value) if isinstance(value, str) else value for value in values])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font, cell.fill = Font(color="FFFFFF", bold=True), PatternFill("solid", fgColor="7138DB")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(50, max(18, max(len(str(cell.value or "")) for cell in column) + 2))
        for cell in column:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    metadata = workbook.create_sheet("ملخص النطاق")
    metadata.sheet_view.rightToLeft = True
    metadata.append(["جميع الأرقام تخص نطاق هذه الجلسة؛ غير المجرد لا يحتسب صفراً"])
    if content["valuation_basis"]:
        basis = content["valuation_basis"]
        repair = workbook.create_sheet("مطابقة طبقات التكلفة")
        repair.sheet_view.rightToLeft = True
        repair.append(["مطابقة أساس FIFO؛ لا يكرر قيد مبيعات سابقة، وقيد الجرد يخص فرق العد الفعلي فقط"])
        basis_reason = basis["reason"]
        repair.append(["السبب", "'" + basis_reason if basis_reason.startswith(("=", "+", "-", "@")) else basis_reason])
        repair.append(["سطر", "الموقع", "المنتج", "المتغير", "رصيد الفرع قبل الجرد", "طبقات قبل المطابقة", "قيمة قبل", "قيمة بعد", "زائد", "غير مغطى"])
        for item in basis["items"]:
            values = [item[key] for key in ("line_id", "location_key", "product_name", "variant_name", "system_quantity",
                "layer_quantity_before", "layer_value_before", "layer_value_after", "removed_quantity", "added_quantity")]
            repair.append([("'" + value if value.startswith(("=", "+", "-", "@")) else value) if isinstance(value, str) else value for value in values])
    for key, label in [("planned","مخططة"),("counted","مجردة"),("uncounted","غير مجرودة"),("excluded","مستبعدة"),
                       ("shortage_units","إجمالي النقص"),("surplus_units","إجمالي الزيادة")]:
        metadata.append([label, content["audit"]["stats"].get(key)])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=f"{row.number}.xlsx")
