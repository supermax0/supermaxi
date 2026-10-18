"""Catalogue import (Excel / CSV / supplier lists) and smart lookup — preview first, nothing is published.

Flow: upload → parse → analyse every row against the catalogue AND the inventory (barcode, then brand+model) →
a preview with a status per row → the employee applies the batch. Applying creates entries in «needs review»
(never published), only fills EMPTY fields of entries that already exist (a verified value is never overwritten),
links rows that match an inventory product, and can run once per batch.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import datetime
from typing import Any

from extensions import db
from modules.catalog import service
from modules.catalog.models import CatalogBarcode, CatalogImportBatch, CatalogImportRow, CatalogProduct

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

# header aliases (Arabic / English, as suppliers write them). Unknown headers become specification attributes.
ALIASES = {
    "name": ("name", "product", "product name", "title", "الاسم", "اسم المنتج", "المنتج", "الوصف المختصر"),
    "brand": ("brand", "make", "manufacturer", "العلامة", "الماركة", "الشركة", "الشركة المصنعة", "البراند"),
    "model_number": ("model", "model number", "model no", "mpn", "الموديل", "رقم الموديل", "موديل"),
    "barcode": ("barcode", "ean", "gtin", "upc", "الباركود", "باركود", "الرمز الشريطي"),
    "category": ("category", "القسم", "الفئة", "الصنف", "التصنيف"),
    "description": ("description", "details", "الوصف", "التفاصيل"),
    "image_url": ("image", "image url", "photo", "picture", "الصورة", "رابط الصورة"),
    "regional_version": ("region", "version", "النسخة", "المنشأ", "الإصدار"),
}
IGNORED = ("price", "cost", "qty", "quantity", "stock", "السعر", "الكلفة", "الكمية", "المخزون", "سعر الشراء", "سعر البيع")


def _canon(header: str) -> str | None:
    key = " ".join(str(header or "").strip().casefold().split())
    for field, names in ALIASES.items():
        if key in names:
            return field
    if key in IGNORED:
        return "__ignore__"  # prices and quantities never come in through the catalogue (Sprint 5 / Inventory)
    return None


def parse(raw: bytes, filename: str) -> list[dict[str, Any]]:
    if not raw:
        raise service.CatalogError("الملف فارغ")
    if len(raw) > MAX_BYTES:
        raise service.CatalogError("حجم الملف أكبر من 5 ميغابايت")
    name = str(filename or "").lower()
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        table = list(csv.reader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise service.CatalogError("تعذّر قراءة ملف Excel") from None
        sheet = book.worksheets[0]
        table = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    else:
        raise service.CatalogError("الملفات المدعومة: Excel (.xlsx) أو CSV")
    # keep the real spreadsheet row number of every row (blank rows are dropped, not renumbered)
    table = [(n, r) for n, r in enumerate(table, start=1) if any(str(c).strip() for c in r)]
    if len(table) < 2:
        raise service.CatalogError("لا توجد صفوف بيانات تحت سطر العناوين")
    headers = [str(h).strip() for h in table[0][1]]
    fields = [_canon(h) for h in headers]
    if "name" not in fields and "model_number" not in fields:
        raise service.CatalogError("يجب أن يحتوي الملف عمود «الاسم» أو «الموديل»")
    if len(table) - 1 > MAX_ROWS:
        raise service.CatalogError(f"الحد الأقصى {MAX_ROWS} صف في الملف الواحد")
    rows = []
    for row_no, values in table[1:]:
        row: dict[str, Any] = {"attributes": {}, "_row_no": row_no}
        for header, field, value in zip(headers, fields, values):
            value = str(value or "").strip()
            if not value or field == "__ignore__":
                continue
            if field:
                row[field] = value
            else:
                row["attributes"][header[:60]] = value[:200]
        rows.append(row)
    return rows


def _inventory_index() -> tuple[dict[str, int], dict[str, int]]:
    """{barcode: product id}, {normalised sku/model: product id} over active products — built once per analysis."""
    from models.product import Product

    rows = (Product.query.execution_options(skip_branch_scope=True).filter(Product.active.is_(True))
            .with_entities(Product.id, Product.barcode, Product.sku).all())
    by_code = {str(code): int(pid) for pid, code, _ in rows if code}
    by_model = {service.norm_key(sku): int(pid) for pid, _, sku in rows if sku}
    return by_code, by_model


def _inventory_match(barcode: str, model_number: str, index=None) -> int | None:
    by_code, by_model = index or _inventory_index()
    if barcode and barcode in by_code:
        return by_code[barcode]
    key = service.norm_key(model_number)
    return by_model.get(key) if key else None


def analyse(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    index = _inventory_index()
    seen_codes: set[str] = set()
    seen_models: set[tuple[str, str]] = set()
    out = []
    for i, row in enumerate(rows, start=2):
        i = int(row.pop("_row_no", i))  # the row number the employee sees in his spreadsheet
        issues: list[str] = []
        code = service.norm_barcode(row.get("barcode"))
        if row.get("barcode") and len(code) < 8:
            issues.append("باركود غير صالح")
            code = ""
        brand, model = str(row.get("brand") or ""), str(row.get("model_number") or "")
        key = (service.norm_key(brand), service.norm_key(model))
        name = str(row.get("name") or "").strip()
        if len(name) < 3:
            name = " ".join(x for x in (brand, model, name) if x).strip()
        if len(name) < 3:
            issues.append("لا يوجد اسم")
        status, entry_id, product_id = "new", None, None
        if (code and code in seen_codes) or (all(key) and key in seen_models):
            status = "duplicate"
            issues.append("مكرر داخل نفس الملف")
        else:
            if code:
                seen_codes.add(code)
            if all(key):
                seen_models.add(key)
            existing = service.find_duplicate(brand=brand, model_number=model,
                                              regional_version=str(row.get("regional_version") or ""),
                                              barcodes=[code] if code else [])
            if existing is not None:
                status, entry_id = "matched", existing.id
            product_id = _inventory_match(code, model, index)
            if product_id is not None:
                if status == "new":
                    status = "linked"
            if status == "new" and not (brand and model):
                status = "needs_review"
                issues.append("العلامة أو الموديل ناقص — يحتاج مراجعة قبل الحفظ")
        if issues and "لا يوجد اسم" in issues:
            status = "invalid"
        out.append({"row_no": i, "data": {**row, "name": name, "barcode": code}, "status": status,
                    "issues": issues, "match_entry_id": entry_id, "match_product_id": product_id})
    return out


def create_batch(raw: bytes, filename: str, *, by: int | None) -> CatalogImportBatch:
    service.ensure_tables()
    rows = analyse(parse(raw, filename))
    digest = hashlib.sha256(raw).hexdigest()
    applied = CatalogImportBatch.query.filter_by(file_sha256=digest, status="applied").first()
    if applied is not None:
        raise service.CatalogError("هذا الملف نفسه استُورد سابقاً", "duplicate_file", 409, batch_id=applied.id)
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    batch = CatalogImportBatch(filename=str(filename)[:200], file_sha256=digest, counts_json=json.dumps(counts),
                               created_by_id=by)
    db.session.add(batch)
    db.session.flush()
    for r in rows:
        db.session.add(CatalogImportRow(batch_id=batch.id, row_no=r["row_no"],
                                        data_json=json.dumps(r["data"], ensure_ascii=False), status=r["status"],
                                        issues_json=json.dumps(r["issues"], ensure_ascii=False),
                                        match_entry_id=r["match_entry_id"], match_product_id=r["match_product_id"]))
    db.session.flush()
    return batch


def batch_view(batch: CatalogImportBatch, *, limit: int = 200) -> dict[str, Any]:
    rows = (CatalogImportRow.query.filter_by(batch_id=batch.id).order_by(CatalogImportRow.row_no).limit(limit).all())
    return {"id": batch.id, "filename": batch.filename, "status": batch.status,
            "counts": json.loads(batch.counts_json or "{}"),
            "rows": [{"id": r.id, "row_no": r.row_no, "status": r.status, "data": json.loads(r.data_json),
                      "issues": json.loads(r.issues_json or "[]"), "match_entry_id": r.match_entry_id,
                      "match_product_id": r.match_product_id, "applied_entry_id": r.applied_entry_id} for r in rows]}


def _sources(filename: str, data: dict) -> dict[str, dict]:
    tag = f"import:{filename}"
    return {f: {"source": tag, "confidence": 60} for f in ("name", "brand", "model_number", "barcode", "category",
                                                         "description") if data.get(f)}


def apply(batch_id: int, *, by: int | None, include_review: bool = False) -> dict[str, int]:
    service.ensure_tables()
    batch = db.session.get(CatalogImportBatch, int(batch_id))
    if batch is None:
        raise service.CatalogError("دفعة الاستيراد غير موجودة", "not_found", 404)
    if batch.status != "preview":
        raise service.CatalogError("نُفّذت هذه الدفعة مسبقاً", "conflict", 409)
    result = {"created": 0, "updated": 0, "linked": 0, "skipped": 0}
    wanted = {"new", "matched", "linked"} | ({"needs_review"} if include_review else set())
    for row in CatalogImportRow.query.filter_by(batch_id=batch.id).order_by(CatalogImportRow.row_no).all():
        if row.status not in wanted:
            result["skipped"] += 1
            continue
        data = json.loads(row.data_json)
        entry = db.session.get(CatalogProduct, row.match_entry_id) if row.match_entry_id else None
        if entry is None:
            # re-check right now: another batch or a colleague may have added it since the preview
            entry = service.find_duplicate(brand=data.get("brand", ""), model_number=data.get("model_number", ""),
                                           regional_version=data.get("regional_version", ""),
                                           barcodes=[data["barcode"]] if data.get("barcode") else [])
        if entry is None:
            payload = {"canonical_name": data["name"], "brand": data.get("brand"), "model_number": data.get("model_number"),
                       "regional_version": data.get("regional_version"), "category": data.get("category"),
                       "description": data.get("description"), "attributes": data.get("attributes") or {},
                       "barcodes": [data["barcode"]] if data.get("barcode") else [], "verification_status": "needs_review",
                       "source": {"type": "supplier", "reference": f"{batch.filename}#{row.row_no}"}}
            if data.get("image_url", "").startswith("https://"):
                payload["media"] = [{"url": data["image_url"], "source_url": data["image_url"], "usage_rights": "unverified"}]
            entry = service.save_entry(payload, by=by)
            entry.field_sources_json = json.dumps(_sources(batch.filename, data), ensure_ascii=False)
            result["created"] += 1
        else:
            filled = False
            for field in ("category", "description", "brand", "model_number"):
                if data.get(field) and not getattr(entry, field):
                    setattr(entry, field, data[field])
                    filled = True
            attrs = json.loads(entry.attributes_json) if entry.attributes_json else {}
            for k, v in (data.get("attributes") or {}).items():
                if k not in attrs:
                    attrs[k], filled = v, True
            entry.attributes_json = json.dumps(attrs, ensure_ascii=False) if attrs else None
            if data.get("barcode") and not CatalogBarcode.query.filter_by(code=data["barcode"]).first():
                db.session.add(CatalogBarcode(catalog_product_id=entry.id, code=data["barcode"]))
                filled = True
            result["updated"] += int(filled)
        if row.match_product_id:
            try:
                service.link(entry.id, row.match_product_id)
                result["linked"] += 1
            except service.CatalogError:
                pass  # the product is already linked to another model: left for review, not forced
        row.applied_entry_id = entry.id
    batch.status, batch.applied_at = "applied", datetime.utcnow()
    db.session.flush()
    return result


def lookup(text: str) -> dict[str, Any]:
    """Smart add, step 1: what do we already know about this barcode / model? Nothing is created here."""
    service.ensure_tables()
    query = str(text or "").strip()
    if len(query) < 3:
        raise service.CatalogError("اكتب الباركود أو رقم الموديل")
    code = service.norm_barcode(query)
    entries, _ = service.search(query, limit=5)
    product_id = _inventory_match(code if len(code) >= 8 else "", query)
    from models.product import Product

    product = Product.query.execution_options(skip_branch_scope=True).get(product_id) if product_id else None
    candidates = []
    from modules.catalog import providers

    for provider in providers.active():
        try:
            for found in provider.lookup(query) or []:
                candidates.append({**found, "provider": provider.name})
        except Exception:  # an external source failing never breaks the screen
            continue
    return {
        "query": query,
        "catalog": [service.serialize(e) for e in entries],
        "inventory": None if product is None else {"id": product.id, "name": product.name, "sku": product.sku or "",
                                                   "barcode": product.barcode or ""},
        "candidates": candidates,
    }
