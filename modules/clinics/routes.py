"""Pages and JSON API for Finora Clinics."""
from __future__ import annotations

from flask import Blueprint, g, jsonify, redirect, render_template, request

from extensions import db
from modules.clinics import service
from modules.clinics.guards import (
    can_use_clinics,
    clinics_guard_response,
    current_branch_id,
    current_tenant_id,
    require_clinics,
)
from modules.clinics.schema import ensure_clinics_schema

clinics_bp = Blueprint("clinics", __name__, url_prefix="/clinics")


@clinics_bp.before_request
def _guard_then_schema():
    # Public infra probe — no auth, no schema side-effects beyond a cheap SELECT.
    if request.path.rstrip("/") == "/clinics/api/health":
        return None
    denied = clinics_guard_response(json_mode=request.path.startswith("/clinics/api/"))
    if denied:
        return denied
    ensure_clinics_schema()


@clinics_bp.get("/api/health")
def api_clinics_health():
    """Safe infrastructure health probe — no secrets, tenants, or stack traces."""
    from sqlalchemy import text

    try:
        db.session.execute(text("SELECT 1"))
        return jsonify({"ok": True, "service": "clinics", "database": "up"})
    except Exception:
        return jsonify({"ok": False, "service": "clinics", "database": "down"}), 503


def _tid() -> int:
    return int(getattr(g, "clinics_tenant_id", None) or current_tenant_id())


def _employee_id():
    row = getattr(g, "clinics_employee", None)
    return row.id if row else None


def _branch_id():
    return getattr(g, "clinics_branch_id", None) or current_branch_id()


def _ok(data=None, **extra):
    payload = {"ok": True}
    if data is not None:
        payload["data"] = data
    payload.update(extra)
    return jsonify(payload)


def _error(exc, status=400):
    db.session.rollback()
    return jsonify({"ok": False, "error": str(exc)}), status


PAGE_CONFIG = {
    "dashboard": ("لوحة العيادة", "clinics_dashboard"),
    "activity": ("سجل النشاط", "clinics_dashboard"),
    "reception": ("الاستقبال", "clinics_reception"),
    "appointments": ("المواعيد", "clinics_appointments"),
    "queue": ("قائمة الانتظار", "clinics_reception"),
    "display": ("شاشة الانتظار", "clinics_reception"),
    "patients": ("المرضى", "clinics_patients"),
    "visits": ("الزيارات", "clinics_visits"),
    "doctor": ("المعاينة الطبية", "clinics_clinical"),
    "clinical": ("السجل الطبي", "clinics_clinical"),
    "labs": ("التحاليل", "clinics_clinical"),
    "radiology": ("الأشعة", "clinics_clinical"),
    "procedures": ("الإجراءات", "clinics_clinical"),
    "billing": ("الفواتير والمدفوعات", "clinics_billing"),
    "settlements": ("تسوية الأطباء", "clinics_billing"),
    "inventory": ("مخزون العيادة", "clinics_inventory"),
    "doctors": ("الأطباء", "clinics_admin"),
    "staff": ("الطاقم والصلاحيات", "clinics_admin"),
    "specialties": ("الاختصاصات", "clinics_admin"),
    "schedules": ("جداول الدوام", "clinics_admin"),
    "services": ("الخدمات", "clinics_admin"),
    "rooms": ("الغرف", "clinics_admin"),
    "reports": ("التقارير", "clinics_reports"),
    "settings": ("إعدادات العيادة", "clinics_settings"),
    "communications": ("التواصل والحجوزات", "clinics_communications_view"),
    "followups": ("المتابعات", "clinics_clinical"),
    "suppliers": ("الموردون", "clinics_inventory"),
    "purchases": ("المشتريات", "purchases.view"),
    "expenses": ("المصروفات", "clinics_billing"),
    "cash": ("الصندوق", "clinics_billing"),
    "attachments": ("المرفقات", "clinics_patients"),
}


@clinics_bp.route("/")
@require_clinics("clinics_dashboard")
def dashboard_page():
    title, _ = PAGE_CONFIG["dashboard"]
    return render_template("clinics/workspace.html", clinics_page="dashboard", clinics_title=title)


@clinics_bp.route("/doctor")
@require_clinics("clinics_clinical")
def doctor_home_page():
    """Resolve the doctor's current open visit, or show waiting state (no error)."""
    from modules.clinics.guards import get_current_clinic_doctor
    from modules.clinics.service_visits_registry import doctor_workspace

    doctor = get_current_clinic_doctor()
    doctor_id = doctor.id if doctor else None
    payload = doctor_workspace(_tid(), doctor_id=doctor_id)
    current = payload.get("current_visit")
    if current and current.get("id"):
        return redirect(f"/clinics/doctor/{int(current['id'])}")
    return render_template(
        "clinics/workspace.html",
        clinics_page="doctor",
        clinics_title="المعاينة الطبية",
        visit_id=None,
    )


@clinics_bp.route("/doctor/<int:visit_id>")
@require_clinics("clinics_clinical")
def doctor_visit_page(visit_id: int):
    return render_template(
        "clinics/workspace.html",
        clinics_page="doctor",
        clinics_title="المعاينة الطبية",
        visit_id=visit_id,
    )


@clinics_bp.get("/api/doctor/workspace")
@require_clinics("clinics_clinical", json_mode=True)
def api_doctor_session_workspace():
    try:
        from models.employee import Employee
        from modules.clinics.guards import get_current_clinic_doctor
        from modules.clinics.service_visits_registry import doctor_workspace

        doctor = get_current_clinic_doctor()
        doctor_id = request.args.get("doctor_id", type=int) or (doctor.id if doctor else None)
        employee = Employee.query.get(_employee_id()) if _employee_id() else None
        return _ok(doctor_workspace(_tid(), doctor_id=doctor_id, employee=employee))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/doctors/<int:doctor_id>")
@require_clinics("clinics_admin")
def doctor_profile_page(doctor_id: int):
    return render_template(
        "clinics/workspace.html",
        clinics_page="doctors",
        clinics_title="ملف الطبيب",
        doctor_id=doctor_id,
    )


@clinics_bp.route("/patients/<int:patient_id>")
@require_clinics("clinics_patients")
def patient_detail_page(patient_id: int):
    return render_template(
        "clinics/workspace.html",
        clinics_page="patient_detail",
        clinics_title="ملف المريض",
        patient_id=patient_id,
    )


@clinics_bp.route("/inventory/items/<int:item_id>")
@require_clinics("clinics_inventory")
def inventory_item_page(item_id: int):
    return render_template(
        "clinics/workspace.html", clinics_page="inventory", clinics_title="تفاصيل الصنف",
        inventory_item_id=item_id, inventory_view="detail",
    )


@clinics_bp.route("/inventory/<view>")
@require_clinics("clinics_inventory")
def inventory_subpage(view: str):
    allowed = {"movements": "حركة المخزون", "batches": "الدفعات والصلاحية", "expiry": "إدارة الصلاحية"}
    if view not in allowed:
        return redirect("/clinics/inventory")
    return render_template(
        "clinics/workspace.html", clinics_page="inventory", clinics_title=allowed[view], inventory_view=view,
    )


@clinics_bp.route("/purchases/new")
@require_clinics("purchases.create")
def purchases_new_page():
    return render_template(
        "clinics/workspace.html",
        clinics_page="purchases",
        clinics_title="فاتورة شراء جديدة",
        purchases_view="new",
    )


@clinics_bp.route("/purchases/reports")
@require_clinics("purchases.export")
def purchases_reports_page():
    return render_template(
        "clinics/workspace.html",
        clinics_page="purchases",
        clinics_title="تقارير المشتريات",
        purchases_view="reports",
    )


@clinics_bp.route("/purchases/<int:purchase_id>/edit")
@require_clinics("purchases.edit")
def purchases_edit_page(purchase_id: int):
    return render_template(
        "clinics/workspace.html",
        clinics_page="purchases",
        clinics_title="تعديل فاتورة الشراء",
        purchase_id=purchase_id,
        purchases_view="edit",
    )


@clinics_bp.route("/purchases/<int:purchase_id>")
@require_clinics("purchases.view")
def purchases_detail_page(purchase_id: int):
    return render_template(
        "clinics/workspace.html",
        clinics_page="purchases",
        clinics_title="تفاصيل فاتورة الشراء",
        purchase_id=purchase_id,
    )


@clinics_bp.route("/calendar")
@require_clinics("clinics_appointments")
def calendar_redirect():
    """Legacy calendar page — unified into Appointment Scheduling Center."""
    return redirect("/clinics/appointments")


@clinics_bp.route("/<page>")
def workspace_page(page):
    if page == "calendar":
        return redirect("/clinics/appointments")
    config = PAGE_CONFIG.get(page)
    if not config:
        return redirect("/clinics/")
    title, permission = config

    @require_clinics(permission)
    def _render():
        return render_template("clinics/workspace.html", clinics_page=page, clinics_title=title)

    return _render()


# ── Bootstrap / Dashboard ──────────────────────────────────────────────────


@clinics_bp.get("/api/bootstrap")
@require_clinics(json_mode=True)
def api_bootstrap():
    try:
        result = service.bootstrap(_tid())
        from modules.clinics.guards import get_current_clinic_doctor
        from modules.clinics.service_dashboard import build_nav_matrix, user_context
        linked = get_current_clinic_doctor()
        result["current_doctor_id"] = linked.id if linked else None
        employee = getattr(g, "clinics_employee", None)
        result["user"] = user_context(employee)
        result["nav"] = build_nav_matrix(
            employee,
            result.get("settings") or {},
            result["current_doctor_id"],
        )
        if not can_use_clinics("clinics_billing") and not can_use_clinics("clinics_admin"):
            for doctor in result["doctors"]:
                doctor.pop("commission_pct", None)
            for svc in result["services"]:
                svc.pop("doctor_commission_pct", None)
            result["settings"].pop("doctor_commission_default_pct", None)
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/dashboard")
@require_clinics("clinics_dashboard", json_mode=True)
def api_dashboard():
    try:
        from modules.clinics.service_dashboard import dashboard_snapshot

        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        doctor_id = request.args.get("doctor_id")
        branch_id = request.args.get("branch_id")
        employee = getattr(g, "clinics_employee", None)
        return _ok(
            dashboard_snapshot(
                _tid(),
                parsed,
                int(doctor_id) if doctor_id else None,
                int(branch_id) if branch_id else None,
                employee,
                view=request.args.get("view", "management"),
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/activity")
@require_clinics("clinics_dashboard", json_mode=True)
def api_activity():
    try:
        from modules.clinics.service_dashboard import activity_log

        date_from = request.args.get("date_from")
        date_to = request.args.get("date_to")
        return _ok(
            activity_log(
                _tid(),
                getattr(g, "clinics_employee", None),
                q=request.args.get("q", ""),
                entity_type=request.args.get("entity_type", ""),
                date_from=service._parse_date(date_from, "من تاريخ") if date_from else None,
                date_to=service._parse_date(date_to, "إلى تاريخ") if date_to else None,
                page=request.args.get("page", 1, type=int) or 1,
                per_page=request.args.get("per_page", 40, type=int) or 40,
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/dashboard/pulse")
@require_clinics("clinics_dashboard", json_mode=True)
def api_dashboard_pulse():
    try:
        from modules.clinics.service_dashboard import dashboard_pulse

        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        doctor_id = request.args.get("doctor_id")
        branch_id = request.args.get("branch_id")
        return _ok(
            dashboard_pulse(
                _tid(),
                parsed,
                int(doctor_id) if doctor_id else None,
                int(branch_id) if branch_id else None,
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/check-phone")
@require_clinics("clinics_patients", json_mode=True)
def api_check_phone():
    try:
        phone = request.args.get("phone", "")
        found = service.find_patient_by_phone(_tid(), phone)
        return _ok({"exists": bool(found), "patient": found})
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/check-similar")
@require_clinics("clinics_patients", json_mode=True)
def api_check_similar():
    try:
        full_name = request.args.get("full_name", "")
        dob_raw = request.args.get("date_of_birth", "")
        dob = service._parse_date(dob_raw, "تاريخ الميلاد") if dob_raw else None
        matches = service.find_similar_patients(_tid(), full_name, dob)
        return _ok({"matches": matches, "has_similar": bool(matches)})
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/registry")
@require_clinics("clinics_patients", json_mode=True)
def api_patients_registry():
    try:
        from models.employee import Employee
        from modules.clinics.service_patients_registry import patients_registry

        employee = Employee.query.get(_employee_id()) if _employee_id() else None
        params = {
            "q": request.args.get("q", ""),
            "page": request.args.get("page", 1),
            "per_page": request.args.get("per_page", 25),
            "sort": request.args.get("sort", "created_at"),
            "status_group": request.args.get("status_group", "active"),
            "gender": request.args.get("gender", ""),
            "doctor_id": request.args.get("doctor_id"),
            "upcoming_only": request.args.get("upcoming_only") == "1",
            "balance_only": request.args.get("balance_only") == "1",
            "followup_only": request.args.get("followup_only") == "1",
            "new_only": request.args.get("new_only") == "1",
            "last_visit": request.args.get("last_visit", ""),
            "no_visit_days": request.args.get("no_visit_days"),
        }
        return _ok(patients_registry(_tid(), params, employee))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/export")
@require_clinics("clinics_patients", json_mode=True)
def api_patients_export():
    try:
        import csv
        import io

        from flask import Response
        from models.employee import Employee
        from modules.clinics.service_patients_registry import patients_export_rows

        employee = Employee.query.get(_employee_id()) if _employee_id() else None
        params = {
            "q": request.args.get("q", ""),
            "sort": request.args.get("sort", "created_at"),
            "status_group": request.args.get("status_group", "active"),
            "gender": request.args.get("gender", ""),
            "doctor_id": request.args.get("doctor_id"),
            "upcoming_only": request.args.get("upcoming_only") == "1",
            "balance_only": request.args.get("balance_only") == "1",
            "followup_only": request.args.get("followup_only") == "1",
            "new_only": request.args.get("new_only") == "1",
            "last_visit": request.args.get("last_visit", ""),
            "no_visit_days": request.args.get("no_visit_days"),
        }
        rows = patients_export_rows(_tid(), params, employee)
        fmt = str(request.args.get("format") or "csv").lower()
        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(
                ["file_number", "full_name", "phone", "age", "gender", "last_visit", "next_appt", "balance", "status"]
            )
            for p in rows:
                lv = p.get("last_visit") or {}
                na = p.get("next_appointment") or {}
                writer.writerow(
                    [
                        p.get("file_number"),
                        p.get("full_name"),
                        p.get("phone"),
                        p.get("age"),
                        p.get("gender"),
                        lv.get("visit_date") or "",
                        f"{na.get('appointment_date') or ''} {na.get('start_time') or ''}".strip(),
                        p.get("balance") if p.get("balance") is not None else "",
                        p.get("operational_status"),
                    ]
                )
            return Response(
                buf.getvalue(),
                mimetype="text/csv; charset=utf-8",
                headers={"Content-Disposition": "attachment; filename=patients.csv"},
            )
        if fmt in ("xlsx", "excel"):
            from io import BytesIO

            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Patients"
            ws.append(
                ["file_number", "full_name", "phone", "age", "gender", "last_visit", "next_appt", "balance", "status"]
            )
            for p in rows:
                lv = p.get("last_visit") or {}
                na = p.get("next_appointment") or {}
                ws.append(
                    [
                        p.get("file_number"),
                        p.get("full_name"),
                        p.get("phone"),
                        p.get("age"),
                        p.get("gender"),
                        lv.get("visit_date") or "",
                        f"{na.get('appointment_date') or ''} {na.get('start_time') or ''}".strip(),
                        p.get("balance") if p.get("balance") is not None else "",
                        p.get("operational_status"),
                    ]
                )
            xbuf = BytesIO()
            wb.save(xbuf)
            return Response(
                xbuf.getvalue(),
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=patients.xlsx"},
            )
        return _error(ValueError("صيغة التصدير غير مدعومة"))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/queue/reorder")
@require_clinics("clinics_reception", json_mode=True)
def api_queue_reorder():
    try:
        payload = request.get_json(silent=True) or {}
        day = payload.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        order = payload.get("order") or []
        result = service.reorder_waiting_queue(_tid(), parsed, [int(x) for x in order], _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/queue/recall")
@require_clinics("clinics_reception", json_mode=True)
def api_queue_recall():
    try:
        payload = request.get_json(silent=True) or {}
        appointment_id = int(payload.get("appointment_id") or 0)
        if not appointment_id:
            raise ValueError("الموعد مطلوب")
        result = service.recall_queue_patient(_tid(), appointment_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/queue/enqueue")
@require_clinics("clinics_reception", json_mode=True)
def api_queue_enqueue():
    try:
        payload = request.get_json(silent=True) or {}
        result = service.enqueue_walkin(_tid(), payload, _employee_id(), _branch_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/appointments/<int:appointment_id>/priority")
@require_clinics("clinics_reception", json_mode=True)
def api_appointment_priority(appointment_id: int):
    try:
        payload = request.get_json(silent=True) or {}
        result = service.set_queue_priority(
            _tid(),
            appointment_id,
            payload.get("priority", 0),
            _employee_id(),
            reason=payload.get("reason"),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/appointments/<int:appointment_id>/emergency")
@require_clinics("clinics_reception", json_mode=True)
def api_appointment_emergency(appointment_id: int):
    try:
        payload = request.get_json(silent=True) or {}
        result = service.set_appointment_emergency(
            _tid(),
            appointment_id,
            bool(payload.get("emergency", True)),
            _employee_id(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/search")
@require_clinics(json_mode=True)
def api_search():
    try:
        from modules.clinics.service_dashboard import global_search

        result = global_search(_tid(), request.args.get("q", ""))
        if not can_use_clinics("clinics_billing"):
            result["invoices"] = []
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/settings", methods=["GET", "PUT", "POST"])
@require_clinics("clinics_settings", json_mode=True)
def api_settings():
    try:
        if request.method == "GET":
            payload = service.settings_payload(_tid())
            if not can_use_clinics("clinics_billing") and not can_use_clinics("clinics_admin"):
                payload.pop("doctor_commission_default_pct", None)
            return _ok(payload)
        data = request.get_json(silent=True) or {}
        if not can_use_clinics("clinics_billing") and not can_use_clinics("clinics_admin"):
            data.pop("doctor_commission_default_pct", None)
        result = service.update_settings(_tid(), data, _employee_id())
        if not can_use_clinics("clinics_billing") and not can_use_clinics("clinics_admin"):
            result.pop("doctor_commission_default_pct", None)
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/specialty-registry")
@require_clinics("clinics_clinical", json_mode=True)
def api_specialty_registry():
    try:
        from modules.clinics.service_specialty_registry import ALIASES, REGISTRY
        return _ok({
            "items": [
                {"code": code, "workspace": item["workspace"], "features": item["features"]}
                for code, item in sorted(REGISTRY.items())
            ],
            "aliases": ALIASES,
        })
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/specialty-resolver")
@require_clinics("clinics_clinical", json_mode=True)
def api_specialty_resolver():
    try:
        from modules.clinics.models import ClinicAppointment, ClinicVisit
        from modules.clinics.service_specialty_registry import resolve_specialty
        visit_id = request.args.get("visit_id", type=int)
        appointment_id = request.args.get("appointment_id", type=int)
        visit = ClinicVisit.query.filter_by(tenant_id=_tid(), id=visit_id).first() if visit_id else None
        appointment = (
            ClinicAppointment.query.filter_by(tenant_id=_tid(), id=appointment_id).first()
            if appointment_id else None
        )
        if visit_id and not visit:
            raise ValueError("الزيارة غير موجودة")
        if appointment_id and not appointment:
            raise ValueError("الموعد غير موجود")
        return _ok(resolve_specialty(
            _tid(), visit=visit, appointment=appointment,
            doctor_id=request.args.get("doctor_id", type=int),
        ))
    except Exception as exc:
        return _error(exc)


def _tenant_slug() -> str:
    from flask import session

    return str(
        session.get("tenant_slug")
        or getattr(g, "tenant", None)
        or ""
    ).strip()


@clinics_bp.get("/api/guides")
@require_clinics("clinics_settings", json_mode=True)
def api_specialty_guides():
    """Catalog of printable specialty operating guides for Settings."""
    try:
        from modules.clinics.specialty_guides import list_guide_catalog
        from modules.clinics.service_specialties import list_specialties

        active_codes = {
            s["code"]
            for s in list_specialties(_tid(), active_only=True)
            if s.get("code")
        }
        rows = []
        for item in list_guide_catalog():
            rows.append(
                {
                    **item,
                    "is_active": item["code"] in active_codes,
                    "url": f"/clinics/guides/{item['code']}",
                }
            )
        # Prefer active specialties first
        rows.sort(key=lambda r: (not r["is_active"], r["name_ar"]))
        return _ok({"items": rows, "tenant_slug": _tenant_slug()})
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/guides/<code>")
@require_clinics()
def specialty_guide_page(code: str):
    """Printable HTML operating guide personalized from clinic settings."""
    try:
        from modules.clinics.specialty_guides import build_guide_context
        from modules.clinics.service_specialties import list_specialties

        settings = service.settings_payload(_tid())
        active_codes = [
            s["code"]
            for s in list_specialties(_tid(), active_only=True)
            if s.get("code")
        ]
        guide = build_guide_context(
            code=code,
            clinic_name=settings.get("clinic_name") or "عيادة Finora",
            organization_mode=settings.get("organization_mode") or "solo_clinic",
            tenant_slug=_tenant_slug(),
            booking_enabled=bool(settings.get("allow_online_booking")),
            active_specialty_codes=active_codes,
        )
        return render_template("clinics/specialty_guide.html", guide=guide)
    except ValueError as exc:
        return redirect("/clinics/settings")
    except Exception:
        return redirect("/clinics/settings")


@clinics_bp.route("/api/patients", methods=["GET", "POST"])
@require_clinics("clinics_patients", json_mode=True)
def api_patients():
    try:
        if request.method == "GET":
            try:
                limit = int(request.args.get("limit") or 100)
            except (TypeError, ValueError):
                limit = 100
            return _ok(service.list_patients(_tid(), q=request.args.get("q", ""), limit=limit))
        result = service.create_patient(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>", methods=["GET", "PUT", "PATCH"])
@require_clinics("clinics_patients", json_mode=True)
def api_patient_detail(patient_id: int):
    try:
        if request.method == "GET":
            return _ok(service.get_patient(_tid(), patient_id, include_clinical=can_use_clinics("clinics_clinical")))
        result = service.update_patient(
            _tid(), patient_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Doctors / Services / Rooms ─────────────────────────────────────────────


@clinics_bp.get("/api/staff")
@require_clinics("clinics_admin", json_mode=True)
def api_staff_center():
    try:
        from modules.clinics.service_staff import staff_center

        return _ok(staff_center(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/staff")
@require_clinics("clinics_admin", json_mode=True)
def api_staff_create():
    try:
        from modules.clinics.service_staff import create_staff
        from utils.plan_guard import users_limit_check

        limit = users_limit_check()
        if not limit.get("ok"):
            raise ValueError(limit.get("error") or "تم بلوغ حد المستخدمين في الخطة")
        result = create_staff(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/staff/<int:employee_id>")
@require_clinics("clinics_admin", json_mode=True)
def api_staff_update(employee_id: int):
    try:
        from modules.clinics.service_staff import update_staff

        result = update_staff(
            _tid(), employee_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/doctors/center")
@require_clinics("clinics_admin", json_mode=True)
def api_doctors_center():
    try:
        from modules.clinics.service_doctors import doctors_center

        return _ok(doctors_center(_tid(), request.args))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/doctors/employees")
@require_clinics("clinics_admin", json_mode=True)
def api_doctor_employee_candidates():
    try:
        from modules.clinics.service_doctors import employee_candidates

        return _ok(employee_candidates(_tid(), request.args.get("q", "")))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/doctors/<int:doctor_id>/workspace")
@require_clinics("clinics_admin", json_mode=True)
def api_doctor_workspace(doctor_id: int):
    try:
        from modules.clinics.service_doctors import doctor_workspace

        employee = getattr(g, "clinics_employee", None)
        can_view_salary = bool(
            employee and (
                employee.role == "admin"
                or employee.has_permission("salary_sensitive_data.view")
                or employee.has_permission("payroll.view")
            )
        )
        return _ok(doctor_workspace(_tid(), doctor_id, can_view_salary=can_view_salary))
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/doctors/<int:doctor_id>/employee-link")
@require_clinics("clinics_admin", json_mode=True)
def api_doctor_employee_link(doctor_id: int):
    try:
        from modules.clinics.service_doctors import link_employee

        payload = request.get_json(silent=True) or {}
        employee_id = int(payload.get("employee_id") or 0) or None
        result = link_employee(_tid(), doctor_id, employee_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/doctors", methods=["GET", "POST"])
@require_clinics("clinics_admin", json_mode=True)
def api_doctors():
    try:
        if request.method == "GET":
            return _ok(service.list_doctors(_tid(), active_only=request.args.get("active") == "1"))
        result = service.create_doctor(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/doctors/<int:doctor_id>", methods=["PUT", "PATCH"])
@require_clinics("clinics_admin", json_mode=True)
def api_doctor_update(doctor_id: int):
    try:
        result = service.update_doctor(
            _tid(), doctor_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/specialties", methods=["GET", "POST"])
@require_clinics("clinics_admin", json_mode=True)
def api_specialties():
    try:
        from modules.clinics.service_specialties import create_specialty, list_specialties
        if request.method == "GET":
            return _ok(list_specialties(_tid(), request.args.get("active") == "1"))
        result = create_specialty(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/specialties/<int:specialty_id>", methods=["GET", "PUT", "PATCH"])
@require_clinics("clinics_admin", json_mode=True)
def api_specialty_detail(specialty_id: int):
    try:
        from modules.clinics.models import ClinicSpecialty
        from modules.clinics.service_specialties import specialty_dict, update_specialty
        if request.method == "GET":
            row = ClinicSpecialty.query.filter_by(tenant_id=_tid(), id=specialty_id).first()
            if not row:
                raise ValueError("الاختصاص غير موجود")
            return _ok(specialty_dict(row))
        result = update_specialty(_tid(), specialty_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/specialties/<int:specialty_id>/templates", methods=["GET", "POST"])
@require_clinics("clinics_admin", json_mode=True)
def api_specialty_templates(specialty_id: int):
    try:
        from modules.clinics.models import ClinicVisitTemplate
        from modules.clinics.service_specialties import create_template, template_dict
        if request.method == "GET":
            rows = ClinicVisitTemplate.query.filter_by(tenant_id=_tid(), specialty_id=specialty_id).order_by(ClinicVisitTemplate.version.desc()).all()
            return _ok([template_dict(_tid(), row.id, include_fields=False) for row in rows])
        result = create_template(_tid(), specialty_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visit-templates/<int:template_id>/publish")
@require_clinics("clinics_admin", json_mode=True)
def api_publish_visit_template(template_id: int):
    try:
        from modules.clinics.service_specialties import publish_template
        result = publish_template(_tid(), template_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/visit-templates/<int:template_id>", methods=["GET", "PUT"])
@require_clinics("clinics_admin", json_mode=True)
def api_visit_template(template_id: int):
    try:
        from modules.clinics.service_specialties import template_dict, update_template
        if request.method == "GET":
            return _ok(template_dict(_tid(), template_id))
        result = update_template(_tid(), template_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visit-templates/<int:template_id>/archive")
@require_clinics("clinics_admin", json_mode=True)
def api_archive_visit_template(template_id: int):
    try:
        from modules.clinics.service_specialties import archive_template
        result = archive_template(_tid(), template_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/services", methods=["GET", "POST"])
@require_clinics("clinics_admin", json_mode=True)
def api_services():
    try:
        if request.method == "GET":
            if request.args.get("center") == "1":
                return _ok(service.services_center(_tid(), request.args.to_dict()))
            return _ok(service.list_services(_tid(), active_only=request.args.get("active") == "1"))
        result = service.create_service(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/services/<int:service_id>", methods=["PUT", "PATCH"])
@require_clinics("clinics_admin", json_mode=True)
def api_service_update(service_id: int):
    try:
        result = service.update_service(
            _tid(), service_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/rooms", methods=["GET", "POST"])
@require_clinics("clinics_admin", json_mode=True)
def api_rooms():
    try:
        if request.method == "GET":
            return _ok(service.list_rooms(_tid(), active_only=request.args.get("active") != "0"))
        result = service.create_room(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Appointments / Reception / Queue ───────────────────────────────────────


@clinics_bp.get("/api/reception")
@require_clinics("clinics_reception", json_mode=True)
def api_reception():
    try:
        from modules.clinics.service_reception import get_reception_board

        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        doctor_id = request.args.get("doctor_id")
        branch_id = request.args.get("branch_id")
        employee = getattr(g, "clinics_employee", None)
        return _ok(
            get_reception_board(
                _tid(),
                parsed,
                int(doctor_id) if doctor_id else None,
                int(branch_id) if branch_id else None,
                employee,
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/appointments/slots")
@require_clinics("clinics_appointments", json_mode=True)
def api_appointment_slots():
    try:
        from modules.clinics.service_booking import available_slots

        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else service._today()
        doctor_id = request.args.get("doctor_id")
        if not doctor_id:
            raise ValueError("الطبيب مطلوب")
        service_id = request.args.get("service_id")
        exclude_id = request.args.get("exclude_id")
        slots = available_slots(
            _tid(),
            int(doctor_id),
            parsed,
            int(service_id) if service_id else None,
            int(exclude_id) if exclude_id else None,
        )
        return _ok({"slots": slots})
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/appointments/<int:appointment_id>")
@require_clinics("clinics_patients", json_mode=True)
def api_appointment_detail(appointment_id: int):
    try:
        return _ok(service.get_appointment(_tid(), appointment_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/appointments/board")
@require_clinics("clinics_appointments", json_mode=True)
def api_appointments_board():
    try:
        from modules.clinics.service_appointments import get_appointments_board

        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        doctor_id = request.args.get("doctor_id")
        status = request.args.get("status")
        service_id = request.args.get("service_id")
        specialty_id = request.args.get("specialty_id")
        branch_id = request.args.get("branch_id")
        view = request.args.get("view") or "day"
        employee = getattr(g, "clinics_employee", None)
        return _ok(
            get_appointments_board(
                _tid(), day=parsed, view=view,
                doctor_id=int(doctor_id) if doctor_id else None,
                status=status or None,
                service_id=int(service_id) if service_id else None,
                specialty_id=int(specialty_id) if specialty_id else None,
                branch_id=int(branch_id) if branch_id else None,
                employee=employee,
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/appointments", methods=["GET", "POST"])
@require_clinics("clinics_appointments", json_mode=True)
def api_appointments():
    try:
        if request.method == "GET":
            day = request.args.get("date")
            parsed = service._parse_date(day, "التاريخ") if day else None
            doctor_id = request.args.get("doctor_id")
            status = request.args.get("status")
            return _ok(
                service.list_appointments(
                    _tid(),
                    parsed,
                    int(doctor_id) if doctor_id else None,
                    status,
                )
            )
        result = service.create_appointment(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/appointments/<int:appointment_id>")
@require_clinics("clinics_reception", json_mode=True)
def api_patch_appointment(appointment_id: int):
    try:
        result = service.patch_appointment(
            _tid(),
            appointment_id,
            request.get_json(silent=True) or {},
            _employee_id(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/appointments/<int:appointment_id>/status")
@require_clinics("clinics_reception", json_mode=True)
def api_appointment_status(appointment_id: int):
    try:
        payload = request.get_json(silent=True) or {}
        result = service.change_appointment_status(
            _tid(),
            appointment_id,
            payload.get("status"),
            _employee_id(),
            payload.get("note"),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Visits / Clinical ──────────────────────────────────────────────────────


@clinics_bp.route("/api/visits", methods=["GET", "POST"])
@require_clinics("clinics_visits", json_mode=True)
def api_visits():
    try:
        if request.method == "GET":
            day = request.args.get("date")
            parsed = service._parse_date(day, "التاريخ") if day else None
            return _ok(service.list_visits(_tid(), parsed, request.args.get("status")))
        payload = request.get_json(silent=True) or {}
        if not payload.get("doctor_id") and getattr(g, "clinics_doctor_scope", None):
            payload["doctor_id"] = g.clinics_doctor_scope
        if payload.get("appointment_id"):
            result = service.start_visit_from_appointment(
                _tid(), int(payload["appointment_id"]), _employee_id()
            )
        else:
            result = service.create_walkin_visit(
                _tid(), payload, _employee_id(), _branch_id()
            )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/registry")
@require_clinics("clinics_visits", json_mode=True)
def api_visits_registry():
    try:
        from models.employee import Employee
        from modules.clinics.service_visits_registry import visits_registry

        employee = Employee.query.get(_employee_id()) if _employee_id() else None
        params = {
            "q": request.args.get("q", ""),
            "date": request.args.get("date"),
            "page": request.args.get("page", 1),
            "per_page": request.args.get("per_page", 25),
            "sort": request.args.get("sort", "time"),
            "status": request.args.get("status", ""),
            "doctor_id": request.args.get("doctor_id"),
            "service_id": request.args.get("service_id"),
            "incomplete_only": request.args.get("incomplete_only") == "1",
            "has_rx": request.args.get("has_rx") == "1",
            "has_lab": request.args.get("has_lab") == "1",
            "has_radiology": request.args.get("has_radiology") == "1",
            "unpaid_only": request.args.get("unpaid_only") == "1",
            "followup_only": request.args.get("followup_only") == "1",
        }
        return _ok(visits_registry(_tid(), params, employee))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_get(visit_id: int):
    try:
        return _ok(service.get_visit_detail(_tid(), visit_id, include_financial=can_use_clinics("clinics_billing")))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/visits/<int:visit_id>", methods=["PUT", "PATCH"])
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_update(visit_id: int):
    try:
        result = service.update_visit(
            _tid(), visit_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/vitals")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_vitals(visit_id: int):
    try:
        result = service.save_vitals(
            _tid(), visit_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.put("/api/visits/<int:visit_id>/specialty-values")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_specialty_values(visit_id: int):
    try:
        from modules.clinics.service_specialties import save_visit_values
        result = save_visit_values(_tid(), visit_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>/dental-chart")
@require_clinics("dental.chart.view", json_mode=True)
def api_visit_dental_chart(visit_id: int):
    try:
        from modules.clinics.service_dental import get_visit_dental_chart
        return _ok(get_visit_dental_chart(_tid(), visit_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.put("/api/visits/<int:visit_id>/dental-chart/tooth/<tooth_number>")
@require_clinics("dental.chart.edit", json_mode=True)
def api_visit_dental_chart_tooth(visit_id: int, tooth_number: str):
    try:
        from modules.clinics.service_dental import upsert_tooth_finding
        result = upsert_tooth_finding(
            _tid(), visit_id, tooth_number, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/dental-history")
@require_clinics("dental.chart.view", json_mode=True)
def api_patient_dental_history(patient_id: int):
    try:
        from modules.clinics.service_dental import get_patient_dental_history
        return _ok(get_patient_dental_history(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/dental-chart")
@require_clinics("dental.chart.view", json_mode=True)
def api_patient_dental_chart(patient_id: int):
    try:
        from modules.clinics.service_dental import get_patient_dental_chart
        return _ok(get_patient_dental_chart(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


def _dental_tp_financial():
    return can_use_clinics("clinics_billing")


@clinics_bp.get("/api/patients/<int:patient_id>/dental-treatment-plans")
@require_clinics("dental.treatment_plan.view", json_mode=True)
def api_patient_dental_treatment_plans(patient_id: int):
    try:
        from modules.clinics.service_dental_treatment import list_patient_plans
        return _ok(list_patient_plans(_tid(), patient_id, include_financial=_dental_tp_financial()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/patients/<int:patient_id>/dental-treatment-plans")
@require_clinics("dental.treatment_plan.create", json_mode=True)
def api_patient_dental_treatment_plan_create(patient_id: int):
    try:
        from modules.clinics.service_dental_treatment import create_plan
        result = create_plan(_tid(), patient_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/dental-treatment-plans/<int:plan_id>")
@require_clinics("dental.treatment_plan.view", json_mode=True)
def api_dental_treatment_plan_get(plan_id: int):
    try:
        from modules.clinics.service_dental_treatment import plan_detail
        return _ok(plan_detail(_tid(), plan_id, include_financial=_dental_tp_financial()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/dental-treatment-plans/<int:plan_id>")
@require_clinics("dental.treatment_plan.edit", json_mode=True)
def api_dental_treatment_plan_patch(plan_id: int):
    try:
        from modules.clinics.service_dental_treatment import update_plan
        result = update_plan(
            _tid(), plan_id, request.get_json(silent=True) or {}, _employee_id(),
            include_financial=_dental_tp_financial(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/dental-treatment-plans/<int:plan_id>/items")
@require_clinics("dental.treatment_plan.create", json_mode=True)
def api_dental_treatment_plan_add_item(plan_id: int):
    try:
        from modules.clinics.service_dental_treatment import add_plan_item
        result = add_plan_item(
            _tid(), plan_id, request.get_json(silent=True) or {}, _employee_id(),
            include_financial=_dental_tp_financial(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/dental-chart/tooth/<tooth_number>/findings")
@require_clinics("dental.chart.edit", json_mode=True)
def api_visit_dental_chart_finding_add(visit_id: int, tooth_number: str):
    try:
        from modules.clinics.service_dental import add_active_tooth_finding
        result = add_active_tooth_finding(
            _tid(), visit_id, tooth_number, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/dental-chart/tooth/<tooth_number>/findings/<int:event_id>/resolve")
@require_clinics("dental.chart.edit", json_mode=True)
def api_visit_dental_chart_finding_resolve(visit_id: int, tooth_number: str, event_id: int):
    try:
        from modules.clinics.service_dental import resolve_active_tooth_finding
        result = resolve_active_tooth_finding(
            _tid(), visit_id, tooth_number, event_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/dental-treatment-plans/<int:plan_id>/options")
@require_clinics("dental.treatment_plan.edit", json_mode=True)
def api_dental_treatment_option_create(plan_id: int):
    try:
        from modules.clinics.service_dental_treatment import create_treatment_option
        result = create_treatment_option(
            _tid(), plan_id, request.get_json(silent=True) or {}, _employee_id(),
            include_financial=_dental_tp_financial(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/dental-treatment-options/<int:option_id>")
@require_clinics("dental.treatment_plan.edit", json_mode=True)
def api_dental_treatment_option_patch(option_id: int):
    try:
        from modules.clinics.service_dental_treatment import update_treatment_option
        result = update_treatment_option(
            _tid(), option_id, request.get_json(silent=True) or {}, _employee_id(),
            include_financial=_dental_tp_financial(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/dental-treatment-plans/<int:plan_id>/accept")
@require_clinics("dental.treatment_plan.edit", json_mode=True)
def api_dental_treatment_plan_accept(plan_id: int):
    try:
        from modules.clinics.service_dental_treatment import accept_plan
        result = accept_plan(
            _tid(), plan_id, request.get_json(silent=True) or {}, _employee_id(),
            include_financial=_dental_tp_financial(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/dental-treatment-plan-items/<int:item_id>")
@require_clinics("dental.treatment_plan.edit", json_mode=True)
def api_dental_treatment_plan_item_patch(item_id: int):
    try:
        from modules.clinics.service_dental_treatment import update_plan_item
        result = update_plan_item(
            _tid(), item_id, request.get_json(silent=True) or {}, _employee_id(),
            include_financial=_dental_tp_financial(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/dental-treatment-plan-items/<int:item_id>/procedure-links")
@require_clinics("dental.treatment_plan.edit", json_mode=True)
def api_dental_treatment_plan_item_link(item_id: int):
    try:
        from modules.clinics.service_dental_treatment import link_item_procedure
        payload = request.get_json(silent=True) or {}
        vp_id = int(payload.get("visit_procedure_id") or 0)
        if not vp_id:
            raise ValueError("visit_procedure_id مطلوب")
        result = link_item_procedure(
            _tid(), item_id, vp_id, _employee_id(), include_financial=_dental_tp_financial()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/dental-treatment-plans/tooth/<tooth_number>")
@require_clinics("dental.treatment_plan.view", json_mode=True)
def api_patient_tooth_treatment_context(patient_id: int, tooth_number: str):
    try:
        from modules.clinics.service_dental_treatment import tooth_plan_context
        return _ok(tooth_plan_context(
            _tid(), patient_id, tooth_number, include_financial=_dental_tp_financial()
        ))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/dental-treatment-plans/<int:plan_id>/print")
@require_clinics("dental.treatment_plan.view", json_mode=True)
def api_dental_treatment_plan_print(plan_id: int):
    try:
        from modules.clinics.service_dental_treatment import print_plan_payload
        return _ok(print_plan_payload(_tid(), plan_id, include_financial=_dental_tp_financial()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/dental-treatment-plan-items/<int:item_id>/sessions", methods=["GET", "POST"])
@require_clinics("dental.sessions.view", json_mode=True)
def api_dental_sessions(item_id: int):
    try:
        from modules.clinics.service_dental_platform import create_session, list_sessions
        if request.method == "GET":
            return _ok(list_sessions(_tid(), item_id=item_id))
        if not can_use_clinics("dental.sessions.create"):
            return jsonify({"ok": False, "error": "لا تملك صلاحية إنشاء جلسات الأسنان"}), 403
        result = create_session(_tid(), item_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/dental-sessions/<int:session_id>")
@require_clinics("dental.sessions.edit", json_mode=True)
def api_dental_session_patch(session_id: int):
    try:
        from modules.clinics.service_dental_platform import update_session
        result = update_session(_tid(), session_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/dental-sessions/<int:session_id>/complete")
@require_clinics("dental.sessions.complete", json_mode=True)
def api_dental_session_complete(session_id: int):
    try:
        from modules.clinics.service_dental_platform import complete_session
        result = complete_session(_tid(), session_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/dental-sessions/<int:session_id>/cancel")
@require_clinics("dental.sessions.edit", json_mode=True)
def api_dental_session_cancel(session_id: int):
    try:
        from modules.clinics.service_dental_platform import cancel_session
        result = cancel_session(_tid(), session_id, _employee_id(), request.get_json(silent=True) or {})
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/dental-images", methods=["GET", "POST"])
@require_clinics("dental.images.view", json_mode=True)
def api_dental_images():
    try:
        from modules.clinics.service_dental_platform import create_image_metadata, list_images
        if request.method == "GET":
            return _ok(list_images(_tid(), request.args.to_dict()))
        if not can_use_clinics("dental.images.upload"):
            return jsonify({"ok": False, "error": "لا تملك صلاحية رفع صور الأسنان"}), 403
        result = create_image_metadata(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.delete("/api/dental-images/<int:image_id>")
@require_clinics("dental.images.delete", json_mode=True)
def api_dental_image_delete(image_id: int):
    try:
        from modules.clinics.service_dental_platform import delete_image_metadata
        result = delete_image_metadata(_tid(), image_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/reports/dental")
@require_clinics("dental.reports.view", json_mode=True)
def api_dental_reports():
    try:
        from modules.clinics.service_dental_platform import dental_reports
        return _ok(dental_reports(_tid(), patient_id=request.args.get("patient_id", type=int)))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/dental-tooth-events")
@require_clinics("dental.chart.view", json_mode=True)
def api_patient_dental_tooth_events(patient_id: int):
    try:
        from modules.clinics.service_dental_completion import list_tooth_events
        return _ok(list_tooth_events(
            _tid(),
            patient_id,
            request.args.get("tooth_number"),
            request.args.get("limit", default=50, type=int),
        ))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/dental-active-treatments")
@require_clinics("dental.treatment_plan.view", json_mode=True)
def api_patient_dental_active_treatments(patient_id: int):
    try:
        from modules.clinics.service_dental_completion import patient_active_treatments
        return _ok(patient_active_treatments(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>/dental-complete-preflight")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_dental_complete_preflight(visit_id: int):
    try:
        from modules.clinics.service_dental_completion import visit_complete_preflight
        return _ok(visit_complete_preflight(_tid(), visit_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/dental-treatment-plans/<int:plan_id>/invoice")
@require_clinics("clinics_billing", json_mode=True)
def api_dental_treatment_plan_invoice(plan_id: int):
    try:
        from modules.clinics.service_dental_completion import create_invoice_from_treatment_plan
        result = create_invoice_from_treatment_plan(
            _tid(), plan_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/pediatric-context")
@require_clinics("clinics_clinical", json_mode=True)
def api_pediatric_context(patient_id: int):
    try:
        from modules.clinics.service_pediatrics import get_pediatric_context
        return _ok(get_pediatric_context(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/pediatric-profile", methods=["GET", "PUT"])
@require_clinics("clinics_clinical", json_mode=True)
def api_pediatric_profile(patient_id: int):
    try:
        from modules.clinics.service_pediatrics import get_pediatric_profile, upsert_pediatric_profile
        if request.method == "GET":
            return _ok(get_pediatric_profile(_tid(), patient_id))
        result = upsert_pediatric_profile(_tid(), patient_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/pediatric-growth")
@require_clinics("clinics_clinical", json_mode=True)
def api_pediatric_growth(patient_id: int):
    try:
        from modules.clinics.service_pediatrics import get_pediatric_growth
        return _ok(get_pediatric_growth(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/vaccinations", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_patient_vaccinations(patient_id: int):
    try:
        from modules.clinics.service_pediatrics import add_vaccination, list_vaccinations
        if request.method == "GET":
            return _ok(list_vaccinations(_tid(), patient_id))
        result = add_vaccination(_tid(), patient_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/pediatric-vaccinations/<int:vax_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_pediatric_vaccination_patch(vax_id: int):
    try:
        from modules.clinics.service_pediatrics import update_vaccination
        result = update_vaccination(_tid(), vax_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/development", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_patient_development(patient_id: int):
    try:
        from modules.clinics.service_pediatrics import add_development, list_development
        if request.method == "GET":
            return _ok(list_development(_tid(), patient_id))
        result = add_development(_tid(), patient_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/pediatric-development/<int:record_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_pediatric_development_patch(record_id: int):
    try:
        from modules.clinics.service_pediatrics import update_development
        result = update_development(_tid(), record_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/gynecology-context")
@require_clinics("clinics_clinical", json_mode=True)
def api_gynecology_context(patient_id: int):
    try:
        from modules.clinics.service_gynecology import get_gynecology_context
        return _ok(get_gynecology_context(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>/gynecology-context")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_gynecology_context(visit_id: int):
    try:
        from modules.clinics.service_gynecology import get_visit_gynecology_context
        return _ok(get_visit_gynecology_context(_tid(), visit_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/gynecology-profile", methods=["GET", "PUT"])
@require_clinics("clinics_clinical", json_mode=True)
def api_gynecology_profile(patient_id: int):
    try:
        from modules.clinics.service_gynecology import get_gynecology_profile, upsert_gynecology_profile
        if request.method == "GET":
            return _ok(get_gynecology_profile(_tid(), patient_id))
        result = upsert_gynecology_profile(_tid(), patient_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/pregnancies", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_patient_pregnancies(patient_id: int):
    try:
        from modules.clinics.service_gynecology import create_pregnancy, list_pregnancies
        if request.method == "GET":
            return _ok(list_pregnancies(_tid(), patient_id))
        payload = request.get_json(silent=True) or {}
        from modules.clinics.guards import get_current_clinic_doctor
        linked = get_current_clinic_doctor()
        doctor_id = payload.get("doctor_id") or (linked.id if linked else None)
        result = create_pregnancy(_tid(), patient_id, payload, _employee_id(), doctor_id=doctor_id)
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/pregnancies/<int:pregnancy_id>", methods=["GET", "PATCH"])
@require_clinics("clinics_clinical", json_mode=True)
def api_pregnancy(pregnancy_id: int):
    try:
        from modules.clinics.service_gynecology import get_pregnancy, update_pregnancy
        if request.method == "GET":
            return _ok(get_pregnancy(_tid(), pregnancy_id))
        result = update_pregnancy(_tid(), pregnancy_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/pregnancies/<int:pregnancy_id>/context")
@require_clinics("clinics_clinical", json_mode=True)
def api_pregnancy_context(pregnancy_id: int):
    try:
        from modules.clinics.service_gynecology import get_pregnancy_context
        return _ok(get_pregnancy_context(_tid(), pregnancy_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/pregnancies/<int:pregnancy_id>/timeline")
@require_clinics("clinics_clinical", json_mode=True)
def api_pregnancy_timeline(pregnancy_id: int):
    try:
        from modules.clinics.service_gynecology import build_pregnancy_timeline
        return _ok(build_pregnancy_timeline(_tid(), pregnancy_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/pregnancies/<int:pregnancy_id>/observations", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_pregnancy_observations(pregnancy_id: int):
    try:
        from modules.clinics.service_gynecology import add_observation, list_observations
        if request.method == "GET":
            return _ok(list_observations(_tid(), pregnancy_id))
        result = add_observation(_tid(), pregnancy_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        # POST returns the created resource; GET remains the collection shape.
        # The service keeps returning the refreshed collection for internal callers.
        return _ok(result["items"][0] if result.get("items") else result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/pregnancy-observations/<int:observation_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_pregnancy_observation_patch(observation_id: int):
    try:
        from modules.clinics.service_gynecology import update_observation
        result = update_observation(_tid(), observation_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/internal-medicine-context")
@require_clinics("clinics_clinical", json_mode=True)
def api_internal_medicine_context(patient_id: int):
    try:
        from modules.clinics.service_internal_medicine import get_internal_medicine_context
        return _ok(get_internal_medicine_context(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>/internal-medicine-context")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_internal_medicine_context(visit_id: int):
    try:
        from modules.clinics.service_internal_medicine import get_visit_internal_medicine_context
        return _ok(get_visit_internal_medicine_context(_tid(), visit_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/chronic-conditions", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_chronic_conditions(patient_id: int):
    try:
        from modules.clinics.service_internal_medicine import create_chronic_condition, list_chronic_conditions
        from modules.clinics.guards import get_current_clinic_doctor
        if request.method == "GET":
            return _ok(list_chronic_conditions(_tid(), patient_id))
        payload = request.get_json(silent=True) or {}
        linked = get_current_clinic_doctor()
        doctor_id = payload.get("doctor_id") or (linked.id if linked else None)
        result = create_chronic_condition(_tid(), patient_id, payload, _employee_id(), doctor_id=doctor_id)
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/chronic-conditions/<int:condition_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_chronic_condition_patch(condition_id: int):
    try:
        from modules.clinics.service_internal_medicine import update_chronic_condition
        result = update_chronic_condition(_tid(), condition_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/patients/<int:patient_id>/chronic-conditions/promote-diagnosis")
@require_clinics("clinics_clinical", json_mode=True)
def api_promote_diagnosis(patient_id: int):
    try:
        from modules.clinics.service_internal_medicine import promote_diagnosis_to_chronic
        from modules.clinics.guards import get_current_clinic_doctor
        payload = request.get_json(silent=True) or {}
        diagnosis_id = int(payload.get("diagnosis_id") or 0)
        if not diagnosis_id:
            raise ValueError("التشخيص مطلوب")
        linked = get_current_clinic_doctor()
        result = promote_diagnosis_to_chronic(
            _tid(), patient_id, diagnosis_id, payload, _employee_id(),
            doctor_id=payload.get("doctor_id") or (linked.id if linked else None),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/medications", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_patient_medications(patient_id: int):
    try:
        from modules.clinics.service_internal_medicine import list_patient_medications, upsert_patient_medication
        from modules.clinics.guards import get_current_clinic_doctor
        if request.method == "GET":
            return _ok(list_patient_medications(_tid(), patient_id))
        payload = request.get_json(silent=True) or {}
        linked = get_current_clinic_doctor()
        result = upsert_patient_medication(
            _tid(), patient_id, payload, _employee_id(),
            doctor_id=payload.get("doctor_id") or (linked.id if linked else None),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/patient-medications/<int:medication_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_patient_medication_patch(medication_id: int):
    try:
        from modules.clinics.service_internal_medicine import upsert_patient_medication
        from modules.clinics.models import ClinicPatientMedication
        row = ClinicPatientMedication.query.filter_by(tenant_id=_tid(), id=medication_id).first()
        if not row:
            raise ValueError("الدواء غير موجود")
        result = upsert_patient_medication(
            _tid(), row.patient_id, request.get_json(silent=True) or {}, _employee_id(),
            medication_id=medication_id,
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/vital-trends")
@require_clinics("clinics_clinical", json_mode=True)
def api_vital_trends(patient_id: int):
    try:
        from modules.clinics.service_internal_medicine import get_vital_trends
        return _ok(get_vital_trends(_tid(), patient_id, request.args.get("range", "6m")))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/lab-trends")
@require_clinics("clinics_clinical", json_mode=True)
def api_lab_trends(patient_id: int):
    try:
        from modules.clinics.service_internal_medicine import get_lab_trends
        raw = request.args.get("analytes") or ""
        analytes = [a.strip() for a in raw.split(",") if a.strip()] or None
        return _ok(get_lab_trends(_tid(), patient_id, analytes=analytes, range_key=request.args.get("range", "1y")))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/visit-comparison")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_comparison(patient_id: int):
    try:
        from modules.clinics.service_internal_medicine import compare_visits
        raw = request.args.get("visit_ids") or ""
        ids = [int(x) for x in raw.split(",") if x.strip().isdigit()]
        return _ok(compare_visits(_tid(), patient_id, ids))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/cardiology-context")
@require_clinics("clinics_clinical", json_mode=True)
def api_cardiology_context(patient_id: int):
    try:
        from modules.clinics.service_cardiology import get_cardiology_context
        return _ok(get_cardiology_context(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>/cardiology-context")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_cardiology_context(visit_id: int):
    try:
        from modules.clinics.service_cardiology import get_visit_cardiology_context
        return _ok(get_visit_cardiology_context(_tid(), visit_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/cardiology-profile", methods=["GET", "PUT"])
@require_clinics("clinics_clinical", json_mode=True)
def api_cardiology_profile(patient_id: int):
    try:
        from modules.clinics.service_cardiology import get_cardiology_profile, upsert_cardiology_profile
        if request.method == "GET":
            return _ok(get_cardiology_profile(_tid(), patient_id))
        result = upsert_cardiology_profile(_tid(), patient_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/cardiology/ecgs", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_cardiology_ecgs(patient_id: int):
    try:
        from modules.clinics.service_cardiology import create_ecg, list_ecgs
        from modules.clinics.guards import get_current_clinic_doctor
        if request.method == "GET":
            return _ok(list_ecgs(_tid(), patient_id))
        payload = request.get_json(silent=True) or {}
        linked = get_current_clinic_doctor()
        result = create_ecg(
            _tid(), patient_id, payload, _employee_id(),
            doctor_id=payload.get("doctor_id") or (linked.id if linked else None),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/cardiology/ecgs/<int:ecg_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_cardiology_ecg_patch(ecg_id: int):
    try:
        from modules.clinics.service_cardiology import update_ecg
        result = update_ecg(_tid(), ecg_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/patients/<int:patient_id>/cardiology/echos", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_cardiology_echos(patient_id: int):
    try:
        from modules.clinics.service_cardiology import create_echo, list_echos
        from modules.clinics.guards import get_current_clinic_doctor
        if request.method == "GET":
            return _ok(list_echos(_tid(), patient_id))
        payload = request.get_json(silent=True) or {}
        linked = get_current_clinic_doctor()
        result = create_echo(
            _tid(), patient_id, payload, _employee_id(),
            doctor_id=payload.get("doctor_id") or (linked.id if linked else None),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/cardiology/echos/<int:echo_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_cardiology_echo_patch(echo_id: int):
    try:
        from modules.clinics.service_cardiology import update_echo
        result = update_echo(_tid(), echo_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/cardiology/visit-comparison")
@require_clinics("clinics_clinical", json_mode=True)
def api_cardiology_visit_comparison(patient_id: int):
    try:
        from modules.clinics.service_cardiology import cardiology_visit_comparison
        raw = request.args.get("visit_ids") or ""
        ids = [int(x) for x in raw.split(",") if x.strip().isdigit()]
        return _ok(cardiology_visit_comparison(_tid(), patient_id, ids))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/specialty-attachment")
@require_clinics("clinics_clinical", json_mode=True)
def api_specialty_attachment(visit_id: int):
    try:
        from modules.clinics.models import ClinicVisit, ClinicVisitTemplateField
        from modules.clinics.service_extras import save_attachment, _detected_mime
        from modules.clinics.service_specialties import save_visit_values
        visit = ClinicVisit.query.filter_by(tenant_id=_tid(), id=visit_id).first()
        if not visit or visit.status == "completed":
            raise ValueError("الزيارة غير متاحة للتعديل")
        field_id = int(request.form.get("field_id") or 0)
        field = ClinicVisitTemplateField.query.filter_by(tenant_id=_tid(), id=field_id, template_id=visit.visit_template_id, is_active=True).first()
        upload = request.files.get("file")
        if not field or field.field_type not in {"file", "image"} or not upload:
            raise ValueError("حقل المرفق أو الملف غير صالح")
        if field.field_type == "image" and not str(_detected_mime(upload.stream)).startswith("image/"):
            raise ValueError("هذا الحقل يقبل صورة فقط")
        upload.stream.seek(0)
        attachment = save_attachment(_tid(), visit.patient_id, upload, {"visit_id": visit.id}, _employee_id())
        save_visit_values(_tid(), visit.id, {"values": {str(field.id): attachment["id"]}}, _employee_id())
        db.session.commit()
        return _ok(attachment)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/diagnoses")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_diagnosis(visit_id: int):
    try:
        result = service.add_diagnosis(
            _tid(), visit_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/prescription")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_rx(visit_id: int):
    try:
        result = service.save_prescription(
            _tid(), visit_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>/prescription-templates")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_prescription_templates(visit_id: int):
    try:
        from modules.clinics.service_prescription_templates import list_for_visit
        return _ok(list_for_visit(_tid(), visit_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/prescription-templates")
@require_clinics("clinics_clinical", json_mode=True)
def api_prescription_template_create():
    try:
        from modules.clinics.service_prescription_templates import create_template
        result = create_template(_tid(), request.get_json(silent=True) or {}, _employee_id(), getattr(g, "clinics_doctor_scope", None))
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.delete("/api/prescription-templates/<int:template_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_prescription_template_archive(template_id: int):
    try:
        from modules.clinics.service_prescription_templates import archive_template
        result = archive_template(_tid(), template_id, _employee_id(), getattr(g, "clinics_doctor_scope", None))
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/complete")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_complete(visit_id: int):
    try:
        from modules.clinics.service_ops import complete_visit_extended

        payload = request.get_json(silent=True) or {}
        discount = int(payload.get("discount") or 0)
        result = complete_visit_extended(
            _tid(),
            visit_id,
            _employee_id(),
            discount=discount,
            create_followup=payload.get("create_followup", True) is not False,
            include_financial=can_use_clinics("clinics_billing"),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/visits/<int:visit_id>/treatment-bill")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_treatment_bill(visit_id: int):
    try:
        from modules.clinics.service_ops import visit_treatment_bill_preview

        return _ok(
            visit_treatment_bill_preview(
                _tid(), visit_id, include_financial=can_use_clinics("clinics_billing")
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/visits/<int:visit_id>/treatment-bill/lines")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_treatment_bill_add(visit_id: int):
    try:
        from modules.clinics.service_ops import add_visit_bill_line

        result = add_visit_bill_line(
            _tid(), visit_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.delete("/api/visits/<int:visit_id>/treatment-bill/lines/<int:line_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_treatment_bill_delete(visit_id: int, line_id: int):
    try:
        from modules.clinics.service_ops import delete_visit_bill_line

        delete_visit_bill_line(_tid(), visit_id, line_id)
        db.session.commit()
        return _ok({"deleted": line_id})
    except Exception as exc:
        return _error(exc)


# ── Billing ────────────────────────────────────────────────────────────────


@clinics_bp.get("/api/invoices")
@require_clinics("clinics_billing", json_mode=True)
def api_invoices():
    try:
        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        return _ok(service.list_invoices(_tid(), parsed))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/invoices/<int:invoice_id>/pay")
@require_clinics("clinics_billing", json_mode=True)
def api_invoice_pay(invoice_id: int):
    try:
        result = service.receive_payment(
            _tid(),
            invoice_id,
            request.get_json(silent=True) or {},
            _employee_id(),
            _branch_id(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/invoices/<int:invoice_id>/void")
@require_clinics("clinics_billing", json_mode=True)
def api_invoice_void(invoice_id: int):
    try:
        result = service.void_invoice(_tid(), invoice_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/patients/<int:patient_id>/statement")
@require_clinics("clinics_billing", json_mode=True)
def api_patient_statement(patient_id: int):
    try:
        from modules.clinics.service_ops import patient_statement

        return _ok(patient_statement(_tid(), patient_id))
    except Exception as exc:
        return _error(exc)


# ── Schedules / Rooms / Time off ───────────────────────────────────────────


@clinics_bp.route("/api/schedules", methods=["GET", "POST"])
@require_clinics("clinics_admin", json_mode=True)
def api_schedules():
    try:
        from modules.clinics.service_ops import list_schedules, upsert_schedule

        if request.method == "GET":
            doctor_id = request.args.get("doctor_id")
            return _ok(list_schedules(_tid(), int(doctor_id) if doctor_id else None))
        result = upsert_schedule(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.delete("/api/schedules/<int:schedule_id>")
@require_clinics("clinics_admin", json_mode=True)
def api_schedule_delete(schedule_id: int):
    try:
        from modules.clinics.service_ops import delete_schedule

        delete_schedule(_tid(), schedule_id)
        db.session.commit()
        return _ok({"deleted": True})
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/time-off", methods=["GET", "POST"])
@require_clinics("clinics_admin", json_mode=True)
def api_time_off():
    try:
        from modules.clinics.service_ops import create_time_off, list_time_off

        if request.method == "GET":
            doctor_id = request.args.get("doctor_id")
            return _ok(list_time_off(_tid(), int(doctor_id) if doctor_id else None))
        result = create_time_off(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/rooms/<int:room_id>", methods=["PUT", "PATCH"])
@require_clinics("clinics_admin", json_mode=True)
def api_room_update(room_id: int):
    try:
        from modules.clinics.service_ops import update_room

        result = update_room(_tid(), room_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Inventory ──────────────────────────────────────────────────────────────


@clinics_bp.route("/api/inventory", methods=["GET", "POST"])
@require_clinics("clinics_inventory", json_mode=True)
def api_inventory():
    try:
        from modules.clinics.service_ops import create_inventory_item, list_inventory

        include_cost = can_use_clinics("purchases.cost.view") or can_use_clinics("clinics_billing")
        if request.method == "GET":
            if request.args.get("center") == "1":
                from modules.clinics.service_ops import inventory_center
                return _ok(inventory_center(_tid(), request.args.to_dict(), include_cost=include_cost))
            return _ok(
                list_inventory(
                    _tid(),
                    q=request.args.get("q", ""),
                    low_only=request.args.get("low") == "1",
                    include_cost=include_cost,
                )
            )
        result = create_inventory_item(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/inventory/<int:item_id>", methods=["GET", "PUT", "PATCH"])
@require_clinics("clinics_inventory", json_mode=True)
def api_inventory_item(item_id: int):
    try:
        from modules.clinics.service_ops import inventory_item_detail, update_inventory_item
        if request.method == "GET":
            include_cost = can_use_clinics("purchases.cost.view") or can_use_clinics("clinics_billing")
            return _ok(inventory_item_detail(_tid(), item_id, include_cost=include_cost))
        result = update_inventory_item(_tid(), item_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/inventory/<int:item_id>/stock-in")
@require_clinics("clinics_inventory", json_mode=True)
def api_inventory_stock_in(item_id: int):
    try:
        from modules.clinics.service_ops import stock_in

        result = stock_in(_tid(), item_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/inventory/<int:item_id>/stock-out")
@require_clinics("clinics_inventory", json_mode=True)
def api_inventory_stock_out(item_id: int):
    try:
        from modules.clinics.service_ops import issue_stock
        result = issue_stock(_tid(), item_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/inventory/batches")
@require_clinics("clinics_inventory", json_mode=True)
def api_inventory_batches():
    try:
        from modules.clinics.service_ops import list_inventory_batches
        item_id = request.args.get("item_id")
        include_cost = can_use_clinics("purchases.cost.view") or can_use_clinics("clinics_billing")
        return _ok(
            list_inventory_batches(
                _tid(),
                int(item_id) if item_id else None,
                request.args.get("status", ""),
                include_cost=include_cost,
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/inventory/movements")
@require_clinics("clinics_inventory", json_mode=True)
def api_inventory_movements():
    try:
        from modules.clinics.service_ops import list_movements

        item_id = request.args.get("item_id")
        include_cost = can_use_clinics("purchases.cost.view") or can_use_clinics("clinics_billing")
        return _ok(list_movements(_tid(), int(item_id) if item_id else None, include_cost=include_cost))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/inventory/alerts")
@require_clinics("clinics_inventory", json_mode=True)
def api_inventory_alerts():
    try:
        from modules.clinics.service_ops import inventory_alerts

        return _ok(inventory_alerts(_tid()))
    except Exception as exc:
        return _error(exc)


# ── Procedures ─────────────────────────────────────────────────────────────


@clinics_bp.route("/api/procedures", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_procedures():
    try:
        from modules.clinics.service_ops import list_procedures
        from modules.clinics.service_clinical_ops import procedures_center, save_procedure

        if request.method == "GET":
            if request.args.get("center") == "1":
                return _ok(procedures_center(_tid(), request.args.to_dict(), include_financial=can_use_clinics("clinics_billing")))
            specialty_id = request.args.get("specialty_id", type=int)
            if request.args.get("visit_id", type=int):
                from modules.clinics.models import ClinicVisit
                visit = ClinicVisit.query.filter_by(tenant_id=_tid(), id=request.args.get("visit_id", type=int)).first()
                if not visit: raise ValueError("الزيارة غير موجودة")
                specialty_id = visit.specialty_id
            return _ok(list_procedures(_tid(), active_only=request.args.get("active") != "0", specialty_id=specialty_id))
        result = save_procedure(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/visits/<int:visit_id>/procedures", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_visit_procedures(visit_id: int):
    try:
        from modules.clinics.service_ops import add_visit_procedure, list_visit_procedures

        if request.method == "GET":
            return _ok(list_visit_procedures(_tid(), visit_id))
        result = add_visit_procedure(
            _tid(), visit_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Labs / Radiology ───────────────────────────────────────────────────────


@clinics_bp.route("/api/lab-tests", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_lab_tests():
    try:
        from modules.clinics.service_ops import create_lab_test, list_lab_tests

        if request.method == "GET":
            return _ok(list_lab_tests(_tid()))
        result = create_lab_test(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/lab-orders", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_lab_orders():
    try:
        from modules.clinics.service_ops import create_lab_order, list_lab_orders

        if request.method == "GET":
            if request.args.get("center") == "1":
                from modules.clinics.service_clinical_ops import lab_center
                return _ok(lab_center(_tid(), request.args.to_dict()))
            visit_id = request.args.get("visit_id")
            return _ok(
                list_lab_orders(
                    _tid(),
                    status=request.args.get("status"),
                    visit_id=int(visit_id) if visit_id else None,
                )
            )
        result = create_lab_order(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/lab-order-items/<int:item_id>/result")
@require_clinics("clinics_clinical", json_mode=True)
def api_lab_result(item_id: int):
    try:
        from modules.clinics.service_ops import update_lab_result

        result = update_lab_result(
            _tid(), item_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/radiology-orders", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_radiology_orders():
    try:
        from modules.clinics.service_ops import create_radiology_order, list_radiology_orders

        if request.method == "GET":
            if request.args.get("center") == "1":
                from modules.clinics.service_clinical_ops import radiology_center
                return _ok(radiology_center(_tid(), request.args.to_dict()))
            visit_id = request.args.get("visit_id")
            return _ok(
                list_radiology_orders(
                    _tid(),
                    status=request.args.get("status"),
                    visit_id=int(visit_id) if visit_id else None,
                )
            )
        result = create_radiology_order(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/radiology-orders/<int:order_id>/result")
@require_clinics("clinics_clinical", json_mode=True)
def api_radiology_result(order_id: int):
    try:
        from modules.clinics.service_ops import update_radiology_result

        result = update_radiology_result(
            _tid(), order_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Settlements / Reports / Public display ─────────────────────────────────


@clinics_bp.route("/api/settlements", methods=["GET", "POST"])
@require_clinics("clinics_billing", json_mode=True)
def api_settlements():
    try:
        from modules.clinics.service_ops import create_settlement, doctor_earnings, list_settlements

        if request.method == "GET":
            doctor_id = request.args.get("doctor_id")
            if doctor_id and request.args.get("period_from") and request.args.get("period_to"):
                return _ok(
                    doctor_earnings(
                        _tid(),
                        int(doctor_id),
                        service._parse_date(request.args.get("period_from"), "من"),
                        service._parse_date(request.args.get("period_to"), "إلى"),
                    )
                )
            return _ok(list_settlements(_tid()))
        result = create_settlement(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/reports/summary")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_reports_summary():
    try:
        from modules.clinics.service_ops import reports_summary

        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        return _ok(reports_summary(_tid(), parsed))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/reports/hub")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_reports_hub():
    try:
        from modules.clinics.service_reports import reports_hub

        return _ok(reports_hub(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


# ── Clinic accounting bridge ──────────────────────────────────────────────


@clinics_bp.get("/api/accounting/reconciliation")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_accounting_reconciliation():
    try:
        from modules.clinics.accounting_service import clinic_accounting_reconciliation
        return _ok(clinic_accounting_reconciliation(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/accounting/postings")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_accounting_postings():
    try:
        from modules.clinics.accounting_service import list_accounting_postings
        return _ok(list_accounting_postings(
            _tid(), status=request.args.get("status"), source_type=request.args.get("source_type"),
            limit=request.args.get("limit", 200, type=int),
        ))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/accounting/postings/<int:posting_id>")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_accounting_posting_detail(posting_id: int):
    try:
        from modules.clinics.accounting_service import posting_dict
        from modules.clinics.models import AccountingPosting
        row = AccountingPosting.query.filter_by(tenant_id=_tid(), id=posting_id).first()
        if not row:
            raise ValueError("سجل الترحيل غير موجود")
        payload = posting_dict(row)
        entry = row.journal_entry
        payload["journal_entry"] = None if not entry else {
            "id": entry.id, "entry_number": entry.entry_number, "description": entry.description,
            "debit_account": entry.debit_account.code, "credit_account": entry.credit_account.code,
            "amount": entry.amount, "entry_date": entry.entry_date.isoformat(),
        }
        return _ok(payload)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/accounting/repost/<int:posting_id>")
@require_clinics("clinics_billing", json_mode=True)
def api_accounting_repost(posting_id: int):
    try:
        from modules.clinics.accounting_service import ClinicAccountingPostingService
        entry = ClinicAccountingPostingService.repost(_tid(), posting_id, _employee_id())
        db.session.commit()
        return _ok({"journal_entry_id": entry.id if entry else None})
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/accounting/migration/preview")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_accounting_migration_preview():
    try:
        from modules.clinics.accounting_service import ClinicAccountingMigrationService
        return _ok(ClinicAccountingMigrationService.preview(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/accounting/migration/run")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_accounting_migration_run():
    try:
        from modules.clinics.accounting_service import ClinicAccountingMigrationService
        result = ClinicAccountingMigrationService.run(_tid(), _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/accounting/migration/validate")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_accounting_migration_validate():
    try:
        from modules.clinics.accounting_service import ClinicAccountingMigrationService
        return _ok(ClinicAccountingMigrationService.validate(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/reports/export")
@require_clinics("clinics_reports", json_mode=True)
@require_clinics("clinics_billing", json_mode=True)
def api_reports_export():
    try:
        import csv
        import io

        from flask import Response
        from modules.clinics.service_reports import reports_export_workbook_data

        data = reports_export_workbook_data(_tid(), request.args.to_dict())
        fmt = str(request.args.get("format") or "xlsx").lower()
        base = data["filename_base"]

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            for sheet_name in ("summary", "by_day", "by_doctor", "payments", "services", "specialties"):
                writer.writerow([f"=== {sheet_name} ==="])
                for row in data[sheet_name]:
                    writer.writerow(row)
                writer.writerow([])
            return Response(
                "\ufeff" + buf.getvalue(),
                mimetype="text/csv; charset=utf-8",
                headers={"Content-Disposition": f"attachment; filename={base}.csv"},
            )

        if fmt in ("xlsx", "excel"):
            from io import BytesIO

            from openpyxl import Workbook

            wb = Workbook()
            titles = {
                "summary": "ملخص",
                "by_day": "حسب اليوم",
                "by_doctor": "حسب الطبيب",
                "specialties": "حسب الاختصاص",
                "payments": "طرق الدفع",
                "services": "الخدمات",
            }
            first = True
            for key, title in titles.items():
                ws = wb.active if first else wb.create_sheet(title)
                if first:
                    ws.title = title
                    first = False
                for row in data[key]:
                    ws.append(list(row))
            xbuf = BytesIO()
            wb.save(xbuf)
            return Response(
                xbuf.getvalue(),
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={base}.xlsx"},
            )
        return _error(ValueError("صيغة التصدير غير مدعومة"))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/public-queue")
@require_clinics("clinics_reception", json_mode=True)
def api_public_queue():
    try:
        from modules.clinics.service_ops import public_queue_display

        day = request.args.get("date")
        parsed = service._parse_date(day, "التاريخ") if day else None
        return _ok(public_queue_display(_tid(), parsed))
    except Exception as exc:
        return _error(exc)


# ── Communications / Follow-ups ────────────────────────────────────────────


@clinics_bp.get("/api/communications")
@require_clinics("clinics_communications_view", json_mode=True)
def api_comms_dashboard():
    try:
        from modules.clinics.service_engagement import engagement_snapshot

        data = engagement_snapshot(_tid())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/communications/conversations", methods=["GET", "POST"])
@require_clinics("clinics_communications_view", json_mode=True)
def api_conversations():
    try:
        from modules.clinics.service_engagement import create_conversation, list_conversations
        if request.method == "GET":
            return _ok(list_conversations(_tid(), request.args.to_dict(), can_use_clinics("clinics_patient_contact_sensitive_view")))
        if not can_use_clinics("clinics_communications_reply"):
            return _error("ليس لديك صلاحية إنشاء محادثة", 403)
        row = create_conversation(_tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id())
        db.session.commit()
        return _ok(row)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/communications/conversations/<int:conversation_id>", methods=["GET", "PUT"])
@require_clinics("clinics_communications_view", json_mode=True)
def api_conversation(conversation_id: int):
    try:
        from modules.clinics.service_engagement import conversation_detail, update_conversation
        if request.method == "GET":
            data = conversation_detail(_tid(), conversation_id, _employee_id(), can_use_clinics("clinics_patient_contact_sensitive_view"))
        else:
            payload = request.get_json(silent=True) or {}
            needed = "clinics_communications_assign" if "assigned_user_id" in payload else "clinics_communications_close"
            if not can_use_clinics(needed):
                return _error("ليس لديك صلاحية هذا الإجراء", 403)
            data = update_conversation(_tid(), conversation_id, payload, _employee_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/communications/conversations/<int:conversation_id>/messages")
@require_clinics("clinics_communications_reply", json_mode=True)
def api_conversation_message(conversation_id: int):
    try:
        from modules.clinics.service_engagement import send_message
        data = send_message(_tid(), conversation_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/communications/conversations/<int:conversation_id>/patient")
@require_clinics("clinics_patients", json_mode=True)
def api_conversation_patient(conversation_id: int):
    try:
        from modules.clinics.service_engagement import attach_patient
        data = attach_patient(_tid(), conversation_id, request.get_json(silent=True) or {}, _employee_id(), _branch_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/communications/conversations/<int:conversation_id>/followups")
@require_clinics(json_mode=True)
def api_conversation_followup(conversation_id: int):
    try:
        # Reception/comms staff create callback follow-ups from chat; clinical is also enough.
        if not (can_use_clinics("clinics_communications_reply") or can_use_clinics("clinics_clinical")):
            return _error("ليس لديك صلاحية إنشاء متابعة من المحادثة", 403)
        from modules.clinics.service_engagement import create_followup
        from sqlalchemy.exc import IntegrityError

        data = create_followup(_tid(), conversation_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(data)
    except IntegrityError:
        db.session.rollback()
        return _error("تعذر حفظ المتابعة — تحقق من ربط المريض أو أعد المحاولة")
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/communications/conversations/<int:conversation_id>/tasks")
@require_clinics("clinics_communications_reply", json_mode=True)
def api_conversation_task(conversation_id: int):
    try:
        from modules.clinics.service_engagement import create_task
        data = create_task(_tid(), conversation_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/communications/templates", methods=["GET", "POST"])
@require_clinics("clinics_communications_view", json_mode=True)
def api_communication_templates():
    try:
        from modules.clinics.service_engagement import save_template, templates
        if request.method == "GET":
            data = templates(_tid())
        else:
            if not can_use_clinics("clinics_templates_manage"):
                return _error("ليس لديك صلاحية إدارة القوالب", 403)
            data = save_template(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/communications/booking-requests")
@require_clinics("clinics_booking_requests_view", json_mode=True)
def api_booking_requests():
    try:
        from modules.clinics.service_engagement import booking_requests
        return _ok(booking_requests(_tid(), request.args.get("status")))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/communications/booking-requests/<int:booking_request_id>/<action>")
@require_clinics("clinics_booking_requests_approve", json_mode=True)
def api_booking_request_decision(booking_request_id: int, action: str):
    try:
        from modules.clinics.service_engagement import decide_booking_request
        data = decide_booking_request(_tid(), booking_request_id, action, request.get_json(silent=True) or {}, _employee_id(), _branch_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/communications/reminders")
@require_clinics("clinics_reminders_view", json_mode=True)
def api_reminder_center():
    try:
        from modules.clinics.service_engagement import reminder_center
        return _ok(reminder_center(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/communications/reminder-rules")
@require_clinics("clinics_reminders_manage", json_mode=True)
def api_reminder_rule_save():
    try:
        from modules.clinics.service_engagement import save_reminder_rule
        data = save_reminder_rule(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/communications/settings")
@require_clinics("clinics_booking_settings_manage", json_mode=True)
def api_communication_settings():
    """Restricted settings surface for booking and public contact options."""
    try:
        payload = request.get_json(silent=True) or {}
        allowed = {
            "booking_mode", "minimum_booking_notice_hours", "maximum_advance_booking_days",
            "allow_same_day_booking", "allow_public_reschedule", "allow_public_cancellation",
            "cancellation_deadline_hours", "max_active_appointments_per_patient", "public_show_prices",
            "public_profile_enabled", "public_chat_enabled", "public_logo_url", "public_location",
            "public_working_hours", "clinic_name",
        }
        if can_use_clinics("clinics_whatsapp_manage"):
            allowed.add("clinic_whatsapp_number")
            allowed.add("whatsapp_notifications_enabled")
        clean = {key: value for key, value in payload.items() if key in allowed}
        data = service.update_settings(_tid(), clean, _employee_id())
        db.session.commit()
        return _ok(data)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/messages")
@require_clinics("clinics_settings", json_mode=True)
def api_messages():
    try:
        from modules.clinics.service_comms import list_message_logs

        return _ok(list_message_logs(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/reminders/process")
@require_clinics("clinics_settings", json_mode=True)
def api_process_reminders():
    try:
        from modules.clinics.service_comms import process_pending_reminders, schedule_appointment_reminders

        scheduled = schedule_appointment_reminders(_tid())
        result = process_pending_reminders(_tid())
        db.session.commit()
        return _ok({**result, "scheduled": scheduled})
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/followups", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_followups():
    try:
        from modules.clinics.service_comms import create_followup_from_visit, list_followups

        if request.method == "GET":
            if request.args.get("center") == "1":
                from modules.clinics.service_clinical_ops import followups_center
                return _ok(followups_center(_tid(), request.args.to_dict()))
            return _ok(list_followups(_tid(), request.args.get("status")))
        payload = request.get_json(silent=True) or {}
        if payload.get("patient_id"):
            from modules.clinics.service_clinical_ops import StateConflict, create_followup
            try:
                result = create_followup(_tid(), payload, _employee_id())
            except StateConflict as exc:
                return _error(exc, 409)
            db.session.commit()
            return _ok(result)
        visit_id = int(payload.get("visit_id") or 0)
        if not visit_id:
            raise ValueError("visit_id مطلوب")
        result = create_followup_from_visit(_tid(), visit_id, payload, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/followups/<int:followup_id>", methods=["PUT", "PATCH"])
@require_clinics("clinics_clinical", json_mode=True)
def api_followup_update(followup_id: int):
    try:
        from modules.clinics.service_comms import update_followup

        result = update_followup(_tid(), followup_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Suppliers / Purchases ──────────────────────────────────────────────────


@clinics_bp.route("/api/suppliers", methods=["GET", "POST"])
@require_clinics("clinics_inventory", json_mode=True)
def api_suppliers():
    try:
        from modules.clinics.service_suppliers import create_supplier, list_suppliers

        if request.method == "GET":
            return _ok(list_suppliers(_tid(), active_only=request.args.get("active") != "0"))
        result = create_supplier(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/suppliers/<int:supplier_id>/statement")
@require_clinics("clinics_inventory", json_mode=True)
def api_supplier_statement(supplier_id: int):
    try:
        from modules.clinics.service_suppliers import supplier_statement

        return _ok(supplier_statement(_tid(), supplier_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/meta")
@require_clinics("purchases.view", json_mode=True)
def api_purchases_meta():
    from modules.clinics.guards import clinics_purchase_permissions

    return _ok({"permissions": clinics_purchase_permissions()})


@clinics_bp.get("/api/purchases/reports/by-item")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports_item():
    try:
        from modules.clinics.service_purchase_extras import purchase_reports_by_item

        return _ok(purchase_reports_by_item(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reports/by-category")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports_category():
    try:
        from modules.clinics.service_purchase_extras import purchase_reports_by_category

        return _ok(purchase_reports_by_category(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reports/pending-receipt")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports_pending():
    try:
        from modules.clinics.service_purchase_extras import purchase_reports_pending_receipt

        return _ok(purchase_reports_pending_receipt(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reports/returns")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports_returns():
    try:
        from modules.clinics.service_purchase_extras import purchase_reports_returns

        return _ok(purchase_reports_returns(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reports/price-changes")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports_price_changes():
    try:
        from modules.clinics.service_purchase_extras import purchase_reports_price_changes

        return _ok(purchase_reports_price_changes(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/kpis")
@require_clinics("purchases.view", json_mode=True)
def api_purchases_kpis():
    try:
        from modules.clinics.service_suppliers import purchase_kpis

        return _ok(purchase_kpis(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/check-duplicate")
@require_clinics("purchases.create", json_mode=True)
def api_purchases_check_duplicate():
    try:
        from modules.clinics.service_suppliers import check_duplicate_invoice

        return _ok(
            check_duplicate_invoice(
                _tid(),
                int(request.args.get("supplier_id") or 0),
                request.args.get("supplier_invoice_no") or "",
                request.args.get("exclude_id"),
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/purchases", methods=["GET", "POST"])
@require_clinics("purchases.view", json_mode=True)
def api_purchases():
    try:
        from modules.clinics.guards import can_use_clinics
        from modules.clinics.service_suppliers import create_purchase, list_purchases

        if request.method == "GET":
            return _ok(list_purchases(_tid(), request.args.to_dict()))
        if not can_use_clinics("purchases.create"):
            return jsonify({"ok": False, "error": "ليس لديك صلاحية لهذه العملية."}), 403
        result = create_purchase(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reports/unpaid")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports_unpaid():
    try:
        from modules.clinics.service_purchase_extras import purchase_reports_unpaid

        return _ok(purchase_reports_unpaid(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reports/by-supplier")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports_supplier():
    try:
        from modules.clinics.service_purchase_extras import purchase_reports_by_supplier

        return _ok(purchase_reports_by_supplier(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/price-history")
@require_clinics("purchases.cost.view", json_mode=True)
def api_purchase_price_history():
    try:
        from modules.clinics.service_purchase_extras import item_purchase_price_history

        item_id = request.args.get("item_id")
        return _ok(item_purchase_price_history(_tid(), int(item_id) if item_id else None))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/price-check")
@require_clinics("purchases.create", json_mode=True)
def api_purchase_price_check():
    try:
        from modules.clinics.service_purchase_extras import check_price_warnings

        payload = request.get_json(silent=True) or {}
        return _ok({"warnings": check_price_warnings(_tid(), payload.get("items") or [])})
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reorder-suggestions")
@require_clinics("purchase_request.create", json_mode=True)
def api_purchase_reorder():
    try:
        from modules.clinics.service_purchase_extras import reorder_suggestions

        return _ok(reorder_suggestions(_tid()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/reports/summary")
@require_clinics("purchases.export", json_mode=True)
def api_purchase_reports():
    try:
        from modules.clinics.service_purchase_extras import purchase_report_summary

        return _ok(purchase_report_summary(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.put("/api/purchases/<int:purchase_id>")
@require_clinics("purchases.edit", json_mode=True)
def api_purchase_update(purchase_id: int):
    try:
        from modules.clinics.service_suppliers import update_purchase

        result = update_purchase(_tid(), purchase_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/<int:purchase_id>/copy")
@require_clinics("purchases.create", json_mode=True)
def api_purchase_copy(purchase_id: int):
    try:
        from modules.clinics.service_suppliers import copy_purchase

        result = copy_purchase(_tid(), purchase_id, _employee_id(), _branch_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/<int:purchase_id>/approve")
@require_clinics("purchases.approve", json_mode=True)
def api_purchase_approve(purchase_id: int):
    try:
        from modules.clinics.service_suppliers import approve_purchase

        result = approve_purchase(_tid(), purchase_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/<int:purchase_id>/validate-receive")
@require_clinics("purchases.receive", json_mode=True)
def api_purchase_validate_receive(purchase_id: int):
    try:
        from modules.clinics.service_purchase_extras import validate_receive_expiry
        from modules.clinics.service_suppliers import get_purchase

        purchase = get_purchase(_tid(), purchase_id)
        item_map = {i["id"]: i for i in purchase.get("items", [])}
        lines = []
        for raw in (request.get_json(silent=True) or {}).get("items") or []:
            pit = item_map.get(int(raw.get("purchase_item_id") or 0))
            if pit:
                lines.append({**raw, "item_id": pit["item_id"]})
        return _ok(validate_receive_expiry(_tid(), lines))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/<int:purchase_id>")
@require_clinics("purchases.view", json_mode=True)
def api_purchase_detail(purchase_id: int):
    try:
        from modules.clinics.service_suppliers import get_purchase

        return _ok(get_purchase(_tid(), purchase_id))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/<int:purchase_id>/payments")
@require_clinics("purchases.payment.create", json_mode=True)
def api_purchase_payment(purchase_id: int):
    try:
        from modules.clinics.service_suppliers import record_purchase_payment

        result = record_purchase_payment(
            _tid(), purchase_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/<int:purchase_id>/cancel")
@require_clinics("purchases.cancel", json_mode=True)
def api_purchase_cancel(purchase_id: int):
    try:
        from modules.clinics.service_suppliers import cancel_purchase

        result = cancel_purchase(_tid(), purchase_id, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/<int:purchase_id>/receive")
@require_clinics("purchases.receive", json_mode=True)
def api_purchase_receive(purchase_id: int):
    try:
        from modules.clinics.service_suppliers import receive_purchase
        result = receive_purchase(_tid(), purchase_id, request.get_json(silent=True) or {}, _employee_id(), _branch_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Purchase requests / returns / attachments ───────────────────────────────


@clinics_bp.get("/api/purchase-requests")
@require_clinics("purchase_request.create", json_mode=True)
def api_purchase_requests():
    try:
        from modules.clinics.service_purchase_extras import list_purchase_requests

        return _ok(list_purchase_requests(_tid(), request.args.to_dict()))
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchase-requests")
@require_clinics("purchase_request.create", json_mode=True)
def api_purchase_requests_create():
    try:
        from modules.clinics.service_purchase_extras import create_purchase_request

        result = create_purchase_request(_tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchase-requests/<int:request_id>/status")
@require_clinics("purchase_request.approve", json_mode=True)
def api_purchase_request_status(request_id: int):
    try:
        from modules.clinics.service_purchase_extras import update_purchase_request_status

        payload = request.get_json(silent=True) or {}
        result = update_purchase_request_status(
            _tid(), request_id, payload.get("status"), _employee_id(), reject_reason=payload.get("reason") or ""
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchase-requests/<int:request_id>/convert")
@require_clinics("purchases.create", json_mode=True)
def api_purchase_request_convert(request_id: int):
    try:
        from modules.clinics.service_purchase_extras import convert_request_to_purchase

        result = convert_request_to_purchase(
            _tid(), request_id, request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/purchases/<int:purchase_id>/returns")
@require_clinics("purchases.return", json_mode=True)
def api_purchase_return(purchase_id: int):
    try:
        from modules.clinics.service_purchase_extras import create_purchase_return

        result = create_purchase_return(
            _tid(), purchase_id, request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/purchases/<int:purchase_id>/attachments", methods=["GET", "POST"])
@require_clinics("purchases.view", json_mode=True)
def api_purchase_attachments(purchase_id: int):
    try:
        from modules.clinics.guards import can_use_clinics
        from modules.clinics.service_purchase_extras import list_purchase_attachments, save_purchase_attachment

        if request.method == "GET":
            return _ok(list_purchase_attachments(_tid(), purchase_id))
        if not can_use_clinics("purchases.edit"):
            return jsonify({"ok": False, "error": "ليس لديك صلاحية لهذه العملية."}), 403
        result = save_purchase_attachment(
            _tid(),
            purchase_id,
            request.files.get("file"),
            {
                "category": request.form.get("category"),
                "notes": request.form.get("notes"),
            },
            _employee_id(),
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/purchases/attachments/<int:attachment_id>/download")
@require_clinics("purchases.view", json_mode=True)
def api_purchase_attachment_download(attachment_id: int):
    try:
        from flask import send_file

        from modules.clinics.service_purchase_extras import purchase_attachment_path

        row, path = purchase_attachment_path(_tid(), attachment_id)
        return send_file(path, as_attachment=True, download_name=row.file_name)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/purchases/<int:purchase_id>/print")
@require_clinics("purchases.view")
def purchase_print_page(purchase_id: int):
    from modules.clinics.service_suppliers import get_purchase

    purchase = get_purchase(_tid(), purchase_id)
    return render_template("clinics/purchase_print.html", purchase=purchase)


@clinics_bp.get("/purchases/<int:purchase_id>/receipts/<int:receipt_id>/print")
@require_clinics("purchases.view")
def purchase_receipt_print_page(purchase_id: int, receipt_id: int):
    from modules.clinics.service_suppliers import get_purchase_receipt_print

    receipt = get_purchase_receipt_print(_tid(), purchase_id, receipt_id)
    return render_template("clinics/purchase_receipt_print.html", receipt=receipt)


@clinics_bp.get("/purchases/<int:purchase_id>/payments/<int:payment_id>/print")
@require_clinics("purchases.payment.view")
def purchase_payment_print_page(purchase_id: int, payment_id: int):
    from modules.clinics.service_suppliers import get_purchase_payment_print

    payment = get_purchase_payment_print(_tid(), purchase_id, payment_id)
    return render_template("clinics/purchase_payment_print.html", payment=payment)


@clinics_bp.get("/purchases/<int:purchase_id>/returns/<int:return_id>/print")
@require_clinics("purchases.view")
def purchase_return_print_page(purchase_id: int, return_id: int):
    from modules.clinics.service_purchase_extras import get_purchase_return_print

    ret = get_purchase_return_print(_tid(), purchase_id, return_id)
    return render_template("clinics/purchase_return_print.html", ret=ret)


@clinics_bp.get("/suppliers/<int:supplier_id>/statement/print")
@require_clinics("purchases.view")
def supplier_statement_print_page(supplier_id: int):
    from modules.clinics.service_suppliers import supplier_statement

    data = supplier_statement(_tid(), supplier_id)
    return render_template("clinics/supplier_statement_print.html", data=data)

# ── Calendar / Reschedule ──────────────────────────────────────────────────


@clinics_bp.get("/api/calendar")
@require_clinics("clinics_appointments", json_mode=True)
def api_calendar():
    try:
        start = service._parse_date(request.args.get("start") or service._today().isoformat(), "من")
        end_raw = request.args.get("end")
        end = service._parse_date(end_raw, "إلى") if end_raw else start
        doctor_id = request.args.get("doctor_id")
        room_id = request.args.get("room_id")
        return _ok(
            service.list_appointments_range(
                _tid(),
                start,
                end,
                int(doctor_id) if doctor_id else None,
                int(room_id) if room_id else None,
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/appointments/<int:appointment_id>/reschedule")
@require_clinics("clinics_appointments", json_mode=True)
def api_reschedule(appointment_id: int):
    try:
        result = service.reschedule_appointment(
            _tid(), appointment_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Expenses / Refunds / Cash ──────────────────────────────────────────────


@clinics_bp.route("/api/expenses", methods=["GET", "POST"])
@require_clinics("clinics_billing", json_mode=True)
def api_expenses():
    try:
        from modules.clinics.service_extras import create_expense, list_expenses

        if request.method == "GET":
            day = request.args.get("date")
            parsed = service._parse_date(day, "التاريخ") if day else None
            return _ok(list_expenses(_tid(), parsed))
        result = create_expense(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/invoices/<int:invoice_id>/refund")
@require_clinics("clinics_billing", json_mode=True)
def api_refund(invoice_id: int):
    try:
        from modules.clinics.service_extras import refund_invoice

        result = refund_invoice(
            _tid(), invoice_id, request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/cash-sessions", methods=["GET", "POST"])
@require_clinics("clinics_billing", json_mode=True)
def api_cash_sessions():
    try:
        from modules.clinics.service_extras import list_cash_sessions, open_cash_session

        if request.method == "GET":
            return _ok(list_cash_sessions(_tid()))
        result = open_cash_session(
            _tid(), request.get_json(silent=True) or {}, _employee_id(), _branch_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/cash-sessions/<int:session_id>/close")
@require_clinics("clinics_billing", json_mode=True)
def api_cash_close(session_id: int):
    try:
        from modules.clinics.service_extras import close_cash_session

        result = close_cash_session(
            _tid(), session_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


# ── Attachments / Consent / Specialty / QR / Jobs ──────────────────────────


@clinics_bp.route("/api/attachments", methods=["GET", "POST"])
@require_clinics("clinics_patients", json_mode=True)
def api_attachments():
    try:
        from modules.clinics.service_extras import list_attachments, save_attachment

        if request.method == "GET":
            return _ok(list_attachments(_tid(), request.args.to_dict()))
        patient_id = int(request.form.get("patient_id") or request.args.get("patient_id") or 0)
        if not patient_id:
            raise ValueError("patient_id مطلوب")
        payload = {
            "visit_id": request.form.get("visit_id"),
            "category": request.form.get("category") or "other",
            "document_date": request.form.get("document_date"),
            "notes": request.form.get("notes"),
        }
        files = request.files.getlist("files") or request.files.getlist("file")
        if not files: raise ValueError("الملف مطلوب")
        if len(files) > 10: raise ValueError("يمكن رفع 10 ملفات كحد أقصى في المرة الواحدة")
        result = [save_attachment(_tid(), patient_id, item, payload, _employee_id()) for item in files]
        db.session.commit()
        return _ok({"items": result, "count": len(result)})
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/attachments/<int:attachment_id>/download")
@require_clinics("clinics_patients", json_mode=True)
def api_attachment_download(attachment_id: int):
    try:
        from flask import send_file

        from modules.clinics.service_extras import attachment_path

        row, path = attachment_path(_tid(), attachment_id)
        return send_file(path, as_attachment=True, download_name=row.file_name, mimetype=row.mime_type)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/attachments/<int:attachment_id>/preview")
@require_clinics("clinics_patients", json_mode=True)
def api_attachment_preview(attachment_id: int):
    try:
        from flask import send_file
        from modules.clinics.service_extras import attachment_path
        row, path = attachment_path(_tid(), attachment_id)
        return send_file(path, as_attachment=False, download_name=row.file_name, mimetype=row.mime_type)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/followups/<int:followup_id>/contact")
@require_clinics("clinics_clinical", json_mode=True)
def api_followup_contact(followup_id: int):
    try:
        from modules.clinics.service_clinical_ops import contact_followup
        result = contact_followup(_tid(), followup_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/radiology-orders/<int:order_id>/workflow")
@require_clinics("clinics_clinical", json_mode=True)
def api_radiology_workflow(order_id: int):
    try:
        from modules.clinics.service_clinical_ops import StateConflict, update_radiology_workflow
        result = update_radiology_workflow(_tid(), order_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except StateConflict as exc:
        return _error(exc, 409)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/lab-orders/<int:order_id>/workflow")
@require_clinics("clinics_clinical", json_mode=True)
def api_lab_workflow(order_id: int):
    try:
        from modules.clinics.service_clinical_ops import StateConflict, transition_lab
        result = transition_lab(_tid(), order_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except StateConflict as exc:
        return _error(exc, 409)
    except Exception as exc:
        return _error(exc)


@clinics_bp.put("/api/procedures/<int:procedure_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_procedure_update(procedure_id: int):
    try:
        from modules.clinics.service_clinical_ops import save_procedure
        result = save_procedure(_tid(), request.get_json(silent=True) or {}, _employee_id(), procedure_id)
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/procedure-executions")
@require_clinics("clinics_clinical", json_mode=True)
def api_procedure_execution_create():
    try:
        from modules.clinics.service_clinical_ops import add_execution
        result = add_execution(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.patch("/api/procedure-executions/<int:execution_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_procedure_execution_update(execution_id: int):
    try:
        from modules.clinics.service_clinical_ops import StateConflict, transition_execution
        result = transition_execution(_tid(), execution_id, request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except StateConflict as exc:
        return _error(exc, 409)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/clinical/recent-patients")
@require_clinics("clinics_clinical", json_mode=True)
def api_clinical_recent_patients():
    try:
        from modules.clinics.service_clinical import recent_patients
        return _ok(recent_patients(_tid(), request.args.get("q", ""), request.args.get("limit", 12, type=int)))
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/clinical/patients/<int:patient_id>")
@require_clinics("clinics_clinical", json_mode=True)
def api_clinical_record(patient_id: int):
    try:
        from modules.clinics.service_clinical import clinical_record
        result = clinical_record(_tid(), patient_id, request.args.get("section", "overview"), request.args.get("page", 1, type=int), request.args.get("per_page", 20, type=int), _employee_id(), request.args.to_dict())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/clinical/patients/<int:patient_id>/print-audit")
@require_clinics("clinics_clinical", json_mode=True)
def api_clinical_print_audit(patient_id: int):
    try:
        from modules.clinics.service import _audit
        _audit(_tid(), _employee_id(), "clinical_record", patient_id, "print", "طباعة ملخص طبي")
        db.session.commit()
        return _ok({"logged": True})
    except Exception as exc:
        return _error(exc)


@clinics_bp.route("/api/consents", methods=["GET", "POST"])
@require_clinics("clinics_clinical", json_mode=True)
def api_consents():
    try:
        from modules.clinics.service_extras import create_consent, list_consents

        if request.method == "GET":
            return _ok(list_consents(_tid(), request.args.get("patient_id", type=int)))
        result = create_consent(_tid(), request.get_json(silent=True) or {}, _employee_id())
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/consents/<int:consent_id>/sign")
@require_clinics("clinics_clinical", json_mode=True)
def api_consent_sign(consent_id: int):
    try:
        from modules.clinics.service_extras import sign_consent

        result = sign_consent(
            _tid(), consent_id, request.get_json(silent=True) or {}, _employee_id()
        )
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/specialty/seed")
@require_clinics("clinics_settings", json_mode=True)
def api_specialty_seed():
    try:
        from modules.clinics.service_specialties import seed_default_specialties

        result = {"specialties": seed_default_specialties(_tid(), _employee_id()),
                  "created": {"services": 0, "procedures": 0, "labs": 0}}
        db.session.commit()
        return _ok(result)
    except Exception as exc:
        return _error(exc)


@clinics_bp.get("/api/booking-qr")
@require_clinics("clinics_settings", json_mode=True)
def api_booking_qr():
    try:
        from flask import session

        from modules.clinics.service_extras import booking_qr_payload

        slug = session.get("tenant_slug") or ""
        base = request.host_url.rstrip("/")
        return _ok(
            booking_qr_payload(
                slug,
                base,
                request.args.get("doctor_id", type=int),
                request.args.get("service_id", type=int),
            )
        )
    except Exception as exc:
        return _error(exc)


@clinics_bp.post("/api/jobs/run")
@require_clinics("clinics_settings", json_mode=True)
def api_jobs_run():
    try:
        from modules.clinics.service_jobs import process_clinic_jobs_for_current_tenant

        return _ok(process_clinic_jobs_for_current_tenant())
    except Exception as exc:
        return _error(exc)

